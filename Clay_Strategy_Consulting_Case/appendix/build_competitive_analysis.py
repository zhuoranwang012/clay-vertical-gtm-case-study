"""Build 04_Competitive_Analysis.xlsx — Clay vertical expansion case study (Phase 2).
Scores and coordinates are analyst judgments [C] with evidence; calculations are live formulas."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L

BASE = "/home/claude/Clay_Strategy_Consulting_Case"
OUT = f"{BASE}/04_Competitive_Analysis.xlsx"
F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLACK = Font(name=F, color="000000", size=10)
GREEN = Font(name=F, color="008000", size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=F, bold=True, size=14, color="1F2A44")
SUB = Font(name=F, italic=True, size=9, color="555555")
HFILL = PatternFill("solid", fgColor="1F2A44")
SFILL = PatternFill("solid", fgColor="E8ECF3")
YFILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
D1 = '0.0'
D2 = '0.00'

wb = Workbook()


def hdr(ws, row, labels, col=1, widths=None, height=30):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR, HFILL, CENTER, BOX
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[L(col + i)].width = w
    ws.row_dimensions[row].height = height


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, bold=False, wrap=True, border=True, center=False):
    c = ws[ref]
    c.value = val
    c.font = Font(name=F, size=10, bold=bold, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = CENTER if center else (WRAP if wrap else Alignment(vertical="top"))
    if border:
        c.border = BOX
    return c


def title(ws, t, s):
    ws["A1"].value, ws["A1"].font = t, TITLE
    ws["A2"].value, ws["A2"].font = s, SUB
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------ Cover
ws = wb.active
ws.title = "Cover"
title(ws, "04 — Competitive Analysis: Does Clay have a defensible advantage by segment?",
      "Independent portfolio case study by Zhuoran (Zora) Wang · Not commissioned by or affiliated with Clay Labs Inc. · Research reference date: 2026-10-18 (accessed 2026-10-19)")
rows = [
    ("Question", "Which alternatives genuinely overlap with Clay in each candidate segment, where is Clay advantaged or disadvantaged, and how strong is its competitive position going into Phase 3 prioritization?"),
    ("Answer (short)", "Clay's differentiation (multi-source and off-LinkedIn data plus AI research) matters most in industrial distribution, where buyers are poorly covered by standard B2B databases. But that segment is also the hardest to self-serve. Professional services is the reverse: Clay is easy to adopt there, but the data advantage is less unique and the segment is already contested, including by Clay's own agency ecosystem."),
    ("Method", "(1) Competitor register with an overlap check against 11 Clay use cases. (2) Comparison matrix on the 9 requested dimensions. (3) Illustrative list-price cost comparison by team size. (4) Persona data-coverage test (H3) using BLS education × Pew LinkedIn usage. (5) Segment competitive scoring feeding Phase 3. (6) Positioning map."),
    ("Evidence classes", "A = verified public fact · B = external benchmark / third-party estimate · C = analyst estimate or judgment · D = scenario assumption. Rating confidence: S = sourced, K = analyst knowledge (to verify in primary research)."),
    ("Color legend", "Blue = hardcoded input · Black = formula · Yellow fill = judgment-based score (editable)."),
    ("Not used", "No interviews, win/loss data or Clay internal data. Vendor marketing claims are labeled as claims. ZoomInfo and HubSpot list prices are not public and are not estimated."),
    ("Tabs", "Summary → Competitor_Register → Overlap_Check → Comparison_Matrix → Pricing_Compare → Persona_Coverage → Segment_Scoring → Positioning_Map → Sources"),
]
for i, (k, v) in enumerate(rows, start=4):
    put(ws, f"A{i}", k, bold=True, fill=SFILL)
    put(ws, f"B{i}", v)
    ws.row_dimensions[i].height = 48
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 125

# ------------------------------------------------------------------ Competitor_Register
cr = wb.create_sheet("Competitor_Register")
title(cr, "Competitor register — direct, partial and alternative solutions",
      "Overlap is verified on the Overlap_Check tab, not assumed. Scale facts are company-reported [A] unless marked as an estimate [B].")
hdr(cr, 4, ["ID", "Competitor / alternative", "Type", "What it is", "Where it plays (segments)", "Pricing architecture",
            "Scale facts (class)", "Relevance to this decision", "Sources"],
    widths=[6, 24, 15, 44, 30, 34, 44, 52, 12])
comps = [
    ("C01", "Apollo.io", "Direct",
     "All-in-one prospecting: contact database, sequencing, dialer. Added waterfall enrichment across 18 providers (Oct 2025), acquired Pocus (Mar 2026), launched 'Apollo Anywhere' in AI assistants (Mar 2026) and a spreadsheet product in beta (Oct 2026).",
     "Horizontal; SMB and mid-market, expanding upmarket.",
     "Per seat: Basic $49, Professional $79, Organization $119 per user/month (annual billing); credit limits on mobile numbers and exports.",
     "ARR ~$265M (Sep 2026) and ~$194M (end-2025) [B, Sacra estimate]; $1.6B valuation (2023).",
     "The closest substitute. It is now copying Clay's waterfall and spreadsheet model while staying turnkey. Regency Supply reported Apollo was ineffective for finding local electrical contractors.",
     "S20, S23, S07"),
    ("C02", "ZoomInfo (NASDAQ: GTM)", "Direct (data) / Partial (workflow)",
     "Proprietary B2B contact and company database with intent, website-visitor tracking and engagement tools; data from web scanning, a contributory network and human research.",
     "Horizontal; 74% of ACV 'upmarket'; detected web-tracking users skew to software (19%), IT services (10%), staffing (4%).",
     "Annual subscription contracts; list pricing not public.",
     "FY2025 revenue $1,249.5M (+3%); NRR 90%; 1,921 customers ≥$100K ACV; serves 35,000+ companies [A].",
     "Incumbent data vendor. Slow growth and 90% NRR signal saturation in its core. Regency reported ZoomInfo was ineffective for local contractors. Can be used inside Clay with the customer's own API key (complement).",
     "S19, S33, S07"),
    ("C03", "HubSpot (Sales Hub + Breeze Intelligence)", "Partial (substitute and platform)",
     "CRM suite with native enrichment, buyer intent and form shortening paid in HubSpot Credits (500 / 3,000 / 5,000 per month for Starter / Pro / Enterprise).",
     "SMB and mid-market across industries; the CRM used in the Regency case.",
     "Per-seat CRM subscription plus a credit allowance; extra credits ~$45/mo per 5,000 [B].",
     "306,446 customers; avg subscription revenue per customer $11,800 (Q2 2026) [A].",
     "Both the system of record Clay plugs into (native integration) and a 'good-enough' built-in enrichment alternative for teams with low builder capacity.",
     "S21, S22, S07"),
    ("C04", "Salesforce / Microsoft Dynamics 365", "Complement (partial substitute)",
     "Enterprise CRMs; Clay lists native integrations with Salesforce and Dynamics 365.",
     "Mid-market and enterprise across industries.",
     "Per seat.",
     "Not quantified here.",
     "Systems of record; their native AI agents could absorb simple enrichment/research jobs.",
     "S36"),
    ("C05", "Dun & Bradstreet (Hoovers)", "Partial (data)",
     "Firmographic and business-credit data; directories by industry including manufacturing.",
     "Enterprise data buyers; strong firmographic coverage of manufacturers.",
     "Enterprise data licenses.",
     "Taken private by Clearlake for $7.7B (2025) [A].",
     "Firmographic depth on industrial firms; weaker on workflow. More likely a data source than a workflow competitor.",
     "S41"),
    ("C06", "Vertical data specialists (Dodge Construction Network, Shovels, Thomasnet)", "Complement / partial substitute",
     "Industry-native data: construction projects and specifications (Dodge), building permits and contractor profiles (Shovels), industrial supplier discovery (Thomasnet).",
     "Construction-facing distributors and building-product manufacturers; industrial sourcing.",
     "Annual subscriptions (Dodge), self-serve API from $599/mo (Shovels), paid supplier placements ~$7K–$10K+/yr (Thomasnet).",
     "Dodge: 677K+ firms, 700K+ projects/yr [A, claim]. Shovels: 178M+ permits, ~3.65M contractor profiles, ~$3M ARR [B]. Thomasnet: 500K+ suppliers; no API [B].",
     "The data industrial sellers actually need. Shovels already distributes through Clay; Dodge and Thomasnet do not appear to. Clay's opportunity is to aggregate them, and the risk is that they bundle their own workflow.",
     "S38, S26, S27"),
    ("C07", "Distribution sales AI (e.g., Proton.ai)", "Partial (budget competitor)",
     "AI CRM and recommendation engine trained on distributors' ERP order history; launched a distributor 'industry cloud' (CRM, PIM, e-commerce recommendations) in Mar 2026.",
     "Mid-to-large distributors (e.g., MSC Industrial Supply).",
     "Enterprise subscription (not public).",
     "~$30M raised; 51–200 employees [B].",
     "Competes for the same distributor sales-productivity budget, focused on installed-base cross-sell rather than net-new prospecting, which is Clay's strength.",
     "S24"),
    ("C08", "ERP-embedded CRM (Epicor Prophet 21/Eclipse, Infor, etc.)", "Alternative (status quo)",
     "CRM modules inside distribution ERPs, where order, pricing and customer data live.",
     "Distribution; Epicor was used by 59% of reporting Top-200 electrical distributors (2014, dated).",
     "Bundled with ERP.",
     "9% of surveyed distributors use an ERP-provided CRM (2022) [B].",
     "Clay has no native ERP connector (CRMs and warehouses only), which is a gap for distributors whose customer data lives in the ERP.",
     "S40, S17, S36"),
    ("C09", "Professional-services vertical CRMs (Deltek Vantagepoint, Unanet CRM)", "Alternative / complement",
     "Project-based CRM/ERP for architecture & engineering and government contractors; pursuit and proposal management.",
     "A&E and GovCon professional services.",
     "Per seat / suite.",
     "Unanet: 3,400+ customers (2022) [B]. A&E median proposal win rate 50% (Deltek Clarity 2025, n≈700) [B].",
     "Systems of record for RFP-driven business development; Clay would need to integrate through API.",
     "S42, S39"),
    ("C10", "AI-native GTM platforms (Unify, Cargo, Common Room, Default, Warmly)", "Direct (software segment)",
     "Agentic outbound, signal and orchestration tools.",
     "Mostly software and AI companies.",
     "Usage- or seat-based.",
     "Not quantified here.",
     "Crowd Clay's core software segment; little visible industrial focus, so they are not a major factor in the industrial decision.",
     "S04"),
    ("C11", "Internal build (Zapier, n8n, Make + LLM APIs)", "Alternative",
     "Self-assembled enrichment and research pipelines.",
     "Firms with technical GTM staff.",
     "Tool subscriptions plus engineer time.",
     "n8n named in 41% and Zapier in 40% of sampled GTM-engineer postings (n=209) [B].",
     "A real substitute where builders exist (professional services, software); not viable where they don't (industrials).",
     "S30"),
    ("C12", "GTM agencies and outsourced lead generation (incl. 'Clay agencies')", "Alternative / channel",
     "Agencies that run Clay or other tools on a client's behalf.",
     "All segments; some agencies market 'Clay for manufacturing' programs.",
     "Retainers / project fees.",
     "174 companies listed as Clay solution partners, mostly marketing services and consulting firms [B]; 181 listings in Clay's expert directory [A].",
     "The likely delivery channel for low-builder-capacity segments. In professional services, these agencies are themselves the customers.",
     "S34, S35, S12"),
    ("C13", "Status quo: spreadsheets, manual research, field reps", "Alternative (do nothing)",
     "Reps build lists by hand; relationship-led selling.",
     "Strongest in distribution and manufacturing.",
     "Labor cost.",
     "Distributors: 68% use spreadsheets, 56% use CRM, only 34% of CRM users log ≥75% of interactions (2026) [B]; 15% have no CRM (2022) [B].",
     "The main competitor in industrials is inertia, not another vendor.",
     "S25, S17"),
]
for i, c in enumerate(comps, start=5):
    for j, v in enumerate(c, start=1):
        put(cr, f"{L(j)}{i}", v, bold=(j == 2))
    cr.row_dimensions[i].height = 92
cr.freeze_panes = "C5"

# ------------------------------------------------------------------ Overlap_Check
oc = wb.create_sheet("Overlap_Check")
title(oc, "Overlap check — which alternatives genuinely cover Clay's use cases?",
      "F = full coverage (1.0), P = partial (0.5), N = none (0), ? = not verified (excluded from the index). Confidence: S = sourced, K = analyst knowledge, verify in primary research.")
uses = ["U1 TAM / account list sourcing", "U2 Contact enrichment (email, phone)", "U3 Multi-provider waterfall enrichment",
        "U4 Off-LinkedIn / local business sourcing (maps, permits, licenses)", "U5 AI web research → custom data points",
        "U6 Signals (job change, intent, news)", "U7 CRM sync and data hygiene", "U8 Inbound lead qualification and routing",
        "U9 Outbound sequencing", "U10 Installed-base cross-sell from ERP / transaction data", "U11 Project, bid and permit intelligence"]
players = ["Clay", "Apollo", "ZoomInfo", "HubSpot Breeze", "Vertical data specialists", "Proton.ai", "Internal build"]
hdr(oc, 4, ["Use case"] + players + ["Evidence / notes"], widths=[46] + [13] * 7 + [90])
grid = [
    # Clay, Apollo, ZoomInfo, HubSpot, Vertical, Proton, Internal, note
    ("F", "F", "F", "P", "P", "N", "P", "Clay Audiences/TAM sourcing (S01); Apollo/ZoomInfo databases (S19, S20); HubSpot enrichment fills existing records rather than building lists (S21) [K]; Dodge/Thomasnet list firms in their niche (S38, S27)."),
    ("F", "F", "F", "P", "N", "N", "P", "Core function of database vendors; Breeze enriches contact/company fields (S21)."),
    ("F", "F", "N", "N", "N", "N", "P", "Clay waterfall across 150–200+ providers (S01); Apollo added an 18-provider waterfall in Oct 2025 (S23). ZoomInfo and HubSpot use their own data [K]."),
    ("F", "N", "N", "N", "F", "N", "P", "Clay Google Maps integration (S37) and Shovels permit/contractor data distributed via Clay (S26). Regency: ZoomInfo and Apollo could not find local contractors (S07). Dodge/Shovels are native here."),
    ("F", "?", "N", "P", "N", "N", "F", "Claygent / account agents (S01). Apollo's AI research depth not verified. Breeze offers AI features inside HubSpot [K]. Internal LLM pipelines can replicate (S30)."),
    ("F", "P", "F", "P", "P", "N", "P", "Clay signals: job change, news, web intent (S01); ZoomInfo intent + visitor tracking (S19, S33); Breeze buyer intent (S21); Dodge spec/bid alerts (S38)."),
    ("F", "F", "F", "F", "P", "P", "P", "Clay native Salesforce/HubSpot/Dynamics sync at Growth+ (S01, S36). HubSpot is the CRM itself. Dodge leads importable to CRM (S38)."),
    ("F", "P", "P", "F", "N", "N", "P", "Regency automated e-commerce registration qualification with Clay (S07); HubSpot form shortening and routing (S21) [K]."),
    ("F", "F", "P", "F", "N", "N", "P", "Clay Sequencer (S01); Apollo sequencing (S20); HubSpot sequences [K]."),
    ("P", "N", "N", "P", "N", "F", "P", "Clay has no ERP connector, only CRMs and warehouses (S36), so partial via warehouse sync/HTTP API. Proton.ai is trained on ERP order history (S24)."),
    ("P", "N", "N", "N", "F", "N", "N", "Clay added public-sector bid signals (GovSpend, Pursuit) in Sep 2026 (S09) and carries Shovels permits (S26); Dodge is the construction-project specialist (S38)."),
]
for i, g in enumerate(grid):
    r = 5 + i
    put(oc, f"A{r}", uses[i])
    for j, v in enumerate(g[:7]):
        put(oc, f"{L(2+j)}{r}", v, font=BLUE, center=True, fill=YFILL)
    put(oc, f"I{r}", g[7])
    oc.row_dimensions[r].height = 44
last = 5 + len(grid) - 1
r = last + 2
labels = [("Use cases rated (excl. '?')", 'COUNTIF({c}5:{c}%d,"F")+COUNTIF({c}5:{c}%d,"P")+COUNTIF({c}5:{c}%d,"N")' % (last, last, last)),
          ("Full coverage (F)", 'COUNTIF({c}5:{c}%d,"F")' % last),
          ("Partial coverage (P)", 'COUNTIF({c}5:{c}%d,"P")' % last),
          ("Coverage index (F + 0.5·P) ÷ rated", None),
          ("Overlap with Clay's full use-case set: (F + 0.5·P) ÷ 11", None)]
for k, (lab, f) in enumerate(labels):
    rr = r + k
    put(oc, f"A{rr}", lab, bold=True, fill=SFILL)
    for j in range(7):
        c = L(2 + j)
        if f:
            put(oc, f"{c}{rr}", "=" + f.format(c=c), center=True)
        elif k == 3:
            put(oc, f"{c}{rr}", f"=IFERROR(({c}{r+1}+0.5*{c}{r+2})/{c}{r},0)", fmt=PCT, center=True)
        else:
            put(oc, f"{c}{rr}", f"=({c}{r+1}+0.5*{c}{r+2})/11", fmt=PCT, center=True)
OC_IDX = r + 4
put(oc, f"A{r+6}", "Reading: Apollo is the only alternative that covers most of Clay's use cases. Its gaps are exactly the industrial-specific ones (U4 off-LinkedIn/local sourcing, U11 project and permit intelligence). Vertical specialists cover those gaps but little else. Clay's defensible overlap position is aggregating the specialists into a workflow that standard vendors don't cover.", border=False)
oc.merge_cells(f"A{r+6}:I{r+6}")
oc.row_dimensions[r + 6].height = 44
oc.freeze_panes = "B5"

# ------------------------------------------------------------------ Comparison_Matrix
cm = wb.create_sheet("Comparison_Matrix")
title(cm, "Comparison matrix — the nine requested dimensions",
      "Qualitative evidence with a 1–5 score (5 = strongest for an industrial mid-market buyer). Scores are analyst judgments [C]; sources on Competitor_Register and Sources.")
cols = ["Clay", "Apollo", "ZoomInfo", "HubSpot (Breeze)", "Vertical specialists", "Internal build"]
hdr(cm, 4, ["Dimension"] + sum([[c, "Score"] for c in cols], []),
    widths=[24] + [34, 7] * 6)
dims = [
    ("Product capabilities",
     [("Data marketplace, waterfall, AI agents, signals, workflows, sequencer, ads audiences.", 5),
      ("Database, sequencing, dialer; now waterfall and a spreadsheet beta.", 4),
      ("Deep proprietary data, intent, engagement tools.", 4),
      ("CRM-native enrichment, intent, form shortening.", 3),
      ("Deep but narrow (projects, permits, suppliers).", 2),
      ("Whatever is built.", 3)]),
    ("Customer segments",
     [("Software/AI-heavy base (~40% software development among detected web-intent users); few industrial references.", 2),
      ("Horizontal SMB and mid-market.", 4),
      ("Upmarket (74% of ACV); diversified base.", 3),
      ("Broad SMB and mid-market (306K customers).", 5),
      ("Construction and industrial niches.", 4),
      ("Firms with technical GTM staff.", 2)]),
    ("Pricing architecture",
     [("Platform fee plus usage (Actions + Data Credits); unlimited seats; Growth ~$5.4K/yr at list.", 4),
      ("Per seat $49–$119/user/mo plus credits.", 3),
      ("Opaque annual contracts.", 2),
      ("Per seat plus credit allowance.", 3),
      ("Subscriptions; Shovels from $599/mo.", 3),
      ("Low cash cost, high labor cost.", 2)]),
    ("Data coverage of industrial buyers",
     [("Aggregates 150–200+ providers incl. Google Maps and Shovels; strongest for off-LinkedIn personas.", 5),
      ("Large database but reported weak for local contractors.", 2),
      ("Large database but reported weak for local contractors.", 2),
      ("Own database; generic coverage.", 2),
      ("Best within their niche.", 4),
      ("Depends on sources licensed.", 3)]),
    ("Workflow automation",
     [("Core strength: tables, workflows, functions, webhooks, CLI/API.", 5),
      ("Sequences and workflow destinations; less flexible.", 3),
      ("Workflow add-ons; less flexible.", 3),
      ("HubSpot workflows inside the CRM.", 4),
      ("Minimal.", 1),
      ("Unlimited but self-maintained.", 4)]),
    ("CRM integration",
     [("Native Salesforce, HubSpot, Dynamics 365, Pipedrive at Growth+; no ERP connector.", 4),
      ("Standard CRM syncs [K].", 4),
      ("Standard CRM syncs [K].", 4),
      ("Is the CRM.", 5),
      ("Dodge imports into CRM; Thomasnet has no API.", 2),
      ("Custom.", 3)]),
    ("Implementation complexity (5 = easiest)",
     [("High: value depends on a builder; Regency's workflows were custom-built.", 2),
      ("Low: turnkey per-seat tool.", 5),
      ("Medium: admin plus training.", 3),
      ("Lowest: already inside the CRM.", 5),
      ("Low: search portals.", 4),
      ("Highest: engineering required.", 1)]),
    ("Differentiation for industrial sellers",
     [("Only option combining off-LinkedIn data, AI research and CRM sync in one workflow.", 5),
      ("Price and simplicity.", 3),
      ("Data breadth for white-collar personas.", 3),
      ("Convenience.", 2),
      ("Industry-native data.", 4),
      ("Full control.", 2)]),
    ("Switching costs once adopted",
     [("High: workflows, functions and history become embedded (unlimited seats spread usage).", 4),
      ("Medium: sequences and seats.", 3),
      ("Medium-high: annual contracts.", 3),
      ("Very high: system of record.", 5),
      ("Low-medium.", 2),
      ("High: sunk build.", 4)]),
]
for i, (dim, vals) in enumerate(dims):
    r = 5 + i
    put(cm, f"A{r}", dim, bold=True)
    for j, (txt, sc) in enumerate(vals):
        put(cm, f"{L(2+2*j)}{r}", txt)
        put(cm, f"{L(3+2*j)}{r}", sc, font=BLUE, fill=YFILL, center=True)
    cm.row_dimensions[r].height = 58
lr = 5 + len(dims) - 1
put(cm, f"A{lr+1}", "Average score (unweighted)", bold=True, fill=SFILL)
for j in range(6):
    c = L(3 + 2 * j)
    put(cm, f"{L(2+2*j)}{lr+1}", None, fill=SFILL)
    put(cm, f"{c}{lr+1}", f"=AVERAGE({c}5:{c}{lr})", fmt=D2, bold=True, fill=SFILL, center=True)
put(cm, f"A{lr+2}", "Implementation complexity score", bold=True, fill=SFILL)
for j in range(6):
    c = L(3 + 2 * j)
    put(cm, f"{L(2+2*j)}{lr+2}", None, fill=SFILL)
    put(cm, f"{c}{lr+2}", f"={c}{5+6}", fmt=D1, center=True, fill=SFILL)
put(cm, f"A{lr+4}", "Reading: Clay leads on capability, industrial data coverage and differentiation, but scores lowest among commercial vendors on implementation ease. The competitive question in industrials is less 'can Clay out-feature Apollo' than 'can Clay be made as easy to adopt as Apollo or HubSpot'.", border=False)
cm.merge_cells(f"A{lr+4}:M{lr+4}")
cm.row_dimensions[lr + 4].height = 32
cm.freeze_panes = "B5"

# ------------------------------------------------------------------ Pricing_Compare
pc = wb.create_sheet("Pricing_Compare")
title(pc, "Illustrative annual list cost by team size — seat-based vs unlimited-seat pricing",
      "List prices, annual billing, before discounts and add-ons. ZoomInfo and HubSpot seat prices are not public / not captured and are excluded rather than estimated. [A] inputs → [C] outputs.")
put(pc, "A4", "Inputs (monthly, annual billing)", bold=True, fill=SFILL)
inp = [
    ("Clay Growth — Actions 480K/yr tier", 185, "S01"), ("Clay Growth — Data Credits 72K/yr tier", 261, "S01"),
    ("Clay Growth — Actions 1.2M/yr tier", 405, "S01"), ("Clay Growth — Data Credits 120K/yr tier", 413, "S01"),
    ("Clay Growth — Actions 2.4M/yr tier", 765, "S01"), ("Clay Growth — Data Credits 240K/yr tier", 792, "S01"),
    ("Apollo Basic — per user", 49, "S20"), ("Apollo Professional — per user", 79, "S20"), ("Apollo Organization — per user", 119, "S20"),
]
hdr(pc, 5, ["Item", "$ / month", "Source"], widths=[46, 14, 10])
IN = {}
for i, (n, v, s) in enumerate(inp):
    r = 6 + i
    put(pc, f"A{r}", n)
    put(pc, f"B{r}", v, font=BLUE, fmt=USD)
    put(pc, f"C{r}", s)
    IN[n] = r
r0 = 6 + len(inp) + 1
put(pc, f"A{r0}", "Annual cost by number of users who need data ($/yr)", bold=True, fill=SFILL)
seats = [3, 9, 28, 50]
hdr(pc, r0 + 1, ["Option"] + [f"{s} users" for s in seats] + ["Note"], widths=[46, 14, 14, 14, 14, 60])
for j, s in enumerate(seats):
    pc.cell(row=r0 + 2, column=2 + j)
opts = [
    ("Clay Growth — entry configuration", lambda s: f"=($B${IN['Clay Growth — Actions 480K/yr tier']}+$B${IN['Clay Growth — Data Credits 72K/yr tier']})*12",
     "Unlimited seats. Entry tier = the Phase 1 SMB ACV floor ($5,352)."),
    ("Clay Growth — mid configuration", lambda s: f"=($B${IN['Clay Growth — Actions 1.2M/yr tier']}+$B${IN['Clay Growth — Data Credits 120K/yr tier']})*12",
     "Same seat count, higher usage."),
    ("Clay Growth — high configuration", lambda s: f"=($B${IN['Clay Growth — Actions 2.4M/yr tier']}+$B${IN['Clay Growth — Data Credits 240K/yr tier']})*12",
     "Upper bound of self-serve Growth before Enterprise."),
    ("Apollo Basic", lambda s: f"=$B${IN['Apollo Basic — per user']}*12*{s}", "Scales linearly with seats."),
    ("Apollo Professional", lambda s: f"=$B${IN['Apollo Professional — per user']}*12*{s}", "Most common tier for sales teams [K]."),
    ("Apollo Organization", lambda s: f"=$B${IN['Apollo Organization — per user']}*12*{s}", ""),
]
for i, (n, fn, note) in enumerate(opts):
    r = r0 + 2 + i
    put(pc, f"A{r}", n)
    for j, s in enumerate(seats):
        put(pc, f"{L(2+j)}{r}", fn(s), fmt=USD)
    put(pc, f"F{r}", note)
ratio_r = r0 + 2 + len(opts) + 1
put(pc, f"A{ratio_r}", "Apollo Professional ÷ Clay Growth mid configuration", bold=True, fill=SFILL)
for j in range(len(seats)):
    c = L(2 + j)
    put(pc, f"{c}{ratio_r}", f"={c}{r0+2+4}/{c}{r0+2+1}", fmt='0.00"x"', fill=SFILL)
put(pc, f"A{ratio_r+2}", "Why 28 users: Phase 1 estimates ~28 sales staff per in-scope distributor vs ~9 per professional-services firm (02_Market_Sizing, Segment_Profile). At distributor team sizes, seat-based tools cost several times Clay's mid configuration at list. That is a structural price advantage, but it only materializes if the distributor can operate Clay's workflows.", border=False)
pc.merge_cells(f"A{ratio_r+2}:F{ratio_r+2}")
pc.row_dimensions[ratio_r + 2].height = 44
PC_RATIO = ratio_r

# ------------------------------------------------------------------ Persona_Coverage
pv = wb.create_sheet("Persona_Coverage")
title(pv, "H3 test — are each segment's buyer personas likely to be covered by LinkedIn-derived B2B databases?",
      "Estimated LinkedIn use = Σ (share of workers by education) × (LinkedIn use rate for that education group). Education: BLS (S29). LinkedIn use: Pew, fielded May–Sep 2023 (S28). Analyst estimate [C].")
put(pv, "A4", "Pew LinkedIn use rate by education (U.S. adults)", bold=True, fill=SFILL)
pew = [("High school or less", 0.10), ("Some college (incl. associate's)", 0.28), ("Bachelor's degree or higher", 0.53)]
for i, (n, v) in enumerate(pew):
    put(pv, f"A{5+i}", n)
    put(pv, f"B{5+i}", v, font=BLUE, fmt=PCT)
PEW = {0: "$B$5", 1: "$B$6", 2: "$B$7"}
hdr(pv, 9, ["Buyer persona (BLS occupation)", "SOC", "Persona of which segment's customers", "< HS", "HS", "Some college", "Associate", "Bachelor's", "Master's", "Doctoral/prof.",
            "Bachelor's+ share", "Est. LinkedIn use", "Index vs all workers"],
    widths=[44, 10, 30, 8, 8, 10, 10, 10, 9, 11, 12, 12, 12])
edu = [
    ("Total, all occupations", "00-0000", "Reference", [7.4, 22.4, 18.2, 9.4, 25.5, 12.1, 5.0]),
    ("Electricians", "47-2111", "A. Distribution (contractor-facing)", [8.4, 37.8, 29.1, 15.4, 7.5, 1.4, 0.3]),
    ("Plumbers, pipefitters & steamfitters", "47-2152", "A. Distribution (contractor-facing)", [14.4, 46.6, 23.4, 8.9, 5.4, 1.0, 0.3]),
    ("HVAC & refrigeration mechanics/installers", "49-9021", "A. Distribution (contractor-facing)", [10.0, 41.6, 27.7, 13.5, 5.8, 1.2, 0.2]),
    ("First-line supervisors of construction trades", "47-1011", "A. Distribution (contractor-facing)", [17.1, 43.1, 22.1, 7.6, 8.3, 1.5, 0.3]),
    ("Maintenance & repair workers, general", "49-9071", "A. Distribution (MRO / plant-facing)", [12.0, 43.3, 24.9, 11.4, 6.8, 1.2, 0.3]),
    ("Construction managers", "11-9021", "A. Distribution (contractor-facing)", [8.1, 26.7, 21.6, 8.9, 27.2, 6.6, 0.9]),
    ("Purchasing managers", "11-3061", "A. Distribution and B. Manufacturing", [1.1, 11.2, 14.8, 8.0, 39.2, 22.3, 3.3]),
    ("Industrial production managers", "11-3051", "A. Distribution and B. Manufacturing", [3.4, 19.5, 18.9, 8.4, 34.4, 13.5, 1.8]),
    ("Civil engineers", "17-2051", "B. Manufacturing (OEM / project buyers)", [0.8, 3.3, 4.1, 3.5, 58.1, 26.2, 4.0]),
    ("General & operations managers", "11-1021", "B. Manufacturing and C. Prof. services", [3.3, 17.0, 23.2, 9.6, 33.6, 11.4, 2.0]),
    ("Chief executives", "11-1011", "C. Prof. services", [1.5, 7.7, 12.5, 5.0, 40.9, 25.3, 7.0]),
    ("Marketing managers", "11-2021", "C. Prof. services", [0.7, 3.9, 8.1, 3.9, 58.3, 22.9, 2.1]),
    ("Facilities managers", "11-3013", "C. Prof. services (A&E clients)", [4.2, 20.9, 24.5, 11.3, 29.7, 8.1, 1.2]),
    ("Software developers", "15-1252", "C. Prof. services (IT services clients)", [0.6, 2.2, 7.3, 3.9, 51.7, 30.2, 4.0]),
]
for i, (n, soc, seg, vals) in enumerate(edu):
    r = 10 + i
    put(pv, f"A{r}", n, bold=(i == 0))
    put(pv, f"B{r}", soc)
    put(pv, f"C{r}", seg)
    for j, v in enumerate(vals):
        put(pv, f"{L(4+j)}{r}", v / 100, font=BLUE, fmt=PCT)
    put(pv, f"K{r}", f"=SUM(H{r}:J{r})", fmt=PCT)
    put(pv, f"L{r}", f"=(D{r}+E{r})*{PEW[0]}+(F{r}+G{r})*{PEW[1]}+(H{r}+I{r}+J{r})*{PEW[2]}", fmt=PCT)
    put(pv, f"M{r}", f"=L{r}/$L$10", fmt='0.00"x"')
pl = 10 + len(edu) - 1
pv.conditional_formatting.add(f"L11:L{pl}", ColorScaleRule(start_type="min", start_color="F8CBAD", end_type="max", end_color="C6E0B4"))
sr = pl + 2
put(pv, f"A{sr}", "Segment summary (simple average of mapped personas)", bold=True, fill=SFILL)
hdr(pv, sr + 1, ["Segment", "Personas averaged", "", "Est. LinkedIn use", "Index vs all workers", "Index vs Prof. services"])
segsum = [("A. Distribution — contractor-facing personas", "Electricians, plumbers, HVAC, construction supervisors", [11, 12, 13, 14]),
          ("A. Distribution — all mapped personas", "Contractor + MRO + construction/purchasing/production managers", [11, 12, 13, 14, 15, 16, 17, 18]),
          ("B. Manufacturing — customer personas", "Purchasing, production managers, civil engineers, G&O managers", [17, 18, 19, 20]),
          ("C. Prof. services — customer personas", "CEOs, G&O, marketing, facilities managers, software developers", [20, 21, 22, 23, 24])]
for k, (s, desc, rows_) in enumerate(segsum):
    r = sr + 2 + k
    put(pv, f"A{r}", s, bold=True)
    put(pv, f"B{r}", desc)
    pv.merge_cells(f"B{r}:C{r}")
    put(pv, f"D{r}", "=AVERAGE(" + ",".join(f"L{x}" for x in rows_) + ")", fmt=PCT, bold=True)
    put(pv, f"E{r}", f"=D{r}/$L$10", fmt='0.00"x"')
    put(pv, f"F{r}", f"=D{r}/$D${sr+5}", fmt='0.00"x"')
PV_SUM = sr + 2
put(pv, f"A{sr+7}", "Caveats: (1) assumes LinkedIn use depends only on education. Occupation likely lowers trades' use further, so the gap shown is conservative. (2) B2B databases also use web crawling and contributory networks (ZoomInfo, S19), so LinkedIn presence is a proxy for coverage, not a measurement. (3) Direct corroboration: Regency Supply found ZoomInfo and Apollo ineffective for local electrical contractors and used Google Maps listings instead (S07).", border=False)
pv.merge_cells(f"A{sr+7}:M{sr+7}")
pv.row_dimensions[sr + 7].height = 44
pv.freeze_panes = "D10"

# ------------------------------------------------------------------ Segment_Scoring
ss = wb.create_sheet("Segment_Scoring")
title(ss, "Segment competitive scoring — inputs to Phase 3 (1 = weak for Clay, 5 = strong)",
      "Factor scores are analyst judgments [C] tied to evidence; weights are assumptions [D]. Competitive position and adoption fit are kept separate so Phase 3 does not double-count ease of execution.")
hdr(ss, 4, ["Factor", "Feeds Phase 3 criterion", "Weight within block", "A. Industrial Distribution", "B. Industrial Manufacturing", "C. B2B Prof. Services", "Evidence"],
    widths=[42, 22, 12, 14, 14, 14, 100])
factors = [
    ("BLOCK 1 — Competitive position", None, None, None, None, None, None),
    ("F1 Relevance of Clay's data differentiation (off-LinkedIn, multi-source, AI research)", "Competitive Position", 1 / 3, 5, 3, 2,
     "Persona_Coverage: contractor personas ≈0.45–0.5× professional-services personas in estimated LinkedIn use; Regency case (S07). Manufacturing's customers (purchasing, engineers) are near-average. Prof.-services buyers are well covered by standard databases, so Clay's edge there is workflow flexibility, not data."),
    ("F2 Incumbent saturation (5 = little entrenched sales-intelligence use)", "Competitive Position", 1 / 3, 4, 4, 2,
     "Distributors: 15% no CRM, 68% spreadsheets, only 34% of CRM users log most interactions (S17, S25). ZoomInfo's detected users skew to software, IT services, staffing, advertising (S33). IT services and agencies are heavy users of sales-intelligence tools [K]."),
    ("F3 Threat from vertical specialists / systems of record (5 = low)", "Competitive Position", 1 / 3, 3, 3, 3,
     "Distribution: Proton.ai and ERP CRMs (S24, S17), but focused on installed base. Manufacturing: Thomasnet/D&B are data, not workflow (S27, S41). Prof. services: Deltek/Unanet own pursuit records in A&E (S39, S42). All largely complementary."),
    ("Competitive position score", None, None, "calc", "calc", "calc", "Weighted average of F1–F3."),
    ("BLOCK 2 — Adoption / execution fit", None, None, None, None, None, None),
    ("F4 Builder capacity to self-serve Clay", "Ease of Execution", 0.5, 2, 1, 5,
     "Computer-occupation index vs all industries: 0.69 / 0.45 / 7.15; marketing-ops index 1.46 / 0.62 / 5.48 (02_Market_Sizing). GTM-engineer postings concentrate in software (43.9%, S31)."),
    ("F5 Existing references and partner channel", "Ease of Execution", 0.5, 2, 2, 4,
     "Distribution: one public case study (Regency, S07). Manufacturing: LG pilot via Marketbridge (S08) and agency marketing (S12). Prof. services: a case-study filter of its own (S06); Clay's solution partners are mostly marketing and consulting firms (S34). No industrial-specialist partners found on the visible directory pages (S35)."),
    ("Adoption / execution fit score", None, None, "calc", "calc", "calc", "Weighted average of F4–F5."),
    ("BLOCK 3 — Flag for Phase 3 market attractiveness", None, None, None, None, None, None),
    ("F6 Incremental headroom (5 = largely unpenetrated by Clay)", "Market Attractiveness (adjustment)", None, 5, 5, 2,
     "Clay's detected web-intent users are ~40% software development, with no industrial industry in the top 5 (S32; small, feature-specific sample). Agencies and consultancies are already Clay customers and partners (S34), so the incremental professional-services pool is smaller than Phase 1's SOM implies."),
]
calc_rows = {}
block_start = None
for i, f in enumerate(factors):
    r = 5 + i
    name, crit, w, a, b, c, ev = f
    if crit is None and a is None:
        put(ss, f"A{r}", name, bold=True, fill=SFILL)
        for col in "BCDEFG":
            put(ss, f"{col}{r}", None, fill=SFILL)
        block_start = r + 1
        continue
    if a == "calc":
        put(ss, f"A{r}", name, bold=True)
        put(ss, f"B{r}", None)
        put(ss, f"C{r}", f"=SUM(C{block_start}:C{r-1})", fmt='0.00')
        for col in "DEF":
            put(ss, f"{col}{r}", f"=SUMPRODUCT($C{block_start}:$C{r-1},{col}{block_start}:{col}{r-1})/SUM($C{block_start}:$C{r-1})", fmt=D2, bold=True, center=True)
        put(ss, f"G{r}", ev)
        calc_rows[name] = r
        continue
    put(ss, f"A{r}", name)
    put(ss, f"B{r}", crit)
    put(ss, f"C{r}", w, font=BLUE, fmt='0.00') if w is not None else put(ss, f"C{r}", "n/a")
    for col, v in zip("DEF", (a, b, c)):
        put(ss, f"{col}{r}", v, font=BLUE, fill=YFILL, center=True)
    put(ss, f"G{r}", ev)
    ss.row_dimensions[r].height = 64
SS_CP = calc_rows["Competitive position score"]
SS_EX = calc_rows["Adoption / execution fit score"]
SS_HR = 5 + len(factors) - 1
nr = SS_HR + 2
put(ss, f"A{nr}", "Reading: Clay's competitive position is strongest where its execution fit is weakest. Distribution scores highest on position because the data gap is real and incumbents are thin, but lowest-but-one on the ability to adopt. Professional services is the mirror image. This is the core trade-off Phase 3 must price explicitly.", border=False)
ss.merge_cells(f"A{nr}:G{nr}")
ss.row_dimensions[nr].height = 44

# ------------------------------------------------------------------ Positioning_Map
pm = wb.create_sheet("Positioning_Map")
title(pm, "Competitive positioning framework — industrial buyer's view",
      "X = coverage of industrial / off-LinkedIn buyer data (1–5). Y = ease of adoption for teams without technical GTM staff (1–5). Coordinates are analyst judgments [C] derived from Comparison_Matrix.")
hdr(pm, 4, ["Player", "X: industrial data coverage", "Y: ease of adoption", "Rationale"], widths=[34, 16, 16, 90])
pts = [
    ("Clay (today)", f"=Comparison_Matrix!C8", f"=Comparison_Matrix!C11", "Strongest industrial data aggregation; adoption requires a builder."),
    ("Apollo", f"=Comparison_Matrix!E8", f"=Comparison_Matrix!E11", "Turnkey; weak on local contractors per Regency."),
    ("ZoomInfo", f"=Comparison_Matrix!G8", f"=Comparison_Matrix!G11", "Large database; enterprise contracting and admin."),
    ("HubSpot (Breeze)", f"=Comparison_Matrix!I8", f"=Comparison_Matrix!I11", "Easiest; generic data."),
    ("Vertical specialists", f"=Comparison_Matrix!K8", f"=Comparison_Matrix!K11", "Deep niche data; portal-style ease."),
    ("Internal build", f"=Comparison_Matrix!M8", f"=Comparison_Matrix!M11", "Any data; needs engineers."),
    ("Target: packaged 'Clay for industrial distribution'", 5, 4, "Hypothesized white space: Clay's data plus pre-built templates and partner delivery (tested in Phases 3–5)."),
]
for i, (n, x, y, why) in enumerate(pts):
    r = 5 + i
    put(pm, f"A{r}", n, bold=True)
    put(pm, f"B{r}", x, font=GREEN if isinstance(x, str) else BLUE, fmt=D1, center=True)
    put(pm, f"C{r}", y, font=GREEN if isinstance(y, str) else BLUE, fmt=D1, center=True)
    put(pm, f"D{r}", why)
ch = ScatterChart()
ch.title = "Industrial data coverage vs ease of adoption"
ch.style = 13
ch.x_axis.title = "Coverage of industrial / off-LinkedIn buyer data →"
ch.y_axis.title = "Ease of adoption for non-technical teams →"
ch.x_axis.scaling.min, ch.x_axis.scaling.max = 0, 5.5
ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 5.5
ch.x_axis.majorUnit = ch.y_axis.majorUnit = 1
ch.x_axis.delete = False
ch.y_axis.delete = False
colors = ["1F2A44", "8E44AD", "C0392B", "E67E22", "16A085", "7F8C8D", "2E86C1"]
for i in range(len(pts)):
    r = 5 + i
    s = Series(Reference(pm, min_col=3, min_row=r), Reference(pm, min_col=2, min_row=r), title_from_data=False, title=pts[i][0])
    s.marker = Marker(symbol="diamond" if i == len(pts) - 1 else "circle", size=11)
    s.marker.graphicalProperties.solidFill = colors[i]
    s.marker.graphicalProperties.line.solidFill = colors[i]
    s.graphicalProperties.line.noFill = True
    ch.series.append(s)
ch.height, ch.width = 11, 20
pm.add_chart(ch, "A14")
put(pm, "A12", "Note: Apollo and HubSpot plot at the same point (2, 5). Both are turnkey but generic for industrial personas.", border=False)
pm.merge_cells("A12:D12")

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary", 1)
title(sm, "Summary — Phase 2 competitive findings", "All values link to the analysis tabs.")
hdr(sm, 4, ["Metric", "A. Industrial Distribution", "B. Industrial Manufacturing", "C. B2B Prof. Services", "Source tab"], widths=[60, 18, 18, 18, 22])
srows = [
    ("Competitive position score (1–5)", [f"=Segment_Scoring!{c}{SS_CP}" for c in "DEF"], D2, "Segment_Scoring"),
    ("Adoption / execution fit score (1–5)", [f"=Segment_Scoring!{c}{SS_EX}" for c in "DEF"], D2, "Segment_Scoring"),
    ("Incremental headroom flag (1–5)", [f"=Segment_Scoring!{c}{SS_HR}" for c in "DEF"], D1, "Segment_Scoring"),
    ("Est. LinkedIn use of customer personas (proxy for database coverage)", [f"=Persona_Coverage!D{PV_SUM+1}", f"=Persona_Coverage!D{PV_SUM+2}", f"=Persona_Coverage!D{PV_SUM+3}"], PCT, "Persona_Coverage"),
    ("   … index vs professional-services personas", [f"=Persona_Coverage!F{PV_SUM+1}", f"=Persona_Coverage!F{PV_SUM+2}", f"=Persona_Coverage!F{PV_SUM+3}"], '0.00"x"', "Persona_Coverage"),
]
for i, (n, fs, fmt, src) in enumerate(srows):
    r = 5 + i
    put(sm, f"A{r}", n)
    for j, f in enumerate(fs):
        put(sm, f"{L(2+j)}{r}", f, font=GREEN, fmt=fmt, center=True)
    put(sm, f"E{r}", src)
r = 5 + len(srows) + 1
put(sm, f"A{r}", "Cross-segment facts", bold=True, fill=SFILL)
r += 1
facts = [
    ("Apollo coverage of Clay's 11 use cases (F + 0.5·P ÷ 11)", f"=Overlap_Check!C{OC_IDX}", PCT, "Overlap_Check"),
    ("Vertical specialists' coverage of Clay's 11 use cases", f"=Overlap_Check!F{OC_IDX}", PCT, "Overlap_Check"),
    ("Apollo Professional ÷ Clay Growth (mid) annual list cost at 28 users", f"=Pricing_Compare!D{PC_RATIO}", '0.00"x"', "Pricing_Compare"),
    ("Apollo Professional ÷ Clay Growth (mid) annual list cost at 9 users", f"=Pricing_Compare!C{PC_RATIO}", '0.00"x"', "Pricing_Compare"),
]
for n, f, fmt, src in facts:
    put(sm, f"A{r}", n)
    put(sm, f"B{r}", f, font=GREEN, fmt=fmt, center=True)
    put(sm, f"E{r}", src)
    r += 1
r += 1
put(sm, f"A{r}", "Key takeaways", bold=True, fill=SFILL)
r += 1
tk = [
    "1. H3 supported for distribution: contractor buyer personas are estimated at roughly half the LinkedIn presence of professional-services buyers, and Regency Supply reports ZoomInfo and Apollo failed to find them. Weak for manufacturing, whose customers are white-collar buyers.",
    "2. H5 supported: Clay's realistic industrial position is complementary. It sits on top of HubSpot, Salesforce or Dynamics and aggregates niche sources (Google Maps, Shovels). It does not replace CRMs, and it lacks an ERP connector.",
    "3. The main competitor in industrials is inertia (spreadsheets, low CRM discipline), followed by turnkey tools (Apollo, HubSpot) that win on ease rather than data.",
    "4. Apollo is converging on Clay's model (waterfall, spreadsheet beta), which narrows Clay's moat in horizontal and software segments more than in industrial niches.",
    "5. Clay's unlimited-seat pricing is structurally cheaper than per-seat tools for seller-dense distributors, but only if the customer can operate it. Packaging and partners, not features, are the binding constraint (H4).",
]
for t in tk:
    put(sm, f"A{r}", t)
    sm.merge_cells(f"A{r}:E{r}")
    sm.row_dimensions[r].height = 34
    r += 1

# ------------------------------------------------------------------ Sources
so = wb.create_sheet("Sources")
title(so, "Sources used in this workbook", "Full register with dates, reliability and limitations: Research_Sources.csv")
hdr(so, 4, ["ID", "Source", "Publisher", "Date / reference", "URL"], widths=[7, 60, 30, 22, 90])
srcs = [
    ("S01", "Clay pricing page", "Clay Labs Inc.", "Live, Oct 2026", "https://www.clay.com/pricing"),
    ("S04", "Clay company profile / ARR estimate", "Sacra", "2026", "https://sacra.com/c/clay/"),
    ("S06", "Clay customer stories (industry filters)", "Clay Labs Inc.", "Live, Oct 2026", "https://www.clay.com/customers"),
    ("S07", "Regency Supply case study", "Clay Labs Inc.", "Undated (~2025)", "https://www.clay.com/customers/regency-supply"),
    ("S08", "Clay–Marketbridge partnership", "B2BMX Insights", "2026-10-06", "https://www.demandgenreport.com/industry-news/news-brief/clay-marketbridge-partner-to-modernize-b2b-gtm-execution/54721/"),
    ("S09", "Clay integrations roundup", "Clay Labs Inc.", "2026-09-29", "https://www.clay.com/changelog/integrations-roundup"),
    ("S12", "Clay agency for manufacturing (city pages)", "DevCommX", "Undated", "https://www.devcommx.com/clay-agency/akron/manufacturing"),
    ("S17", "State of CRM in Distribution", "Distribution Strategy Group", "2022", "https://archive.distributionstrategy.com/wp-content/uploads/2022/10/DSG-The-State-of-CRM-in-Distribution-Making-the-Software-Work-for-You.pdf"),
    ("S19", "ZoomInfo Q4 and FY2025 results", "ZoomInfo via Business Wire", "2026-02-09", "https://www.businesswire.com/news/home/20260209677262/en/"),
    ("S20", "Apollo.io pricing in 2026 (third-party summary)", "Docket", "Verified Apr 2026", "https://docket.io/resources/research/apollo-pricing"),
    ("S21", "HubSpot Breeze Intelligence pricing guide", "eesel AI", "Edited 2026-10-14", "https://www.eesel.ai/blog/hubspot-breeze-intelligence-pricing"),
    ("S22", "HubSpot Q2 2026 results", "HubSpot via Business Wire", "2026-08-05", "https://www.businesswire.com/news/home/20260805185463/en/"),
    ("S23", "Apollo.io at $265M ARR", "Sacra", "2026-10-14", "https://sacra.com/research/apollo-at-265m-arr/"),
    ("S24", "Proton.ai company profile", "Simplify", "2026", "https://simplify.jobs/c/ProtonAi"),
    ("S25", "State of Sales in Distribution 2026 (Phocas Software)", "Distribution Strategy Group", "2026-09-21", "http://distributionstrategy.com/2026/09/nearly-half-of-distributors-use-ai-for-sales-but-many-still-dont-track-key-metrics/"),
    ("S26", "Shovels company profile", "Sacra", "2026", "https://sacra.com/c/shovels/"),
    ("S27", "Thomasnet tool profile", "The GTM Directory", "2026", "https://thegtmdirectory.com/tools/thomasnet/md"),
    ("S28", "Americans' Social Media Use", "Pew Research Center", "2024-01-31 (fielded May–Sep 2023)", "https://www.pewresearch.org/internet/2024/01/31/americans-social-media-use/"),
    ("S29", "Educational attainment for workers 25+ by detailed occupation", "U.S. Bureau of Labor Statistics", "2025 National Employment Matrix", "https://www.bls.gov/emp/tables/educational-attainment.htm"),
    ("S30", "The Rise of GTM Engineering, in 3 Charts", "Onfire AI", "2026-08-31", "https://www.onfire.ai/blog/gtm-engineering-data"),
    ("S31", "GTM Engineer Salary & Market Report 2026", "RevGuild (Glozo data)", "2026-07-17 / 08-21", "https://revguild.org/blog/gtm-engineer-salary-market-report-2026/"),
    ("S32", "Companies that use Clay (web-intent detection)", "Bloomberry", "2026-10-18", "https://bloomberry.com/data/clay/"),
    ("S33", "Companies that use ZoomInfo (web-tracking detection)", "Bloomberry", "2026-10-15", "https://bloomberry.com/data/zoominfo/"),
    ("S34", "Companies listed as Clay Solution Partners", "Bloomberry", "2026-10-16", "https://bloomberry.com/data/clay-solution-partners"),
    ("S35", "Clay solution partner directory", "Clay Labs Inc.", "Live, Oct 2026", "https://www.clay.com/experts"),
    ("S36", "Clay CRM integrations", "Clay Labs Inc.", "Live, Oct 2026", "https://www.clay.com/integrations?provider-categories=crms-db5ec"),
    ("S37", "Google Maps prospecting in Clay", "Clay University", "Undated", "https://clay.com/university/lesson/scrape-google-maps"),
    ("S38", "Dodge for suppliers & distributors", "Dodge Construction Network", "Live, 2026", "https://www.construction.com/suppliers-distributors/"),
    ("S39", "46th Deltek Clarity A&E Industry Study", "Deltek", "2025-05-13", "https://www.deltek.com/en/about/media-center/press-releases/2025/what-the-46th-annual-deltek-clarity-ae-study-reveals-about-the-industry"),
    ("S40", "Most popular ERP systems for Top 200 distributors", "Electrical Wholesaling", "2014-06-23 (dated)", "https://www.ewweb.com/top-200/most-popular-erp-systems-top-200-distributors"),
    ("S41", "Clearlake completes $7.7B acquisition of Dun & Bradstreet", "Pulse 2.0", "2025-08-28", "https://pulse2.com/clearlake-7-7-billion-acquisition-of-dun-bradstreet-completed/amp/"),
    ("S42", "Unanet Q1 momentum (customer count)", "Unanet", "2022-04-28", "https://unanet.com/news/unanet-demonstrates-continued-momentum-in-first-quarter-winning-erp-and-crm-market-share-with-innovation-and-customer-service"),
]
for i, s in enumerate(srcs, start=5):
    for j, v in enumerate(s, start=1):
        put(so, f"{L(j)}{i}", v)

order = ["Cover", "Summary", "Competitor_Register", "Overlap_Check", "Comparison_Matrix", "Pricing_Compare",
         "Persona_Coverage", "Segment_Scoring", "Positioning_Map", "Sources"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Cover": "1F2A44", "Summary": "1F2A44", "Sources": "A6A6A6"}.get(w.title, "4472C4")
wb.save(OUT)
print("saved", OUT)
