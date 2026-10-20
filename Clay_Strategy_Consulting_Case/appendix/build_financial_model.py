"""Build 07_Financial_Model.xlsx — v2 (post partner review, 2026-10-19).

Excel engine = live formulas that mirror appendix/clay_engine.py formula for formula.
Python engine = Monte Carlo + gate policy + break-evens (values written to the 'Monte Carlo' sheet,
with a run-by-run reconciliation against the Excel engine).
v1 is archived at appendix/archive/. The Change Log sheet maps every change to the review issue it fixes.
"""
import math
import random
import sys
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import clay_engine as ce

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = f"{BASE}/07_Financial_Model.xlsx"
F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLACK = Font(name=F, color="000000", size=10)
GREEN = Font(name=F, color="008000", size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=F, bold=True, size=14, color="1F2A44")
SUB = Font(name=F, italic=True, size=9, color="555555")
HFILL = PatternFill("solid", fgColor="1F2A44")
SFILL = PatternFill("solid", fgColor="E8ECF3")
YFILL = PatternFill("solid", fgColor="FFF2CC")
KFILL = PatternFill("solid", fgColor="C6E0B4")
RFILL = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="top")
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
N1 = '#,##0.0;(#,##0.0);"-"'
N0 = '#,##0;(#,##0);"-"'
N3 = '0.000'
MO = '0.0" mo";(0.0" mo");"-"'

QCOLS = [L(3 + i) for i in range(12)]          # C..N
QLAB = [f"Q{(i % 4) + 1} {2027 + i // 4}" for i in range(12)]
ANN = [("P", "2027", 0, 4), ("Q", "2028", 4, 8), ("R", "2029", 8, 12)]
TOT = "S"
SEGS = ["dist", "ps"]
SEGLAB = {"dist": "Industrial distribution", "ps": "Professional services"}
SCN = ["Downside", "Base", "Upside"]
MOT = ["lean", "sales", "play"]
MOTLAB = {"lean": "Lean partner-led pilot", "sales": "Sales-led vertical team", "play": "Solution play (no vertical team)"}

wb = Workbook()


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, bold=False, wrap=True, border=True, center=False, italic=False):
    c = ws[ref]
    c.value = val
    c.font = Font(name=F, size=10, bold=bold, italic=italic, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = CENTER if center else (WRAP if wrap else Alignment(vertical="top"))
    if border:
        c.border = BOX
    return c


def hdr(ws, row, labels, col=1, widths=None, height=30):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR, HFILL, CENTER, BOX
    if widths:
        for i, w in enumerate(widths):
            if w:
                ws.column_dimensions[L(col + i)].width = w
    ws.row_dimensions[row].height = height


def title(ws, t, s):
    ws["A1"].value, ws["A1"].font = t, TITLE
    ws["A2"].value, ws["A2"].font = s, SUB
    ws.sheet_view.showGridLines = False


def note(ws, rng, text, height=34):
    a = rng.split(":")[0]
    put(ws, a, text, border=False)
    ws.merge_cells(rng)
    ws.row_dimensions[int("".join(ch for ch in a if ch.isdigit()))].height = height


def band_fill(ws, row, cols, fill):
    for c in cols:
        put(ws, f"{c}{row}", None, fill=fill)


cover = wb.active
cover.title = "Cover"
exo = wb.create_sheet("Executive Output")
asm = wb.create_sheet("Assumptions")
ms = wb.create_sheet("Market Sizing")
acvs = wb.create_sheet("ACV Build")
fun = wb.create_sheet("GTM Funnel")
rev = wb.create_sheet("Revenue Model")
ue = wb.create_sheet("Unit Economics")
scn = wb.create_sheet("Scenarios")
sen = wb.create_sheet("Sensitivity")
mc = wb.create_sheet("Monte Carlo")
eng = wb.create_sheet("Engine Runs")
chg = wb.create_sheet("Change Log")
src = wb.create_sheet("Sources")

# ============================================================ ASSUMPTIONS
title(asm, "Assumptions — every input, by segment and scenario (v2)",
      "Edit blue cells only. 'Live' feeds the main model via the three selectors. Scenarios, Sensitivity and Monte Carlo run their own copies. Segment-level market inputs are fixed across scenarios: the universe does not flex, only performance does.")
sel = [("D4", "Scenario (1 = Downside, 2 = Base, 3 = Upside)", 2, 3, '=CHOOSE($D$4,"Downside","Base","Upside")'),
       ("D5", "Segment (1 = Industrial distribution, 2 = Professional services)", 1, 2, '=CHOOSE($D$5,"Industrial distribution","Professional services")'),
       ("D6", "Motion (1 = Lean partner-led pilot, 2 = Sales-led vertical team, 3 = Solution play, no vertical team)", 1, 3, '=CHOOSE($D$6,"Lean partner-led pilot","Sales-led vertical team","Solution play")')]
for ref, lab, v, mx, f in sel:
    r = int(ref[1:])
    put(asm, f"A{r}", lab, bold=True, fill=SFILL)
    asm.merge_cells(f"A{r}:C{r}")
    put(asm, ref, v, font=BLUE, fill=YFILL, center=True, bold=True)
    put(asm, f"E{r}", f, center=True, bold=True)
    asm.merge_cells(f"E{r}:F{r}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(mx))
    asm.add_data_validation(dv)
    dv.add(ref)
SEL_S, SEL_G, SEL_M = "Assumptions!$D$4", "Assumptions!$D$5", "Assumptions!$D$6"
hdr(asm, 8, ["ID", "Input", "Unit", "Dist: Downside", "Dist: Base", "Dist: Upside", "PS: Downside", "PS: Base", "PS: Upside", "Live", "Class", "MC group", "Source / rationale"],
    widths=[6, 52, 9, 11, 11, 11, 11, 11, 11, 11, 7, 9, 80])
A = {}
SCOL = {("dist", 0): "D", ("dist", 1): "E", ("dist", 2): "F", ("ps", 0): "G", ("ps", 1): "H", ("ps", 2): "I"}
SECTION_BREAKS = {"f1": "MARKET (fixed across scenarios)", "cov": "DEMAND AND FUNNEL", "salesshare": "ACV BUILD (usage-based pricing)",
                  "grrS": "RETENTION AND EXPANSION (annual)", "quota": "SALES CAPACITY AND COST RATES", "disc": "VALUATION, INCREMENTALITY AND OPTION 0"}
r = 9
for x in ce.INPUTS:
    name, unit, cls, grp, lab, srcx = x[0], x[7], x[8], x[9], x[10], x[11]
    if name in SECTION_BREAKS:
        put(asm, f"A{r}", SECTION_BREAKS[name], bold=True, fill=SFILL)
        band_fill(asm, r, "BCDEFGHIJKLM", SFILL)
        r += 1
    A[name] = r
    fmt = PCT if unit == "%" else (USD if unit.startswith("$") else ('0.00' if unit in ("x", "x/yr") else N0))
    put(asm, f"A{r}", f"A{len(A)}")
    put(asm, f"B{r}", lab)
    put(asm, f"C{r}", unit, center=True)
    for k, col in enumerate("DEFGHI"):
        put(asm, f"{col}{r}", x[1 + k], font=BLUE, fmt=fmt, center=True, fill=YFILL if cls in ("D", "E", "B/D") else None)
    put(asm, f"J{r}", f"=INDEX(D{r}:I{r},1,({SEL_G}-1)*3+{SEL_S})", fmt=fmt, center=True, bold=True)
    put(asm, f"K{r}", cls, center=True)
    put(asm, f"L{r}", grp or "fixed", center=True)
    put(asm, f"M{r}", srcx)
    r += 1
r += 1
put(asm, f"A{r}", "MOTION DESIGN (team, programs, scope)", bold=True, fill=SFILL)
band_fill(asm, r, "BCDEFGHIJKLM", SFILL)
r += 1
hdr(asm, r, ["ID", "Input", "Unit", "Lean pilot", "Sales-led", "Solution play", "", "", "", "Live", "Class", "", "Rationale"])
r += 1
M = {}
MOTION_ROWS = [("fte_lead", "Vertical lead (FTE)", "FTE", "Owns the vertical P&L; none in the solution play"),
               ("pm1", "Partner managers — 2027 (FTE)", "FTE", ""), ("pm2", "Partner managers — 2028+ (FTE)", "FTE", ""),
               ("eng1", "GTM / solutions engineers — 2027 (FTE)", "FTE", ""), ("eng2", "GTM / solutions engineers — 2028+ (FTE)", "FTE", ""),
               ("pmm", "Product marketing (FTE)", "FTE", ""),
               ("prog", "Program marketing per quarter", "$", "Above $25K buys extra MQLs at the CPL (diminishing); below $25K scales inbound down"),
               ("shared", "Shared AE capacity borrowed from the core team", "AE-eq.", "Lean and play close with existing AEs at a variable cost per $ closed"),
               ("varae", "Charge variable AE cost (1 = yes)", "flag", ""),
               ("covmult", "Automated outbound coverage multiplier", "x", "Play: core team runs half the outbound"),
               ("partmult", "Partner network multiplier", "x", "Play: no one owns partner recruitment"),
               ("buildmult", "Template & connector build multiplier", "x", "Play: templates only, no connector")]
for name, lab, unit, why in MOTION_ROWS:
    M[name] = r
    fmt = USD if unit == "$" else '0.00'
    put(asm, f"A{r}", f"M{len(M)}")
    put(asm, f"B{r}", lab)
    put(asm, f"C{r}", unit, center=True)
    for col, mo in zip("DEF", MOT):
        put(asm, f"{col}{r}", ce.MOTIONS[mo][name], font=BLUE, fmt=fmt, center=True, fill=YFILL)
    put(asm, f"J{r}", f"=INDEX(D{r}:F{r},1,{SEL_M})", fmt=fmt, center=True, bold=True)
    put(asm, f"K{r}", "D/E", center=True)
    put(asm, f"M{r}", why)
    r += 1
r += 1
put(asm, f"A{r}", "Dedicated AE headcount by quarter (FTE)", bold=True, fill=SFILL)
asm.merge_cells(f"A{r}:C{r}")
for i, q in enumerate(QLAB):
    put(asm, f"{L(4 + i)}{r}", q, bold=True, center=True, fill=SFILL)
r += 1
HIRE = {}
for key, lab in [("lean", "Lean pilot"), ("sales", "Sales-led team"), ("play", "Solution play"), ("scale", "Lean pilot + scale-up after Gate 1 (gated runs)")]:
    put(asm, f"B{r}", lab, bold=True)
    for i, v in enumerate(ce.HIRES[key]):
        put(asm, f"{L(4 + i)}{r}", v, font=BLUE, center=True, fill=YFILL)
    HIRE[key] = r
    r += 1
put(asm, f"B{r}", "Live (by motion selector)", bold=True)
for i in range(12):
    c = L(4 + i)
    put(asm, f"{c}{r}", f"=INDEX({c}{HIRE['lean']}:{c}{HIRE['play']},{SEL_M},1)", center=True, bold=True)
HIRE["live"] = r
r += 2
TH = {"dist": dict(opps=30, cust=4, pl=31, acv=7700, win=.18, min_dec=10, s_acv=10000, s_cust=14, s_win=.25),
      "ps": dict(opps=67, cust=8, pl=41, acv=6300, win=.155, min_dec=10, s_acv=6600, s_cust=25, s_win=.215)}
put(asm, f"A{r}", "GATE 1 RULE (Q4 2027) — used by the Monte Carlo staged policy. Thresholds = midpoints between the Downside/Base (stop) and Base/Upside (scale) model paths.", bold=True, fill=SFILL)
band_fill(asm, r, "BCDEFGHIJKLM", SFILL)
r += 1
hdr(asm, r, ["", "Test", "Unit", "Distribution", "Prof. services", "", "", "", "", "", "Class", "", "Logic"])
r += 1
for key, lab, unit, fmt, logic in [("opps", "STOP test: opportunities created in 2027 below", "#", N0, "STOP if any two STOP tests fail"),
                                   ("cust", "STOP test: new customers in 2027 below", "#", N0, ""),
                                   ("pl", "STOP test: partner introductions in 2027 below", "#", N0, ""),
                                   ("acv", "STOP test: average first-year ACV below", "$", USD, ""),
                                   ("win", "STOP test: observed win rate below (only with ≥ 10 decided opportunities)", "%", PCT, ""),
                                   ("min_dec", "Minimum decided opportunities to use the win rate", "#", N0, ""),
                                   ("s_acv", "SCALE test: average first-year ACV at least", "$", USD, "SCALE (hire 2 AEs in H1 2028) if every SCALE test passes and no STOP"),
                                   ("s_cust", "SCALE test: new customers in 2027 at least", "#", N0, ""),
                                   ("s_win", "SCALE test: observed win rate at least", "%", PCT, "")]:
    put(asm, f"B{r}", lab)
    put(asm, f"C{r}", unit, center=True)
    put(asm, f"D{r}", TH["dist"][key], font=BLUE, fmt=fmt, center=True, fill=YFILL)
    put(asm, f"E{r}", TH["ps"][key], font=BLUE, fmt=fmt, center=True, fill=YFILL)
    put(asm, f"K{r}", "E", center=True)
    put(asm, f"M{r}", logic)
    r += 1
asm.freeze_panes = "D9"

# ============================================================ MARKET SIZING
title(ms, "Market Sizing — unique firms with 20+ employees by wave and size band (Census 2022, S14)",
      "Distribution waves: 4236 electrical → 4237 PHVAC → 4238 machinery/MRO. Professional-services waves: 5413 A&E → 5416 consulting → 5415 IT services (5418 advertising/PR excluded: Clay's agency partner channel). Firms, not establishments.")
hdr(ms, 4, ["NAICS", "Sub-industry", "Firms 20–49", "50–99", "100–249", "250–499", "500+", "Emp. 20–49", "50–99", "100–249", "250–499", "500+"],
    widths=[10, 36, 11, 10, 10, 10, 10, 12, 11, 11, 11, 11])
NAMES = {"4236": "Electrical & electronic goods wholesalers", "4237": "Hardware, plumbing & HVAC wholesalers", "4238": "Machinery, equipment & supplies wholesalers",
         "5413": "Architectural & engineering services", "5416": "Management & technical consulting", "5415": "Computer systems design (IT services)"}
CROW = {}
rr = 5
for n in ["4236", "4237", "4238", "5413", "5416", "5415"]:
    CROW[n] = rr
    put(ms, f"A{rr}", n, bold=True)
    put(ms, f"B{rr}", NAMES[n])
    for j, v in enumerate(ce.CENSUS[n][0] + ce.CENSUS[n][1]):
        put(ms, f"{L(3 + j)}{rr}", v, font=BLUE, fmt=N0, center=True)
    rr += 1
rr += 1
put(ms, f"A{rr}", "De-duplication factor by band (sector firms ÷ Σ sub-industry firms)", bold=True)
ms.merge_cells(f"A{rr}:B{rr}")
rr += 1
DDROW = {}
for seg, lab in [("dist", "Sector 423 (50–499 not measurable: suppressed cells → 1.0)"), ("ps", "Sector 54")]:
    DDROW[seg] = rr
    put(ms, f"A{rr}", seg.upper())
    put(ms, f"B{rr}", lab)
    for j, v in enumerate(ce.DEDUP[seg]):
        put(ms, f"{L(3 + j)}{rr}", v, font=BLUE, fmt=N3, center=True)
    rr += 1
rr += 1
hdr(ms, rr, ["Wave", "Sub-industry", "SMB 20–99", "Mid-Market 100–499", "Enterprise 500+", "Total"])
rr += 1
WROW = {}
for seg in SEGS:
    for w, n in enumerate(ce.SEGMENTS[seg]["waves"]):
        WROW[(seg, w)] = rr
        c, d = CROW[n], DDROW[seg]
        put(ms, f"A{rr}", f"{seg.upper()} W{w + 1}", bold=True)
        put(ms, f"B{rr}", f"=B{c}")
        put(ms, f"C{rr}", f"=C{c}*C${d}+D{c}*D${d}", fmt=N0, center=True)
        put(ms, f"D{rr}", f"=E{c}*E${d}+F{c}*F${d}", fmt=N0, center=True)
        put(ms, f"E{rr}", f"=G{c}*G${d}", fmt=N0, center=True)
        put(ms, f"F{rr}", f"=SUM(C{rr}:E{rr})", fmt=N0, center=True, bold=True)
        rr += 1
rr += 1
hdr(ms, rr, ["Segment", "Band", "Firms (3 waves)", "Employees", "Avg employees per firm", ""])
rr += 1
EMP = {}
for seg in SEGS:
    rows_ = [CROW[n] for n in ce.SEGMENTS[seg]["waves"]]
    for b, fc, ec in (("S", "CD", "HI"), ("M", "EF", "JK"), ("E", "G", "L")):
        put(ms, f"A{rr}", SEGLAB[seg])
        put(ms, f"B{rr}", {"S": "SMB 20–99", "M": "Mid-Market 100–499", "E": "Enterprise 500+"}[b])
        put(ms, f"C{rr}", "=" + "+".join(f"{c}{x}" for c in fc for x in rows_), fmt=N0, center=True)
        put(ms, f"D{rr}", "=" + "+".join(f"{c}{x}" for c in ec for x in rows_), fmt=N0, center=True)
        put(ms, f"E{rr}", f"=D{rr}/C{rr}", fmt=N1, center=True, bold=True)
        EMP[(seg, b)] = f"'Market Sizing'!$E${rr}"
        rr += 1
rr += 1
note(ms, f"A{rr}:L{rr}", "Serviceable accounts (SAM) are computed inside each model run as firms × filter 1 × filter 2 × (1 − existing Clay customers), with the filters fixed across scenarios (v2, review M1/M2). "
     "Known limitation (review M3): Census sizes firms by total firm employment, and NAICS wholesale includes manufacturers' sales branches, so 500+ 'distributors' include manufacturers' sales arms: 4–22% of the band per Census type of operation (02 Firm_Counts Block 7). Removing them moves the lean-pilot Base NPV by −$0.01M to −$0.03M (appendix/msbo_sensitivity.py).", 46)


def FIRM(seg, w, b):
    return f"'Market Sizing'!${'CDE'['SME'.index(b)]}${WROW[(seg, w)]}"


def acv_formula(seg, band, g):
    emp = EMP[(seg, band)]
    cp = f"({g('dcprice')}*12/{g('dcsize')})"
    ap = f"({g('acprice')}*12/{g('acsize')})"
    acc = f"({emp}*{g('salesshare')}*{g('aps')}*{g('usage')})"
    con = f"({acc}*{g('cpa')})"
    cred = f"({con}*{g('refresh')}*{g('crc')}+{acc}*{g('ai')}+{acc}*{g('newacc')}*{g('crnew')})"
    ops = f"({con}*{g('refresh')}+{acc}+{acc}*{g('newacc')})"
    use = f"({cred}*{cp}+{ops}*{g('apo')}*{ap})"
    if band == "E":
        return f"MAX({g('entcr')}*{cp}+{g('entac')}*{ap},{use})*(1-{g('entdisc')})"
    return f"MAX(({g('dcprice')}+{g('acprice')})*12,{use})"


def g_col(col):
    return lambda n: f"Assumptions!${col}${A[n]}"


title(acvs, "ACV Build — usage-based first-year ACV by segment, band and scenario",
      "Sellers per firm × accounts tracked × Clay's published usage meters (S01). SMB/Mid-Market floored at the Growth plan; Enterprise floored at the Enterprise minimum, less discount.")
hdr(acvs, 4, ["Band", "Dist: Downside", "Dist: Base", "Dist: Upside", "PS: Downside", "PS: Base", "PS: Upside", "Live"], widths=[34, 14, 14, 14, 14, 14, 14, 14])
for k, b in enumerate("SME"):
    rr = 5 + k
    put(acvs, f"A{rr}", {"S": "SMB (20–99)", "M": "Mid-Market (100–499)", "E": "Enterprise (500+)"}[b], bold=True)
    for j, (seg, s) in enumerate([("dist", 0), ("dist", 1), ("dist", 2), ("ps", 0), ("ps", 1), ("ps", 2)]):
        put(acvs, f"{L(2 + j)}{rr}", "=" + acv_formula(seg, b, g_col(SCOL[(seg, s)])), fmt=USD, center=True)
    put(acvs, f"H{rr}", f"=INDEX(B{rr}:G{rr},1,({SEL_G}-1)*3+{SEL_S})", fmt=USD, center=True, bold=True)
put(acvs, "A9", "Sellers per firm (Base)", bold=True)
for k, b in enumerate("SME"):
    put(acvs, f"A{10 + k}", {"S": "SMB", "M": "Mid-Market", "E": "Enterprise"}[b])
    put(acvs, f"C{10 + k}", f"={EMP[('dist', b)]}*Assumptions!$E${A['salesshare']}", fmt=N1, center=True)
    put(acvs, f"F{10 + k}", f"={EMP[('ps', b)]}*Assumptions!$H${A['salesshare']}", fmt=N1, center=True)
note(acvs, "A14:H14", "Reading: at list prices, a distributor's seller density lifts ACV above the Growth floor only from Mid-Market up. Professional-services firms have ~4× fewer sellers per firm, so their SMB and Mid-Market ACV both sit at the floor. Base Enterprise ACV for distribution ≈ Vendr's median Clay buyer ($59,996, S16).", 46)


# ============================================================ ENGINE
class Engine:
    def __init__(self, stages, get, mget, hire, label, seg, firm_fn, acv_fn):
        self.stages, self.get, self.mget, self.hire, self.label, self.seg = stages, get, mget, hire, label, seg
        self.firm, self.acvf = firm_fn, acv_fn
        self.reg, self.cur, self.out = {}, None, {}

    def _r(self, key, col, absolute=False):
        ws, row = self.reg[key]
        cur_ws = self.stages[self.cur][0]
        c = f"${col}${row}" if absolute else f"{col}{row}"
        return c if ws.title == cur_ws.title else f"'{ws.title}'!{c}"

    def rng(self, key):
        ws, row = self.reg[key]
        cur_ws = self.stages[self.cur][0]
        c = f"$C${row}:$N${row}"
        return c if ws.title == cur_ws.title else f"'{ws.title}'!{c}"

    def S(self, key):
        return self._r(key, "C", absolute=True)

    def Q(self, key, i):
        return self._r(key, QCOLS[i])

    def header(self, stage, text):
        ws, row = self.stages[stage]
        put(ws, f"A{row}", text, bold=True, font=Font(color="FFFFFF"), fill=HFILL)
        put(ws, f"B{row}", "Unit / note", font=Font(color="FFFFFF"), fill=HFILL, bold=True)
        for i, q in enumerate(QLAB):
            put(ws, f"{QCOLS[i]}{row}", q, font=Font(color="FFFFFF"), fill=HFILL, bold=True, center=True)
        for col, y, _, _ in ANN:
            put(ws, f"{col}{row}", y, font=Font(color="FFFFFF"), fill=HFILL, bold=True, center=True)
        put(ws, f"{TOT}{row}", "2027–29", font=Font(color="FFFFFF"), fill=HFILL, bold=True, center=True)
        self.stages[stage][1] += 1

    def scalar(self, stage, key, label, formula, fmt, note_=""):
        self.cur = stage
        ws, row = self.stages[stage]
        self.reg[key] = (ws, row)
        put(ws, f"A{row}", label)
        put(ws, f"B{row}", note_, italic=True)
        put(ws, f"C{row}", formula() if callable(formula) else formula, fmt=fmt, center=True)
        self.stages[stage][1] += 1

    def row(self, stage, key, label, unit, fn, fmt, agg="sum", bold=False):
        self.cur = stage
        ws, row = self.stages[stage]
        self.reg[key] = (ws, row)
        put(ws, f"A{row}", label, bold=bold)
        put(ws, f"B{row}", unit, italic=True)
        for i in range(12):
            put(ws, f"{QCOLS[i]}{row}", fn(i), fmt=fmt, wrap=False, bold=bold)
            ws[f"{QCOLS[i]}{row}"].alignment = RIGHT
        for col, y, a, b in ANN:
            f = {"sum": f"=SUM({QCOLS[a]}{row}:{QCOLS[b - 1]}{row})", "end": f"={QCOLS[b - 1]}{row}", "avg": f"=AVERAGE({QCOLS[a]}{row}:{QCOLS[b - 1]}{row})"}.get(agg)
            put(ws, f"{col}{row}", f, fmt=fmt, bold=True, wrap=False)
            ws[f"{col}{row}"].alignment = RIGHT
        f = {"sum": f"=SUM(P{row}:R{row})", "end": f"=R{row}"}.get(agg)
        put(ws, f"{TOT}{row}", f, fmt=fmt, bold=True, wrap=False, fill=SFILL)
        ws[f"{TOT}{row}"].alignment = RIGHT
        self.stages[stage][1] += 1

    def outrow(self, stage, key, label, formula, fmt):
        self.cur = stage
        ws, row = self.stages[stage]
        put(ws, f"A{row}", label, bold=True)
        put(ws, f"C{row}", formula, fmt=fmt, bold=True, fill=KFILL, center=True)
        ws.merge_cells(f"C{row}:D{row}")
        self.reg[key] = (ws, row)
        self.out[key] = f"'{ws.title}'!$C${row}"
        self.stages[stage][1] += 1

    def build(self):
        g, mg, S, Q, seg = self.get, self.mget, self.S, self.Q, self.seg

        def prev(key, i):
            return Q(key, i - 1) if i > 0 else "0"

        self.header("fun", self.label + " — run constants")
        for w in range(3):
            for b in "SME":
                self.scalar("fun", f"sam{w}{b}", f"SAM accounts — wave {w + 1}, {b}", f"={self.firm(seg, w, b)}*{g('f1')}*{g('f2')}*(1-{g('exist')})", N1, "firms × filters")
            self.scalar("fun", f"samt{w}", f"SAM accounts — wave {w + 1}, total", lambda w=w: f"={S(f'sam{w}S')}+{S(f'sam{w}M')}+{S(f'sam{w}E')}", N1, "")
        for b in "SME":
            self.scalar("fun", f"acv{b}", f"First-year ACV — {b}", "=" + self.acvf(seg, b, g), USD, "usage build")
        self.scalar("fun", "cov", "Automated outbound coverage per quarter", f"={g('cov')}*{mg('covmult')}", PCT, "coverage × motion multiplier")
        self.scalar("fun", "aeload", "AE fully loaded annual cost", f"={g('aeote')}/{g('wshare')}", USD, "OTE ÷ wage share")
        self.scalar("fun", "capae", "ARR one ramped AE closes per quarter", f"={g('quota')}*{g('prod')}/4", USD, "quota × productivity ÷ 4")
        self.scalar("fun", "costper", "AE cost per $ of first-year ACV closed", lambda: f"={S('aeload')}/{g('quota')}/{g('prod')}", '0.000', "")
        self.scalar("fun", "extra", "Program spend above base, per quarter", f"=MAX(0,{mg('prog')}-{g('progbase')})", USD, "")

        self.header("fun", "Demand and funnel")
        self.row("fun", "qn", "Quarter number", "#", lambda i: i + 1, N0, agg=None)
        self.row("fun", "run", "Program running (1) / stopped (0)", "flag", lambda i: f"=IF({Q('qn', i)}<={g('stop')},1,0)", N0, agg=None)
        for w, nm in enumerate(["w1", "w2", "w3"]):
            self.row("fun", f"ramp{w}", f"Wave {w + 1} active (0.5 in launch quarter)", "factor", lambda i, nm=nm: f"=IF({Q('qn', i)}={g(nm)},0.5,IF({Q('qn', i)}>{g(nm)},1,0))", '0.0', agg=None)
        for w, nm in enumerate(["w1", "w2", "w3"]):
            self.row("fun", f"fat{w}", f"Wave {w + 1} outbound fatigue factor", "factor", lambda i, nm=nm: f"={g('fatigue')}^(MAX(0,{Q('qn', i)}-{g(nm)})/4)", '0.00', agg=None)
        for b in "SME":
            self.row("fun", f"act{b}", f"Active SAM accounts — {b}", "accounts",
                     lambda i, b=b: "=" + "+".join(f"{S(f'sam{w}{b}')}*{Q(f'ramp{w}', i)}" for w in range(3)), N0, agg="end")
        self.row("fun", "act", "Active SAM accounts — total", "accounts", lambda i: f"={Q('actS', i)}+{Q('actM', i)}+{Q('actE', i)}", N0, agg="end")
        self.row("fun", "efw", "Fatigue-weighted active accounts", "accounts", lambda i: "=" + "+".join(f"{S(f'samt{w}')}*{Q(f'ramp{w}', i)}*{Q(f'fat{w}', i)}" for w in range(3)), N0, agg=None)
        self.row("fun", "lo", "Meetings — automated outbound", "meetings", lambda i: f"={S('cov')}*{g('meet')}*{Q('run', i)}*{Q('efw', i)}", N1)
        self.row("fun", "ae", "Dedicated AEs (headcount)", "FTE", lambda i: f"={self.hire(i)}", N1, agg="end")

        def ramped(i):
            if i == 0:
                return "=0"
            if i == 1:
                return f"=0.5*{Q('ae', 0)}"
            return f"={Q('ae', i - 2)}+0.5*({Q('ae', i - 1)}-{Q('ae', i - 2)})"
        self.row("fun", "aer", "Ramped dedicated AEs (≈6-month ramp)", "FTE", ramped, N1, agg="end")
        self.row("fun", "aeacc", "Accounts worked by dedicated AEs", "accounts", lambda i: f"=MIN({Q('aer', i)}*{g('aeacc')},{Q('act', i)}*(1-{S('cov')}))", N0)
        self.row("fun", "lae", "Meetings — AE self-sourced outreach", "meetings",
                 lambda i: f"={Q('aeacc', i)}*{g('meet')}*{g('aemult')}*IF({Q('act', i)}>0,{Q('efw', i)}/{Q('act', i)},0)*{Q('run', i)}", N1)
        self.row("fun", "xm", "Extra MQLs bought by program spend above base", "MQLs",
                 lambda i: f"=IF(AND({Q('act', i)}>0,{S('extra')}>0),{Q('act', i)}*{g('inbcap')}*(1-EXP(-{S('extra')}/({g('cpl')}*{Q('act', i)}*{g('inbcap')}))),0)", N1)
        self.row("fun", "li", "MQLs — inbound / templates", "MQLs", lambda i: f"=({Q('act', i)}*{g('inb')}*MIN(1,{mg('prog')}/{g('progbase')})+{Q('xm', i)})*{Q('run', i)}", N1)
        self.row("fun", "parts", "Active partners", "partners", lambda i: f"={Q('run', i)}*MIN({g('pcap')},{g('p0')}+{g('padd')}*({Q('qn', i)}-1))*{mg('partmult')}", N1, agg="end")
        self.row("fun", "lp", "Partner introductions", "intros", lambda i: f"={Q('parts', i)}*{g('lpp')}", N1)
        self.row("fun", "oo", "Opportunities — outbound (automated + AE)", "opps", lambda i: f"=({Q('lo', i)}+{Q('lae', i)})*{g('m2o')}", N1)
        self.row("fun", "oi", "Opportunities — inbound", "opps", lambda i: f"={Q('li', i)}*{g('i2o')}", N1)
        self.row("fun", "op", "Opportunities — partner", "opps", lambda i: f"={Q('lp', i)}*{g('p2o')}", N1)
        self.row("fun", "opp", "Opportunities — total", "opps", lambda i: f"={Q('oo', i)}+{Q('oi', i)}+{Q('op', i)}", N1, bold=True)
        self.row("fun", "wS", "Band weight — SMB (accounts)", "weight", lambda i: f"={Q('actS', i)}", N1, agg=None)
        self.row("fun", "wM", "Band weight — Mid-Market (accounts × tilt × propensity)", "weight", lambda i: f"={Q('actM', i)}*{g('tiltM')}*{g('propM')}", N1, agg=None)
        self.row("fun", "wE", "Band weight — Enterprise (accounts × tilt × propensity)", "weight", lambda i: f"={Q('actE', i)}*{g('tiltE')}*{g('propE')}", N1, agg=None)
        for b in "SME":
            self.row("fun", f"mix{b}", f"Opportunity mix — {b} (derived, not assumed)", "%",
                     lambda i, b=b: f"=IFERROR({Q(f'w{b}', i)}/({Q('wS', i)}+{Q('wM', i)}+{Q('wE', i)}),0)", PCT, agg=None)

        def lagged(b, lagname, i):
            return (f"=IF({Q('qn', i)}-{g(lagname)}<1,0,INDEX({self.rng('opp')},1,{Q('qn', i)}-{g(lagname)})"
                    f"*INDEX({self.rng('mix' + b)},1,{Q('qn', i)}-{g(lagname)})*{g('win')})")
        self.row("fun", "dwS", "Wins (demand) — SMB", "wins", lambda i: lagged("S", "lagSM", i), N1)
        self.row("fun", "dwM", "Wins (demand) — Mid-Market", "wins", lambda i: lagged("M", "lagSM", i), N1)
        self.row("fun", "dwE", "Wins (demand) — Enterprise (longer cycle)", "wins", lambda i: lagged("E", "lagE", i), N1)
        self.row("fun", "psh", "Partner-sourced share of closing opportunities", "%",
                 lambda i: f"=IF({Q('qn', i)}-{g('lagSM')}<1,0,IFERROR(INDEX({self.rng('op')},1,{Q('qn', i)}-{g('lagSM')})/INDEX({self.rng('opp')},1,{Q('qn', i)}-{g('lagSM')}),0))", PCT, agg=None)
        self.row("fun", "ssw", "Self-serve SMB wins (no AE)", "wins", lambda i: f"={Q('dwS', i)}*{g('ss')}", N1)
        self.row("fun", "sadS", "Sales-assisted demand — SMB", "wins", lambda i: f"={Q('dwS', i)}*(1-{g('ss')})", N1)
        self.row("fun", "sadv", "Sales-assisted demand (first-year ACV)", "$", lambda i: f"={Q('sadS', i)}*{S('acvS')}+{Q('dwM', i)}*{S('acvM')}+{Q('dwE', i)}*{S('acvE')}", USD)
        self.row("fun", "capv", "Closing capacity (dedicated + shared AEs)", "$", lambda i: f"=({Q('aer', i)}+{mg('shared')}*{Q('run', i)})*{S('capae')}", USD)
        self.row("fun", "fac", "Share of demand closed = MIN(1, capacity ÷ demand)", "%", lambda i: f"=IF({Q('sadv', i)}>0,MIN(1,{Q('capv', i)}/{Q('sadv', i)}),1)", PCT, agg=None)
        self.row("fun", "lost", "First-year ACV lost to the capacity constraint", "$", lambda i: f"={Q('sadv', i)}*(1-{Q('fac', i)})", USD)
        self.row("fun", "nS", "New customers — SMB", "customers", lambda i: f"={Q('ssw', i)}+{Q('sadS', i)}*{Q('fac', i)}", N1)
        self.row("fun", "nM", "New customers — Mid-Market", "customers", lambda i: f"={Q('dwM', i)}*{Q('fac', i)}", N1)
        self.row("fun", "nE", "New customers — Enterprise", "customers", lambda i: f"={Q('dwE', i)}*{Q('fac', i)}", N1)
        self.row("fun", "new", "New customers — total", "customers", lambda i: f"={Q('nS', i)}+{Q('nM', i)}+{Q('nE', i)}", N1, bold=True)

        self.header("rev", "Customers, ARR and revenue")
        for b, gl in (("S", "SMB"), ("M", "Mid-Market"), ("E", "Enterprise")):
            self.row("rev", f"c{b}", f"Customers — {gl}", "customers", lambda i, b=b: f"={prev('c' + b, i)}*{g('grr' + b)}^(1/4)+{Q('n' + b, i)}", N1, agg="end")
        self.row("rev", "cT", "Customers — total", "customers", lambda i: f"={Q('cS', i)}+{Q('cM', i)}+{Q('cE', i)}", N1, agg="end", bold=True)
        for b, gl in (("S", "SMB"), ("M", "Mid-Market"), ("E", "Enterprise")):
            self.row("rev", f"a{b}", f"ARR — {gl}", "$", lambda i, b=b: f"={prev('a' + b, i)}*{g('nrr' + b)}^(1/4)+{Q('n' + b, i)}*{S('acv' + b)}", USD, agg="end")
        self.row("rev", "aT", "ARR — total (end of quarter)", "$", lambda i: f"={Q('aS', i)}+{Q('aM', i)}+{Q('aE', i)}", USD, agg="end", bold=True)
        self.row("rev", "na", "New ARR booked (first-year ACV)", "$", lambda i: f"={Q('nS', i)}*{S('acvS')}+{Q('nM', i)}*{S('acvM')}+{Q('nE', i)}*{S('acvE')}", USD)
        self.row("rev", "rv", "Recognized revenue (average ARR ÷ 4)", "$", lambda i: f"=({prev('aT', i)}+{Q('aT', i)})/8", USD, bold=True)
        self.row("rev", "gp", "Gross profit", "$", lambda i: f"={Q('rv', i)}*{g('gm')}", USD)

        self.header("ue", "Incremental program cost")
        self.row("ue", "fl", "FTE — vertical lead", "FTE", lambda i: f"={mg('fte_lead')}*{Q('run', i)}", N1, agg="avg")
        self.row("ue", "fa", "FTE — dedicated AEs", "FTE", lambda i: f"={Q('ae', i)}*{Q('run', i)}", N1, agg="avg")
        self.row("ue", "fp", "FTE — partner managers", "FTE", lambda i: f"=IF({Q('qn', i)}<=4,{mg('pm1')},{mg('pm2')})*{Q('run', i)}", N1, agg="avg")
        self.row("ue", "fe", "FTE — GTM / solutions engineers", "FTE", lambda i: f"=IF({Q('qn', i)}<=4,{mg('eng1')},{mg('eng2')})*{Q('run', i)}", N1, agg="avg")
        self.row("ue", "fm", "FTE — product marketing", "FTE", lambda i: f"={mg('pmm')}*{Q('run', i)}", N1, agg="avg")
        self.row("ue", "fc", "FTE — customer success", "FTE", lambda i: f"=({Q('cM', i)}+{Q('cE', i)})/{g('csmM')}+{Q('cS', i)}/{g('csmS')}", N1, agg="avg")
        self.row("ue", "kl", "Cost — vertical lead + partner managers", "$", lambda i: f"=({Q('fl', i)}+{Q('fp', i)})*{g('wmgr')}/{g('wshare')}/4", USD)
        self.row("ue", "ka", "Cost — dedicated AEs", "$", lambda i: f"={Q('fa', i)}*{S('aeload')}/4", USD)
        self.row("ue", "kv", "Cost — variable AE cost (shared capacity)", "$",
                 lambda i: f"=MAX(0,({Q('na', i)}-{Q('ssw', i)}*{S('acvS')})-{Q('aer', i)}*{S('capae')})*{S('costper')}*{mg('varae')}", USD)
        self.row("ue", "km", "Cost — product marketing", "$", lambda i: f"={Q('fm', i)}*{g('wmkt')}/{g('wshare')}/4", USD)
        self.row("ue", "kp", "Cost — programs and launches", "$",
                 lambda i: f"=({mg('prog')}+{g('launchsp')}*(({Q('qn', i)}={g('w1')})+({Q('qn', i)}={g('w2')})+({Q('qn', i)}={g('w3')})))*{Q('run', i)}", USD)
        self.row("ue", "kr", "Cost — partner fees", "$", lambda i: f"={Q('na', i)}*{Q('psh', i)}*{g('ref')}", USD)
        self.row("ue", "kn", "Cost — partner enablement", "$", lambda i: f"={Q('parts', i)}*{g('enab')}", USD)
        self.row("ue", "ko", "Cost — onboarding subsidy (MM & Ent)", "$", lambda i: f"=({Q('nM', i)}*{S('acvM')}+{Q('nE', i)}*{S('acvE')})*{g('onb')}", USD)
        self.row("ue", "sm", "Sales & marketing cost (CAC basis)", "$", lambda i: "=" + "+".join(Q(k, i) for k in ("kl", "ka", "kv", "km", "kp", "kr", "kn", "ko")), USD, bold=True)
        self.row("ue", "ke", "Cost — GTM / solutions engineers", "$", lambda i: f"={Q('fe', i)}*{g('wdev')}/{g('wshare')}/4", USD)
        self.row("ue", "kb", "Cost — template & connector build", "$", lambda i: f"=IF({Q('qn', i)}<={g('buildq')},{g('build')}*{mg('buildmult')}/{g('buildq')},0)", USD)
        self.row("ue", "kc", "Cost — customer success", "$", lambda i: f"={Q('fc', i)}*{g('wcs')}/{g('wshare')}/4", USD)
        self.row("ue", "kt", "Total incremental program cost", "$", lambda i: f"={Q('sm', i)}+{Q('ke', i)}+{Q('kb', i)}+{Q('kc', i)}", USD, bold=True)
        self.header("ue", "Program P&L, NPV and the Option 0 benchmark")
        self.row("ue", "pl", "Program P&L (gross profit − incremental cost)", "$", lambda i: f"={Q('gp', i)}-{Q('kt', i)}", USD, bold=True)
        self.row("ue", "cum", "Cumulative program P&L", "$", lambda i: f"={prev('cum', i)}+{Q('pl', i)}", USD, agg="end", bold=True)
        self.row("ue", "df", "Discount factor (end of quarter)", "x", lambda i: f"=1/(1+{g('disc')})^({Q('qn', i)}/4)", '0.000', agg=None)
        self.row("ue", "pvi", "PV of incremental P&L (gross profit × (1 − haircut) − cost)", "$", lambda i: f"=({Q('gp', i)}*(1-{g('hair')})-{Q('kt', i)})*{Q('df', i)}", USD)
        self.row("ue", "n0", "Option 0: core ARR bought with the same non-CS spend", "$", lambda i: f"=({Q('kt', i)}-{Q('kc', i)})*12/({g('pb0')}*{g('gm')})*{g('eff0')}", USD)
        self.row("ue", "a0", "Option 0: core ARR (end of quarter)", "$", lambda i: f"={prev('a0', i)}*{g('nrr0')}^(1/4)+{Q('n0', i)}", USD, agg="end")
        self.row("ue", "pl0", "Option 0: P&L (gross profit − spend − CS)", "$",
                 lambda i: f"=({prev('a0', i)}+{Q('a0', i)})/8*{g('gm')}-({Q('kt', i)}-{Q('kc', i)})-{Q('a0', i)}*{g('csr0')}/4", USD)
        self.row("ue", "pv0", "Option 0: PV of P&L", "$", lambda i: f"={Q('pl0', i)}*{Q('df', i)}", USD)

        def tot(key):
            return self._r(key, TOT)
        self.scalar("ue", "grrb", "ARR-weighted GRR of new customers", lambda: (
            f"=IFERROR(({tot('nS')}*{S('acvS')}*{g('grrS')}+{tot('nM')}*{S('acvM')}*{g('grrM')}+{tot('nE')}*{S('acvE')}*{g('grrE')})"
            f"/({tot('nS')}*{S('acvS')}+{tot('nM')}*{S('acvM')}+{tot('nE')}*{S('acvE')}),0.9)"), PCT, "")
        self.scalar("ue", "csr", "CS cost as % of ARR at exit", lambda: f"=IFERROR({self._r('kc', 'N')}*4/{self._r('aT', 'N')},0)", PCT, "")
        self.scalar("ue", "x", "grr ÷ (1 + discount rate)", lambda: f"={S('grrb')}/(1+{g('disc')})", '0.0000', "")
        self.scalar("ue", "x0", "Option 0: grr0 ÷ (1 + discount rate)", f"={g('grr0')}/(1+{g('disc')})", '0.0000', "")
        self.header("ue", "Key outputs of this run")
        self.cur = "ue"
        aN = self._r("aT", "N")
        self.outrow("ue", "o_exit", "Exit ARR, Q4 2029", f"={aN}", USD)
        self.outrow("ue", "o_cust", "Customers, Q4 2029", f"={self._r('cT', 'N')}", N0)
        self.outrow("ue", "o_newc", "New customers 2027–29", f"={tot('new')}", N1)
        self.outrow("ue", "o_rev", "Recognized revenue 2027–29", f"={tot('rv')}", USD)
        self.outrow("ue", "o_acv", "Average first-year ACV", f"=IFERROR({tot('na')}/{tot('new')},0)", USD)
        self.outrow("ue", "o_cac", "CAC, fully loaded (S&M ÷ new customers)", f"=IFERROR({tot('sm')}/{tot('new')},0)", USD)
        self.outrow("ue", "o_mcac", "CAC, marginal (variable AE, programs, partner costs, subsidies)",
                    f"=IFERROR(({tot('kv')}+{tot('kp')}+{tot('kr')}+{tot('kn')}+{tot('ko')})/{tot('new')},0)", USD)
        self.outrow("ue", "o_pb", "CAC payback, 2027–29 (months)", f"=IFERROR({self.out['o_cac']}/({self.out['o_acv']}*{g('gm')}/12),0)", MO)
        self.outrow("ue", "o_pb29", "CAC payback, 2029 cohort (months)",
                    f"=IFERROR(({self._r('sm', 'R')}/{self._r('new', 'R')})/(({self._r('na', 'R')}/{self._r('new', 'R')})*{g('gm')}/12),0)", MO)
        self.outrow("ue", "o_cost", "Total incremental cost 2027–29", f"={tot('kt')}", USD)
        self.outrow("ue", "o_cum", "Cumulative program P&L 2027–29 (undiscounted, gross)", f"={self._r('cum', 'N')}", USD)
        self.outrow("ue", "o_peak", "Peak investment (lowest cumulative P&L)", f"=MIN({self.rng('cum')})", USD)
        self.outrow("ue", "o_emb", "Undiscounted post-2029 value of the 2029 ARR base (gross)",
                    f"={aN}*({g('gm')}-{S('csr')})*{S('grrb')}*(1-{S('grrb')}^{g('life')})/(1-{S('grrb')})", USD)
        self.outrow("ue", "o_pvemb", "PV of post-2029 value, after incrementality haircut",
                    f"={aN}*({g('gm')}-{S('csr')})*(1-{g('hair')})/(1+{g('disc')})^2.5*{S('x')}*(1-{S('x')}^{g('life')})/(1-{S('x')})", USD)
        self.outrow("ue", "o_npv", "NPV of the program (incremental, at hurdle rate)", f"=SUM({self.rng('pvi')})+{self.out['o_pvemb']}", USD)
        self.outrow("ue", "o_npv0", "NPV of Option 0 (same spend in the core motion)",
                    f"=SUM({self.rng('pv0')})+{self._r('a0', 'N')}*({g('gm')}-{g('csr0')})/(1+{g('disc')})^2.5*{S('x0')}*(1-{S('x0')}^{g('life')})/(1-{S('x0')})", USD)
        self.outrow("ue", "o_excess", "NPV advantage over Option 0 (program NPV − Option 0 NPV)", f"={self.out['o_npv']}-{self.out['o_npv0']}", USD)
        self.outrow("ue", "o_lost", "First-year ACV lost to capacity, 2027–29", f"={tot('lost')}", USD)
        self.outrow("ue", "o_gopps", "Gate metric: opportunities created in 2027", f"={self._r('opp', 'P')}", N1)
        self.outrow("ue", "o_gcust", "Gate metric: new customers in 2027", f"={self._r('new', 'P')}", N1)
        self.outrow("ue", "o_gpl", "Gate metric: partner introductions in 2027", f"={self._r('lp', 'P')}", N1)
        self.outrow("ue", "o_gacv", "Gate metric: average first-year ACV, 2027", f"=IFERROR({self._r('na', 'P')}/{self._r('new', 'P')},0)", USD)
        return self.out


# ============================================================ LIVE MODEL
for ws_, t, s_ in [(fun, "GTM Funnel — live model (selected segment, scenario and motion)", "Accounts → meetings, MQLs and partner introductions → opportunities (mix by band derived from accounts) → wins, with a closing-capacity constraint."),
                   (rev, "Revenue Model — customers, ARR and recognized revenue", "ARR = prior ARR × NRR^(1/4) + new customers × ACV. Recognized revenue = average ARR ÷ 4 (kept separate from ARR)."),
                   (ue, "Unit Economics — cost, P&L, NPV and Option 0", "Incremental cost only. NPV discounts quarterly at the hurdle rate and haircuts gross profit for wins that would have happened anyway. Option 0 redeploys the same non-CS spend into Clay's core motion.")]:
    title(ws_, t, s_)
    put(ws_, "A3", '="Live run: "&Assumptions!$E$5&" · "&Assumptions!$E$4&" · "&Assumptions!$E$6', bold=True, border=False)
    ws_.column_dimensions["A"].width = 54
    ws_.column_dimensions["B"].width = 14
    for c in QCOLS:
        ws_.column_dimensions[c].width = 11
    ws_.column_dimensions["O"].width = 2
    for c in "PQRS":
        ws_.column_dimensions[c].width = 13
    ws_.freeze_panes = "C5"


def firm_live(seg, w, b):
    return f"IF({SEL_G}=1,{FIRM('dist', w, b)},{FIRM('ps', w, b)})"


def acv_live(seg, b, g):
    return f"IF({SEL_G}=1,{acv_formula('dist', b, g)},{acv_formula('ps', b, g)})"


LIVE = Engine({"fun": [fun, 5], "rev": [rev, 5], "ue": [ue, 5]}, lambda n: f"Assumptions!$J${A[n]}", lambda n: f"Assumptions!$J${M[n]}",
              lambda i: f"Assumptions!${L(4 + i)}${HIRE['live']}", "Live model", "live", firm_live, acv_live).build()

# ============================================================ ENGINE RUNS
title(eng, "Engine Runs — full copies of the model for every scenario, option and sensitivity case",
      "Generated by the same code as the live model. Each block lists its overrides; the Python engine (appendix/clay_engine.py) reproduces every scenario block (Monte Carlo › reconciliation).")
eng.column_dimensions["A"].width = 54
eng.column_dimensions["B"].width = 14
for c in QCOLS:
    eng.column_dimensions[c].width = 11
for c in "PQRS":
    eng.column_dimensions[c].width = 13
eng_row = [4]
RUNSPEC = {}


def run(key, label, seg, scen, motion="lean", overrides=None, hire=None):
    overrides = overrides or {}
    col = SCOL[(seg, scen)]
    mcol = "DEF"[MOT.index(motion)]
    g = lambda n: str(overrides[n]) if n in overrides else f"Assumptions!${col}${A[n]}"
    mg = lambda n: f"Assumptions!${mcol}${M[n]}"
    hr = HIRE[hire or motion]
    start = eng_row[0]
    put(eng, f"A{start}", "RUN: " + label, bold=True, fill=YFILL if overrides or hire else KFILL)
    eng.merge_cells(f"A{start}:N{start}")
    eng_row[0] += 1
    if overrides or hire:
        put(eng, f"A{eng_row[0]}", "Overrides: " + "; ".join([f"{k} = {v}" for k, v in overrides.items()] + ([f"AE hiring = {hire}"] if hire else [])), italic=True, border=False)
        eng.merge_cells(f"A{eng_row[0]}:N{eng_row[0]}")
        eng_row[0] += 1
    e = Engine({"fun": [eng, eng_row[0]], "rev": [eng, 0], "ue": [eng, 0]}, g, mg, lambda i: f"Assumptions!${L(4 + i)}${hr}", label, seg, FIRM, acv_formula)
    e.stages["rev"] = e.stages["fun"]
    e.stages["ue"] = e.stages["fun"]
    out = e.build()
    eng_row[0] = e.stages["fun"][1] + 2
    out["_reg"] = e.reg
    RUNSPEC[key] = dict(seg=seg, scen=scen, motion=motion, overrides=overrides, hire=hire, label=label)
    return out


RUNS = {}
for seg in SEGS:
    for s in range(3):
        RUNS[(seg, "lean", s)] = run((seg, "lean", s), f"{SEGLAB[seg]} — lean pilot — {SCN[s]}", seg, s, "lean")
    RUNS[(seg, "gated", 0)] = run((seg, "gated", 0), f"{SEGLAB[seg]} — lean pilot — Downside, stopped at Gate 1 (Q4 2027)", seg, 0, "lean", {"stop": 4})
    RUNS[(seg, "scaled", 2)] = run((seg, "scaled", 2), f"{SEGLAB[seg]} — lean pilot — Upside, 2 AEs hired after Gate 1", seg, 2, "lean", hire="scale")
    for s in range(3):
        RUNS[(seg, "sales", s)] = run((seg, "sales", s), f"{SEGLAB[seg]} — sales-led vertical team — {SCN[s]}", seg, s, "sales")
    for s in range(3):
        RUNS[(seg, "play", s)] = run((seg, "play", s), f"{SEGLAB[seg]} — solution play — {SCN[s]}", seg, s, "play")

TORNADO = [
    ("Automated outbound coverage", ["cov"], None), ("Meeting rate", ["meet"], None), ("Outbound fatigue", ["fatigue"], None),
    ("Inbound MQL rate", ["inb"], None), ("Partner network (start, adds, cap, intros)", ["p0", "padd", "pcap", "lpp"], None),
    ("Funnel conversion (meeting / intro / MQL → opp)", ["m2o", "p2o", "i2o"], None), ("Win rate", ["win"], None),
    ("Enterprise propensity (mix)", ["propE"], None), ("ICP tilt to Mid-Market / Enterprise (1× ↔ 3×)", ["tiltM", "tiltE"], (1, 3)),
    ("Usage intensity (ACV)", ["usage"], None), ("Gross revenue retention", ["grrS", "grrM", "grrE"], None),
    ("Net revenue retention", ["nrrS", "nrrM", "nrrE"], None), ("Gross margin", ["gm"], None), ("Wave timing", ["w2", "w3"], None),
    ("Partner fee and onboarding subsidy", ["ref", "onb"], None), ("Self-serve share of SMB", ["ss"], None),
    ("Incrementality haircut", ["hair"], None), ("Discount rate (15% ↔ 8%)", ["disc"], (0.15, 0.08)),
]
TRUNS = []
for nm, keys, lit in TORNADO:
    if lit:
        lo = run(("t", nm, 0), f"Sensitivity — {nm}: low", "dist", 1, "lean", {k: lit[0] for k in keys})
        hi = run(("t", nm, 1), f"Sensitivity — {nm}: high", "dist", 1, "lean", {k: lit[1] for k in keys})
    else:
        lo = run(("t", nm, 0), f"Sensitivity — {nm}: Downside value", "dist", 1, "lean", {k: f"Assumptions!$D${A[k]}" for k in keys})
        hi = run(("t", nm, 1), f"Sensitivity — {nm}: Upside value", "dist", 1, "lean", {k: f"Assumptions!$F${A[k]}" for k in keys})
    TRUNS.append((nm, keys, lit, lo, hi))
GW = [0.15, 0.21, 0.30, 0.40]
GU = [0.75, 1.0, 1.5, 2.5]
GRUNS = {(w, u): run(("g", w, u), f"Grid — win {w:.0%} × usage {u}x", "dist", 1, "lean", {"win": w, "usage": u}) for w in GW for u in GU}
AEB = [0, 150, 300, 600]
ARUNS = {a: run(("ae", a), f"Sales-led team — Base — AE accounts per quarter = {a}", "dist", 1, "sales", {"aeacc": a}) for a in AEB}

# ============================================================ SCENARIOS
title(scn, "Scenarios — options compared on one engine (deterministic Downside / Base / Upside)",
      "Every figure links to a full run on 'Engine Runs'. NPV at the hurdle rate, after the incrementality haircut. Illustrative scenario outputs — not Clay forecasts.")
COLS = [(seg, mo, s) for seg in SEGS for mo in MOT for s in range(3)]
hdr(scn, 4, ["Output", "Unit"] + [f"{'Dist' if c[0] == 'dist' else 'PS'} · {MOTLAB[c[1]].split(' (')[0]} · {SCN[c[2]]}" for c in COLS], widths=[48, 8] + [13] * len(COLS), height=58)
OUTS = [("o_npv", "NPV of the program at hurdle rate", "$", USD), ("o_npv0", "NPV of Option 0 (same spend in core)", "$", USD), ("o_excess", "NPV advantage over Option 0", "$", USD),
        ("o_exit", "Exit ARR, Q4 2029", "$", USD), ("o_cust", "Customers, Q4 2029", "#", N0), ("o_newc", "New customers 2027–29", "#", N0),
        ("o_acv", "Average first-year ACV", "$", USD), ("o_cac", "CAC, fully loaded", "$", USD), ("o_mcac", "CAC, marginal", "$", USD),
        ("o_pb29", "CAC payback, 2029 cohort", "months", MO), ("o_cost", "Total incremental cost 2027–29", "$", USD),
        ("o_cum", "Cumulative P&L 2027–29 (undiscounted)", "$", USD), ("o_peak", "Peak investment", "$", USD), ("o_lost", "ACV lost to capacity", "$", USD),
        ("o_gopps", "Gate metric: 2027 opportunities", "#", N1), ("o_gcust", "Gate metric: 2027 new customers", "#", N1),
        ("o_gpl", "Gate metric: 2027 partner introductions", "#", N1), ("o_gacv", "Gate metric: 2027 average ACV", "$", USD)]
SROW = {}
for i, (k, lab, unit, fmt) in enumerate(OUTS):
    rr = 5 + i
    SROW[k] = rr
    put(scn, f"A{rr}", lab, bold=k in ("o_npv", "o_excess", "o_exit"))
    put(scn, f"B{rr}", unit, center=True)
    for j, c in enumerate(COLS):
        put(scn, f"{L(3 + j)}{rr}", f"={RUNS[c][k]}", fmt=fmt, font=GREEN, center=True)
LASTC = L(2 + len(COLS))
for k in ("o_npv", "o_excess"):
    scn.conditional_formatting.add(f"C{SROW[k]}:{LASTC}{SROW[k]}", CellIsRule(operator="lessThan", formula=["0"], fill=RFILL))
    scn.conditional_formatting.add(f"C{SROW[k]}:{LASTC}{SROW[k]}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=KFILL))
rr = 5 + len(OUTS) + 1
put(scn, f"A{rr}", "Gated variants (lean pilot)", bold=True, fill=SFILL)
rr += 1
hdr(scn, rr, ["Output", "Unit", "Dist · Downside stopped at Gate 1", "Dist · Upside + 2 AEs after Gate 1", "PS · Downside stopped at Gate 1", "PS · Upside + 2 AEs after Gate 1"], height=46)
rr += 1
for k, lab, unit, fmt in OUTS[:14]:
    put(scn, f"A{rr}", lab)
    put(scn, f"B{rr}", unit, center=True)
    for j, c in enumerate([("dist", "gated", 0), ("dist", "scaled", 2), ("ps", "gated", 0), ("ps", "scaled", 2)]):
        put(scn, f"{L(3 + j)}{rr}", f"={RUNS[c][k]}", fmt=fmt, font=GREEN, center=True)
    rr += 1
rr += 1
put(scn, f"A{rr}", "Reconciliation: live model vs matching engine run (should be 0)", bold=True, fill=SFILL)
rr += 1
colidx = f"(({SEL_G}-1)*9+({SEL_M}-1)*3+{SEL_S})"
put(scn, f"A{rr}", "NPV difference")
put(scn, f"C{rr}", f"={LIVE['o_npv']}-INDEX($C${SROW['o_npv']}:${LASTC}${SROW['o_npv']},1,{colidx})", fmt=USD, center=True)
RECON1 = rr
rr += 1
put(scn, f"A{rr}", "Exit ARR difference")
put(scn, f"C{rr}", f"={LIVE['o_exit']}-INDEX($C${SROW['o_exit']}:${LASTC}${SROW['o_exit']},1,{colidx})", fmt=USD, center=True)
RECON2 = rr
rr += 2
put(scn, f"A{rr}", "ARR trajectory by quarter — distribution (for charts)", bold=True, fill=SFILL)
for i, q in enumerate(QLAB):
    put(scn, f"{L(3 + i)}{rr}", q, bold=True, center=True, fill=SFILL)
TRJ = rr
traj = [("Lean pilot — Downside", ("dist", "lean", 0)), ("Lean pilot — Base", ("dist", "lean", 1)), ("Lean pilot — Upside", ("dist", "lean", 2)),
        ("Sales-led team — Base", ("dist", "sales", 1)), ("Solution play — Base", ("dist", "play", 1))]
for k, (lab, key) in enumerate(traj):
    r2 = rr + 1 + k
    put(scn, f"A{r2}", lab)
    _, regr = RUNS[key]["_reg"]["aT"]
    for i in range(12):
        put(scn, f"{L(3 + i)}{r2}", f"='Engine Runs'!{QCOLS[i]}{regr}", fmt=USD, font=GREEN)
scn.freeze_panes = "C5"

# ============================================================ SENSITIVITY
title(sen, "Sensitivity — distribution lean pilot, Base: what moves NPV, and what would have to be true",
      "Block A: one input at a time between its Downside and Upside values (ICP tilt and discount rate use the stated values). Block B: win rate × usage grid. Block C: can a sales-led team beat the pilot if its AEs prospect harder?")
hdr(sen, 4, ["Input swung", "Low value", "High value", "NPV @ low", "NPV @ high", "Swing (NPV)", "Advantage vs Option 0 @ low", "@ high", "Exit ARR @ low", "@ high"],
    widths=[44, 12, 12, 15, 15, 15, 17, 15, 14, 14])
for i, (nm, keys, lit, lo, hi) in enumerate(TRUNS):
    rr = 5 + i
    put(sen, f"A{rr}", nm, bold=True)
    k0 = keys[0]
    unit = next(x[7] for x in ce.INPUTS if x[0] == k0)
    fmt = PCT if unit == "%" else ('0.00' if unit in ("x", "x/yr") else N0)
    if lit:
        put(sen, f"B{rr}", lit[0], fmt=fmt, center=True)
        put(sen, f"C{rr}", lit[1], fmt=fmt, center=True)
    else:
        put(sen, f"B{rr}", f"=Assumptions!$D${A[k0]}", fmt=fmt, font=GREEN, center=True)
        put(sen, f"C{rr}", f"=Assumptions!$F${A[k0]}", fmt=fmt, font=GREEN, center=True)
    put(sen, f"D{rr}", f"={lo['o_npv']}", fmt=USD, center=True)
    put(sen, f"E{rr}", f"={hi['o_npv']}", fmt=USD, center=True)
    put(sen, f"F{rr}", f"=ABS(E{rr}-D{rr})", fmt=USD, center=True, bold=True)
    put(sen, f"G{rr}", f"={lo['o_excess']}", fmt=USD, center=True)
    put(sen, f"H{rr}", f"={hi['o_excess']}", fmt=USD, center=True)
    put(sen, f"I{rr}", f"={lo['o_exit']}", fmt=USD, center=True)
    put(sen, f"J{rr}", f"={hi['o_exit']}", fmt=USD, center=True)
TL = 5 + len(TRUNS) - 1
sen.conditional_formatting.add(f"F5:F{TL}", DataBarRule(start_type="num", start_value=0, end_type="max", color="4472C4"))
put(sen, f"A{TL + 1}", "Base case", bold=True, fill=SFILL)
put(sen, f"D{TL + 1}", f"={RUNS[('dist', 'lean', 1)]['o_npv']}", fmt=USD, center=True, fill=SFILL)
put(sen, f"G{TL + 1}", f"={RUNS[('dist', 'lean', 1)]['o_excess']}", fmt=USD, center=True, fill=SFILL)
chart = BarChart()
chart.type = "bar"
chart.title = "Swing in NPV (one input, low ↔ high)"
chart.add_data(Reference(sen, min_col=6, min_row=4, max_row=TL), titles_from_data=True)
chart.set_categories(Reference(sen, min_col=1, min_row=5, max_row=TL))
chart.height, chart.width, chart.legend = 10, 18, None
chart.x_axis.delete = False
chart.y_axis.delete = False
sen.add_chart(chart, "L4")
gb = TL + 4
put(sen, f"A{gb}", "Block B — NPV of the distribution lean pilot: win rate (rows) × usage intensity (columns), all else Base", bold=True, border=False)
put(sen, f"A{gb + 1}", "Win rate ↓ / usage →", bold=True, fill=SFILL)
for j, u in enumerate(GU):
    put(sen, f"{L(2 + j)}{gb + 1}", u, font=BLUE, fmt='0.00"x"', center=True, fill=SFILL)
for i, w in enumerate(GW):
    put(sen, f"A{gb + 2 + i}", w, font=BLUE, fmt=PCT, center=True, fill=SFILL)
    for j, u in enumerate(GU):
        put(sen, f"{L(2 + j)}{gb + 2 + i}", f"={GRUNS[(w, u)]['o_npv']}", fmt=USD, center=True)
GR = (gb + 2, gb + 1 + len(GW))
gb2 = GR[1] + 2
put(sen, f"A{gb2}", "Block B2 — Same grid: NPV advantage over Option 0", bold=True, border=False)
put(sen, f"A{gb2 + 1}", "Win rate ↓ / usage →", bold=True, fill=SFILL)
for j, u in enumerate(GU):
    put(sen, f"{L(2 + j)}{gb2 + 1}", u, font=BLUE, fmt='0.00"x"', center=True, fill=SFILL)
for i, w in enumerate(GW):
    put(sen, f"A{gb2 + 2 + i}", w, font=BLUE, fmt=PCT, center=True, fill=SFILL)
    for j, u in enumerate(GU):
        put(sen, f"{L(2 + j)}{gb2 + 2 + i}", f"={GRUNS[(w, u)]['o_excess']}", fmt=USD, center=True)
GR2 = (gb2 + 2, gb2 + 1 + len(GW))
for a_, b_ in (GR, GR2):
    sen.conditional_formatting.add(f"B{a_}:E{b_}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=KFILL))
    sen.conditional_formatting.add(f"B{a_}:E{b_}", CellIsRule(operator="lessThan", formula=["0"], fill=RFILL))
cb = GR2[1] + 3
put(sen, f"A{cb}", "Block C — Sales-led team (Base) with more AE prospecting vs the lean pilot (Base)", bold=True, border=False)
hdr(sen, cb + 1, ["Accounts each ramped AE works per quarter", "New customers", "Exit ARR", "NPV", "Lean pilot NPV", "Sales-led − lean", "", "", "", ""])
for i, a in enumerate(AEB):
    r3 = cb + 2 + i
    put(sen, f"A{r3}", a, font=BLUE, fmt=N0, center=True)
    put(sen, f"B{r3}", f"={ARUNS[a]['o_newc']}", fmt=N1, center=True)
    put(sen, f"C{r3}", f"={ARUNS[a]['o_exit']}", fmt=USD, center=True)
    put(sen, f"D{r3}", f"={ARUNS[a]['o_npv']}", fmt=USD, center=True)
    put(sen, f"E{r3}", f"={RUNS[('dist', 'lean', 1)]['o_npv']}", fmt=USD, center=True)
    put(sen, f"F{r3}", f"=D{r3}-E{r3}", fmt=USD, center=True, bold=True)
CB = (cb + 2, cb + 1 + len(AEB))
note(sen, f"A{CB[1] + 2}:J{CB[1] + 2}", "Reading (v2): the sales-led team now generates its own pipeline, so the comparison is no longer true by construction (review C1). It adds customers and ARR, but in a ~5,000-account market the extra ACV does not pay for six AEs. "
     "Block B: at Base volumes, NPV turns positive only with win rates well above the 21% benchmark and usage far above the base build.", 46)

# ============================================================ MONTE CARLO (Python)
title(mc, "Monte Carlo and gate policy — computed by appendix/clay_engine.py (a Python mirror of this workbook)",
      "Each draw samples every non-fixed input from triangular(Downside, Base, Upside); inputs in the same group (demand, value, cost, timing) are correlated through a Gaussian copula (ρ = 0.5). Values are static; rerun build_financial_model.py to refresh.")
N_MC, SEED = 5000, 42
sims = {}
for seg in SEGS:
    sims[(seg, "lean")] = ce.simulate(seg, N_MC, SEED, TH[seg], motion="lean")
    sims[(seg, "sales")] = ce.simulate(seg, N_MC, SEED, None, motion="sales")
    sims[(seg, "play")] = ce.simulate(seg, N_MC, SEED, None, motion="play")
hdr(mc, 4, ["Option / policy", "Mean NPV", "P10", "P50", "P90", "P(NPV > 0)", "Mean advantage vs Option 0", "P(beats Option 0)", "Stopped at Gate 1", "Scaled at Gate 1"],
    widths=[56, 15, 15, 15, 15, 12, 18, 14, 13, 13])
MCROWS = {}
rr = 5
for seg in SEGS:
    for key, lab, field, xfield in [("lean", "Lean pilot — run all 12 quarters (ungated)", "ungated", "ungated_x"),
                                    ("lean", "Lean pilot — staged: Gate 1 with observation noise", "staged", "staged_x"),
                                    ("sales", "Sales-led vertical team", "ungated", "ungated_x"),
                                    ("play", "Solution play (no vertical team)", "ungated", "ungated_x")]:
        res = sims[(seg, key)]
        s = ce.summarize([x[field] for x in res])
        sx = ce.summarize([x[xfield] for x in res])
        put(mc, f"A{rr}", f"{SEGLAB[seg]} — {lab}", bold=(field == "staged" or key == "play"))
        for col, v, fmt in [("B", s["mean"], USD), ("C", s["p10"], USD), ("D", s["p50"], USD), ("E", s["p90"], USD), ("F", s["ppos"], PCT),
                            ("G", sx["mean"], USD), ("H", sx["ppos"], PCT)]:
            put(mc, f"{col}{rr}", round(v, 4) if fmt == PCT else round(v), fmt=fmt, center=True)
        if field == "staged":
            n = len(res)
            put(mc, f"I{rr}", round(sum(1 for x in res if x["decision"] == "stop") / n, 4), fmt=PCT, center=True)
            put(mc, f"J{rr}", round(sum(1 for x in res if x["decision"] == "scale") / n, 4), fmt=PCT, center=True)
        MCROWS[(seg, key, field)] = rr
        rr += 1
rr += 1
put(mc, f"A{rr}", "Robustness of P(NPV > 0), distribution lean pilot (ungated): correlation and range width (2,000 draws each)", bold=True, fill=SFILL)
rr += 1
hdr(mc, rr, ["Specification", "Within-group ρ", "Cross-group γ", "Range × (around Base)", "Mean NPV", "P90 NPV", "P(NPV > 0)", "", "", ""])
rr += 1
PCTN = {"f1", "f2", "exist", "cov", "meet", "fatigue", "inb", "m2o", "p2o", "i2o", "win", "ss", "gm", "ref", "onb", "hair", "grrS", "grrM", "grrE", "prod"}


def sample_ext(seg, rng, rho, gam, k):
    off = 0 if seg == "dist" else 3
    G = rng.gauss(0, 1)
    groups = {g_: math.sqrt(gam) * G + math.sqrt(1 - gam) * rng.gauss(0, 1) for g_ in ("demand", "value", "cost", "timing")}
    p = {}
    for x in ce.INPUTS:
        name, grp = x[0], x[9]
        d, b, u = x[1 + off], x[2 + off], x[3 + off]
        if grp is None:
            p[name] = b
            continue
        z = rng.gauss(0, 1) if grp == "indep" else math.sqrt(rho) * groups[grp] + math.sqrt(1 - rho) * rng.gauss(0, 1)
        q = ce._phi(z)
        if grp == "timing":
            p[name] = d if q < 1 / 3 else (b if q < 2 / 3 else u)
            continue
        lo_, hi_ = (d, u) if d <= u else (u, d)
        lo2, hi2 = b - k * (b - lo_), b + k * (hi_ - b)
        if name in PCTN:
            lo2, hi2 = max(0.001, lo2), min(0.99, hi2)
        lo2 = max(lo2, 0.0001)
        p[name] = ce._triq(q if d <= u else 1 - q, lo2, b, hi2)
    return p


ROB = []
for lab, rho, gam, k in [("Base specification", .5, 0, 1.0), ("Strong correlation", .8, .5, 1.0), ("Near-comonotonic (≈ joint scenarios)", .9, .9, 1.0),
                         ("Ranges 1.5× wider", .5, 0, 1.5), ("Ranges 1.5× wider + strong correlation", .8, .5, 1.5), ("Ranges 2× wider + strong correlation", .8, .5, 2.0)]:
    rng_ = random.Random(SEED)
    v = [ce.run("dist", sample_ext("dist", rng_, rho, gam, k))["npv"] for _ in range(2000)]
    s = ce.summarize(v)
    ROB.append((lab, s))
    put(mc, f"A{rr}", lab)
    for col, val, fmt in [("B", rho, '0.0'), ("C", gam, '0.0'), ("D", k, '0.0"x"'), ("E", round(s["mean"]), USD), ("F", round(s["p90"]), USD), ("G", round(s["ppos"], 4), PCT)]:
        put(mc, f"{col}{rr}", val, fmt=fmt, center=True)
    rr += 1
rr += 1
put(mc, f"A{rr}", "What would have to be true — single-input break-evens (Python bisection, all else Base)", bold=True, fill=SFILL)
rr += 1
hdr(mc, rr, ["Break-even", "Distribution lean", "Distribution play", "PS lean", "PS play", "Base value", "", "", "", ""])
rr += 1


def bisect(seg, key, target, lo, hi, motion, scen=1):
    def f(v):
        p = ce.params(seg, scen)
        p[key] = v
        return ce.run(seg, p, motion=motion)[target]
    flo, fhi = f(lo), f(hi)
    if flo * fhi > 0:
        return None
    for _ in range(60):
        mid = (lo + hi) / 2
        if (f(mid) > 0) == (fhi > 0):
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


BE = {}
for lab, key, target, lo, hi, scen, fmt, base in [("Usage intensity needed for NPV = 0 (Base)", "usage", "npv", .5, 12, 1, '0.00"x"', "1.00x"),
                                                  ("Win rate needed for NPV = 0 (Base; capped by borrowed AE capacity)", "win", "npv", .05, .95, 1, PCT, "21% / 18%"),
                                                  ("Option 0 efficiency at which the program ties Option 0 (Upside)", "eff0", "excess", 0.0, 2.0, 2, '0.00', "0.60"),
                                                  ("Option 0 efficiency at which the program ties Option 0 (Base)", "eff0", "excess", -1.0, 2.0, 1, '0.00', "0.60")]:
    put(mc, f"A{rr}", lab)
    for j, (seg, mo) in enumerate([("dist", "lean"), ("dist", "play"), ("ps", "lean"), ("ps", "play")]):
        v = bisect(seg, key, target, lo, hi, mo, scen)
        BE[(key, scen, seg, mo)] = v
        put(mc, f"{L(2 + j)}{rr}", round(v, 4) if v is not None else "not reachable", fmt=fmt, center=True)
    put(mc, f"F{rr}", base, center=True)
    rr += 1
rr += 1
put(mc, f"A{rr}", "Reconciliation — Python engine vs Excel engine, scenario runs (difference should be 0)", bold=True, fill=SFILL)
rr += 1
hdr(mc, rr, ["Run", "Python NPV", "Excel NPV", "Difference", "Python exit ARR", "Excel exit ARR", "Difference", "", "", ""])
rr += 1
REC0 = rr
for key, out in RUNS.items():
    spec = RUNSPEC[key]
    p = ce.params(spec["seg"], spec["scen"])
    for k, v in spec["overrides"].items():
        p[k] = v
    py = ce.run(spec["seg"], p, motion=spec["motion"], hires=ce.HIRES[spec["hire"]] if spec["hire"] else None)
    put(mc, f"A{rr}", spec["label"])
    put(mc, f"B{rr}", round(py["npv"], 2), fmt=USD, center=True)
    put(mc, f"C{rr}", f"={out['o_npv']}", fmt=USD, center=True, font=GREEN)
    put(mc, f"D{rr}", f"=ROUND(C{rr}-B{rr},0)", fmt=USD, center=True)
    put(mc, f"E{rr}", round(py["exit_arr"], 2), fmt=USD, center=True)
    put(mc, f"F{rr}", f"={out['o_exit']}", fmt=USD, center=True, font=GREEN)
    put(mc, f"G{rr}", f"=ROUND(F{rr}-E{rr},0)", fmt=USD, center=True)
    rr += 1
REC1 = rr - 1
put(mc, f"A{rr}", "Maximum absolute difference", bold=True)
put(mc, f"D{rr}", f"=MAX(MAX(D{REC0}:D{REC1}),-MIN(D{REC0}:D{REC1}))", fmt=USD, center=True, bold=True, fill=KFILL)
put(mc, f"G{rr}", f"=MAX(MAX(G{REC0}:G{REC1}),-MIN(G{REC0}:G{REC1}))", fmt=USD, center=True, bold=True, fill=KFILL)
RECMAX = rr
rr += 2
put(mc, f"A{rr}", "NPV distribution — distribution segment (share of draws per $250K bin; end bins are open)", bold=True, fill=SFILL)
rr += 1
hdr(mc, rr, ["NPV bin (lower edge)", "Lean pilot, staged", "Sales-led team", "Solution play", "", "", "", "", "", ""])
rr += 1
H0 = rr
edges = [-6.5e6 + 0.25e6 * k for k in range(0, 35)]
series = {"lean": [x["staged"] for x in sims[("dist", "lean")]], "sales": [x["ungated"] for x in sims[("dist", "sales")]], "play": [x["ungated"] for x in sims[("dist", "play")]]}
for idx, e in enumerate(edges):
    put(mc, f"A{rr}", e, fmt='$#,##0,"K"', center=True)
    for j, k in enumerate(["lean", "sales", "play"]):
        v = series[k]
        if idx == 0:
            cnt = sum(1 for x in v if x < e + 0.25e6)
        elif idx == len(edges) - 1:
            cnt = sum(1 for x in v if x >= e)
        else:
            cnt = sum(1 for x in v if e <= x < e + 0.25e6)
        put(mc, f"{L(2 + j)}{rr}", round(cnt / len(v), 4), fmt=PCT, center=True)
    rr += 1
H1 = rr - 1
hc = BarChart()
hc.type = "col"
hc.grouping = "clustered"
hc.title = "Simulated NPV, distribution segment (5,000 draws)"
hc.add_data(Reference(mc, min_col=2, max_col=4, min_row=H0 - 1, max_row=H1), titles_from_data=True)
hc.set_categories(Reference(mc, min_col=1, min_row=H0, max_row=H1))
hc.height, hc.width = 9, 22
hc.x_axis.delete = False
hc.y_axis.delete = False
hc.y_axis.numFmt = '0%'
mc.add_chart(hc, "L4")
note(mc, f"A{rr + 1}:J{rr + 1}", "Gate thresholds (Assumptions › Gate 1 rule) sit at the midpoints between the deterministic Downside/Base and Base/Upside 2027 paths. Observed 2027 metrics are drawn with Poisson/binomial noise, so the gate can misclassify (review C3). "
     "Staged-policy NPV includes 2027 spend even when the program is stopped.", 46)

# ============================================================ EXECUTIVE OUTPUT
title(exo, "Executive Output — is any vertical program worth funding in 2027–2029? (v2, after partner review)",
      "Scenario and simulation outputs are illustrative estimates built from public data and labeled assumptions. They are not Clay forecasts and use no Clay internal data.")
put(exo, "A3", '="Live model: "&Assumptions!$E$5&" · "&Assumptions!$E$4&" · "&Assumptions!$E$6&"  (selectors on Assumptions D4:D6)"', bold=True, border=False)
hdr(exo, 5, ["Option (Base case, deterministic)", "NPV at 12%", "NPV advantage vs Option 0", "Exit ARR Q4 2029", "Average first-year ACV", "CAC payback, 2029 cohort", "Cumulative P&L 2027–29", "Peak investment"],
    widths=[50, 15, 18, 15, 15, 15, 17, 15])
rr = 6
for seg in SEGS:
    for mo in MOT:
        key = (seg, mo, 1)
        put(exo, f"A{rr}", f"{SEGLAB[seg]} — {MOTLAB[mo]}", bold=(mo == "play"))
        for col, k, fmt in [("B", "o_npv", USD), ("C", "o_excess", USD), ("D", "o_exit", USD), ("E", "o_acv", USD), ("F", "o_pb29", MO), ("G", "o_cum", USD), ("H", "o_peak", USD)]:
            put(exo, f"{col}{rr}", f"=Scenarios!{L(3 + COLS.index(key))}{SROW[k]}", fmt=fmt, font=GREEN, center=True)
        rr += 1
exo.conditional_formatting.add(f"B6:C{rr - 1}", CellIsRule(operator="lessThan", formula=["0"], fill=RFILL))
exo.conditional_formatting.add(f"B6:C{rr - 1}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=KFILL))
rr += 1
hdr(exo, rr, ["Simulation (5,000 draws; correlated inputs)", "Mean NPV", "P10", "P90", "P(NPV > 0)", "P(beats Option 0)", "Stopped at Gate 1", "Scaled at Gate 1"])
rr += 1
for seg in SEGS:
    for key, field, lab in [("lean", "staged", "Lean pilot, staged"), ("sales", "ungated", "Sales-led team"), ("play", "ungated", "Solution play")]:
        mr = MCROWS[(seg, key, field)]
        put(exo, f"A{rr}", f"{SEGLAB[seg]} — {lab}")
        for col, src_col, fmt in [("B", "B", USD), ("C", "C", USD), ("D", "E", USD), ("E", "F", PCT), ("F", "H", PCT), ("G", "I", PCT), ("H", "J", PCT)]:
            put(exo, f"{col}{rr}", f"='Monte Carlo'!{src_col}{mr}", fmt=fmt, font=GREEN, center=True)
        rr += 1
rr += 1
put(exo, f"A{rr}", "Model checks", bold=True, fill=SFILL)
rr += 1
put(exo, f"A{rr}", "Live model reconciles to engine runs (0 = pass)")
put(exo, f"B{rr}", f"=ABS(Scenarios!C{RECON1})+ABS(Scenarios!C{RECON2})", fmt=USD, font=GREEN, center=True)
CHK1 = rr
rr += 1
put(exo, f"A{rr}", "Python engine reconciles to Excel engine (max abs. NPV difference, $)")
put(exo, f"B{rr}", f"='Monte Carlo'!D{RECMAX}", fmt=USD, font=GREEN, center=True)
rr += 2
put(exo, f"A{rr}", "What the model shows (v2)", bold=True, fill=SFILL)
rr += 1
Bres = {(seg, mo, s): ce.run(seg, ce.params(seg, s), motion=mo) for seg in SEGS for mo in MOT for s in range(3)}
st = ce.summarize([x["staged"] for x in sims[("dist", "lean")]])
plv = ce.summarize([x["ungated"] for x in sims[("dist", "play")]])
fm = lambda v: ("−" if v < 0 else "+") + f"${abs(v) / 1e6:.2f}M"
b = Bres
msgs = [
    f"1. No option clears a 12% hurdle in Base. Distribution lean pilot: NPV {fm(b[('dist', 'lean', 1)]['npv'])}, exit ARR ${b[('dist', 'lean', 1)]['exit_arr'] / 1e6:.2f}M, average first-year ACV ${b[('dist', 'lean', 1)]['acv'] / 1e3:.1f}K once the win mix is derived from the account base (v1 assumed $13.8K).",
    f"2. Upside is NPV-positive ({fm(b[('dist', 'lean', 2)]['npv'])}) only when favorable values occur together. With correlated uncertainty, the staged pilot's P(NPV > 0) is {st['ppos']:.0%} and its P90 is {fm(st['p90'])}; even near-comonotonic inputs or 2× wider ranges keep P(NPV > 0) at or below {max(s_['ppos'] for _, s_ in ROB):.0%}.",
    f"3. Against Option 0 (the same spend in Clay's core motion at a discounted benchmark payback) every staffed option loses, in Base and in Upside. The bar a vertical must clear is the core's return, not zero.",
    f"4. A sales-led team now builds its own pipeline and adds ARR (exit ${b[('dist', 'sales', 1)]['exit_arr'] / 1e6:.2f}M vs ${b[('dist', 'lean', 1)]['exit_arr'] / 1e6:.2f}M lean), but NPV falls to {fm(b[('dist', 'sales', 1)]['npv'])}: a ~5,000-account market cannot carry six AEs.",
    f"5. Professional services is no better on the same engine: list-price ACV sits at the Growth floor for SMB and Mid-Market, and more of its wins would have come anyway (Base NPV {fm(b[('ps', 'lean', 1)]['npv'])}).",
    f"6. The cheapest test — a solution play with templates and data providers but no vertical team — loses least (distribution Base NPV {fm(b[('dist', 'play', 1)]['npv'])}; simulated P(NPV > 0) {plv['ppos']:.0%}) and is the only option that beats Option 0 in Upside. Recommendation: validate before staffing anything.",
]
for t in msgs:
    put(exo, f"A{rr}", t)
    exo.merge_cells(f"A{rr}:H{rr}")
    exo.row_dimensions[rr].height = 34
    rr += 1
ch1 = LineChart()
ch1.title = "Distribution: ARR by quarter, 2027–2029"
ch1.y_axis.numFmt = '$#,##0,"K"'
for k in range(len(traj)):
    ch1.series.append(Series(Reference(scn, min_col=3, max_col=14, min_row=TRJ + 1 + k), title=traj[k][0]))
ch1.set_categories(Reference(scn, min_col=3, max_col=14, min_row=TRJ))
ch1.height, ch1.width = 9, 17
ch1.x_axis.delete = False
ch1.y_axis.delete = False
exo.add_chart(ch1, f"A{rr + 1}")
hc2 = BarChart()
hc2.type = "col"
hc2.title = "Simulated NPV, distribution (share of draws)"
hc2.add_data(Reference(mc, min_col=2, max_col=4, min_row=H0 - 1, max_row=H1), titles_from_data=True)
hc2.set_categories(Reference(mc, min_col=1, min_row=H0, max_row=H1))
hc2.height, hc2.width = 9, 17
hc2.x_axis.delete = False
hc2.y_axis.delete = False
hc2.y_axis.numFmt = '0%'
exo.add_chart(hc2, f"E{rr + 1}")

# ============================================================ CHANGE LOG
title(chg, "Change Log — v1 → v2 (response to the partner review of 2026-10-19)", "Review: review/Partner_Review_2026-10-19.md. v1 workbook and build script archived in appendix/archive/.")
hdr(chg, 4, ["Review issue", "v1 behaviour", "v2 change", "Effect (distribution, Base)"], widths=[16, 55, 70, 48])
v1 = {"acv": 13801, "exit": 861204, "tv": -1041530}
CHANGES = [
    ("C2 win mix", "Fixed 50/38/12 SMB/MM/Ent mix: enterprise accounts converted 5.4× SMB", "Mix derived each quarter from active accounts × ICP tilt × band propensity; enterprise cycle 2 quarters",
     f"Average first-year ACV ${v1['acv'] / 1e3:.1f}K → ${b[('dist', 'lean', 1)]['acv'] / 1e3:.1f}K"),
    ("C1 sales-led", "Demand fixed regardless of motion; extra AEs could only add cost", "Dedicated AEs work extra accounts (meeting-rate multiplier); program spend above $25K/qtr buys MQLs at a CPL with diminishing returns",
     f"v1: sales-led ARR = lean ${v1['exit'] / 1e6:.2f}M. v2: sales-led ${b[('dist', 'sales', 1)]['exit_arr'] / 1e6:.2f}M vs lean ${b[('dist', 'lean', 1)]['exit_arr'] / 1e6:.2f}M"),
    ("C3 valuation", "Undiscounted 'total value'; 35/45/20 weights on joint scenarios; perfect gate; no incrementality", "NPV at 12%; incrementality haircut on gross profit; Monte Carlo with correlated inputs; gate with observation noise",
     f"v1 total value {v1['tv'] / 1e6:+.2f}M → v2 NPV {b[('dist', 'lean', 1)]['npv'] / 1e6:+.2f}M"),
    ("C4 alternatives", "Only distribution modeled, against zero", "Professional services on the same engine; Option 0 benchmark in every run; solution-play motion added", "See Scenarios and Executive Output"),
    ("M2 SAM", "CRM-ready 65% Base; SAM flexed 52–76% across scenarios", "CRM in use 56% (S25, 2026) fixed across scenarios; only performance flexes", "Serviceable accounts no longer flex"),
    ("M4 funnel", "Booked meetings run through MQL→SQL and SQL→opp benchmarks; win rate 28%", "Stages: meeting / intro / MQL → opportunity → win; Base win rate 21% (HubSpot average)", ""),
    ("m1 CAC", "Fully loaded only", "Fully loaded and marginal CAC", ""),
    ("m3 fatigue", "Constant meeting rate on re-sequenced accounts", "Meeting rate decays by wave age (75% retention per year, Base)", ""),
    ("De-dup", "0.973 on the 20–49 band for distribution only", "Band-specific factors for both sectors (54: 0.951–0.988); 423 mid bands remain unmeasurable", ""),
    ("Pending", "—", "Scorecard rebuild (06, review C5), sizing refresh (02/03, M1), documents 08/09 and the deck", "Next step"),
]
for i, row in enumerate(CHANGES):
    for j, v in enumerate(row):
        put(chg, f"{L(1 + j)}{5 + i}", v)
    chg.row_dimensions[5 + i].height = 40

# ============================================================ COVER
title(cover, "07 — Financial Model v2: should Clay fund a U.S. vertical program in 2027–2029?",
      "Independent portfolio case study by Zhuoran (Zora) Wang · Not commissioned by or affiliated with Clay Labs Inc. · v2 rebuilt 2026-10-19 after partner review")
rows = [
    ("Question", "Does any vertical program (distribution or professional services; lean pilot, sales-led team or solution play) create value above a 12% hurdle and above Option 0 — the same money spent in Clay's core motion?"),
    ("Structure", "Assumptions (3 selectors) → Market Sizing → ACV Build → GTM Funnel → Revenue Model → Unit Economics (live) → Scenarios, Sensitivity (full model copies on Engine Runs) → Monte Carlo (Python mirror, reconciled run by run) → Executive Output. Change Log maps v1 → v2."),
    ("Time and units", "Quarterly, Q1 2027–Q4 2029; annual totals P–R, 2027–29 in S. US$ nominal. ARR is end-of-quarter run-rate; recognized revenue = average ARR ÷ 4."),
    ("Valuation", "NPV = Σ PV[(gross profit × (1 − incrementality haircut)) − incremental cost] + PV of 5 years of gross profit from the 2029 ARR base (net of CS, decaying at GRR, haircut). Option 0 NPV is computed on the same non-CS spend."),
    ("Evidence classes", "A verified fact · B benchmark · C analyst estimate · D scenario assumption · E recommendation. Each input row carries its class, Monte Carlo group and source."),
    ("Color legend", "Blue = input · black = formula · green = link to another tab · yellow fill = judgment · green fill = key output · red fill = negative NPV."),
    ("Not a forecast", "Outputs are scenario and simulation estimates, not Clay forecasts. Monte Carlo values are static results of appendix/clay_engine.py (seed 42)."),
]
for i, (k, v) in enumerate(rows, start=4):
    put(cover, f"A{i}", k, bold=True, fill=SFILL)
    put(cover, f"B{i}", v)
    cover.row_dimensions[i].height = 46
cover.column_dimensions["A"].width = 18
cover.column_dimensions["B"].width = 125

# ============================================================ SOURCES
title(src, "Sources", "Full register with dates, reliability and limitations: Research_Sources.csv")
hdr(src, 4, ["ID", "Source", "Publisher", "Date / reference", "Used for", "URL"], widths=[7, 52, 28, 22, 44, 70])
SRCS = [
    ("S01", "Clay pricing page (plans, Actions & Data Credits, Enterprise minimums, unlimited seats)", "Clay Labs Inc.", "Live, Oct 2026", "ACV Build", "https://www.clay.com/pricing"),
    ("S14", "2022 Economic Census — Employment Size of Firms", "U.S. Census Bureau", "Ref. year 2022", "Firm counts, employees per firm", "https://data.census.gov/table/ECNSIZE2022.EC2200SIZEEMPFIRM"),
    ("S15", "OEWS May 2025 industry-specific estimates", "U.S. BLS", "May 2025", "Sales share of employment", "https://data.bls.gov/oes/"),
    ("S16", "Clay buyer pricing data", "Vendr", "Feb 2026", "Enterprise discount, ACV check", "https://www.vendr.com/marketplace/clay"),
    ("S25", "State of Sales in Distribution 2026 (Phocas), reported by Distribution Strategy Group", "Distribution Strategy Group", "2026", "CRM use 56% (SAM filter 2)", "see Research_Sources.csv"),
    ("S43", "2025 B2B SaaS Retention Benchmarks", "SaaS Capital", "2025-09", "GRR / NRR; Option 0 core retention", "https://www.saas-capital.com/wp-content/uploads/2025/09/RB32WS1-2025-B2B-SaaS-Retention-Benchmarks.pdf"),
    ("S44", "2026 AE Models, Motions & Metrics", "The Bridge Group", "2026-06-22", "AE quota, OTE, attainment", "https://www.bridgegroupinc.com/research/2026-ae-models-motions-metrics"),
    ("S45", "Employer Costs for Employee Compensation, June 2026", "U.S. BLS (TED)", "2026-09", "Cost loading", "https://www.bls.gov/opub/ted/2026/compensation-costs-for-private-industry-workers-averaged-46-89-per-hour-worked-in-june-2026.htm"),
    ("S46", "OEWS May 2025 — Software Publishers (513200) wages", "U.S. BLS", "May 2025", "Role wages", "https://data.bls.gov/oes/#/industry/513200/2025"),
    ("S47", "CAC payback benchmarks (Maxio, Benchmarkit, MetricHQ compilation)", "SaaS Hero", "Updated 2026-07-18", "Option 0 payback", "https://www.saashero.net/strategy/maxio-cac-benchmarks-b2b-saas/"),
    ("S48", "Win rate statistics (HubSpot 2024; Ebsta 2024; Norwest 2025)", "Webtonic (compilation)", "2026", "Win rate, cycle", "https://www.webtonic.io/blog/win-rate-statistics"),
    ("S49", "Pipeline performance benchmarks 2025", "The Digital Bloom (compilation)", "2025-10-22", "Funnel conversion", "https://thedigitalbloom.com/learn/pipeline-performance-benchmarks-2025/"),
    ("S50", "What margin should we give partners? (Freshworks, Bessemer)", "Echai Ventures", "Undated", "Partner fees, subsidy", "https://echai.ventures/gtm/gtm-channel-ecosystem/what-margin-should-we-give-partners-and-how-do-we-stop-channel-conflict"),
    ("S52", "Clay company profile", "Sacra", "2026", "Gross-margin context", "https://sacra.com/c/clay/"),
    ("S55", "B2B cost per lead and conversion benchmarks 2025", "The Starr Conspiracy", "2026-05-22", "CPL ($200–350 commonly cited for SaaS)", "https://thestarrconspiracy.com/insights/benchmarks/b2b-cost-per-lead-conversion-benchmarks-2025"),
]
for i, s_ in enumerate(SRCS, start=5):
    for j, v in enumerate(s_, start=1):
        put(src, f"{L(j)}{i}", v)

order = ["Cover", "Executive Output", "Assumptions", "Market Sizing", "ACV Build", "GTM Funnel", "Revenue Model", "Unit Economics", "Scenarios", "Sensitivity", "Monte Carlo", "Engine Runs", "Change Log", "Sources"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Cover": "1F2A44", "Executive Output": "1F2A44", "Assumptions": "FFC000", "Monte Carlo": "70AD47", "Engine Runs": "A6A6A6", "Change Log": "A6A6A6", "Sources": "A6A6A6"}.get(w.title, "4472C4")
wb.save(OUT)
print("saved", OUT, "engine rows", eng_row[0], "runs", len(RUNSPEC))
for m_ in msgs:
    print(m_)
