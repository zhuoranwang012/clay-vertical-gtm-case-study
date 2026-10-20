"""Export the v2 workbook outputs the deck charts use (07, 06, 02) to appendix/deck_data_from_model.json.
Run after recalculating the workbooks:  python3 appendix/export_deck_data.py
Then:  node appendix/build_deck.js appendix/deck_data_from_model.json Executive_Strategy_Deck.pptx"""
import json
import os
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = {}
w7 = load_workbook(os.path.join(BASE, "07_Financial_Model.xlsx"), data_only=True)
s = w7["Scenarios"]
heads = [s.cell(4, c).value for c in range(1, s.max_column + 1)]
out["scen"] = {}
for r in range(5, 23):
    lab = s.cell(r, 1).value
    if lab:
        out["scen"][lab] = {heads[c - 1]: s.cell(r, c).value for c in range(3, len(heads) + 1) if heads[c - 1]}
for r in range(30, 60):
    if s.cell(r, 1).value and str(s.cell(r, 1).value).startswith("ARR trajectory"):
        out["arr"] = {s.cell(rr, 1).value: [s.cell(rr, c).value for c in range(3, 15)] for rr in range(r + 1, r + 6)}
        break
mc = w7["Monte Carlo"]
out["mc"] = {mc.cell(r, 1).value: {"mean": mc.cell(r, 2).value, "p10": mc.cell(r, 3).value, "p50": mc.cell(r, 4).value, "p90": mc.cell(r, 5).value,
                                   "ppos": mc.cell(r, 6).value, "x_mean": mc.cell(r, 7).value, "x_ppos": mc.cell(r, 8).value,
                                   "stop": mc.cell(r, 9).value, "scale": mc.cell(r, 10).value} for r in range(5, 13)}
out["rob"] = [[mc.cell(r, c).value for c in range(1, 8)] for r in range(16, 22)]
hist = None
for r in range(1, mc.max_row + 1):
    if mc.cell(r, 1).value and str(mc.cell(r, 1).value).startswith("NPV distribution"):
        hist = r + 2
        break
out["hist"] = []
r = hist
while mc.cell(r, 1).value is not None and isinstance(mc.cell(r, 1).value, (int, float)):
    out["hist"].append([mc.cell(r, c).value for c in range(1, 5)])
    r += 1
se = w7["Sensitivity"]
out["tornado"] = [[se.cell(r, 1).value, se.cell(r, 6).value, se.cell(r, 4).value, se.cell(r, 5).value] for r in range(5, 23) if se.cell(r, 6).value is not None]
w6 = load_workbook(os.path.join(BASE, "06_Segment_Prioritization.xlsx"), data_only=True)
sc = w6["Scorecard"]
out["score"] = {}
for r in range(5, 25):
    a, b = sc.cell(r, 1).value, sc.cell(r, 2).value
    if a and b == "score 1–5":
        out["score"][a] = [sc.cell(r, c).value for c in (3, 4, 5)]
    if a == "Weighted total":
        out["score"]["total"] = [sc.cell(r, c).value for c in (3, 4, 5)]
ex6 = w6["Executive_Summary"]
out["robust6"] = {ex6.cell(r, 1).value: [ex6.cell(r, c).value for c in (2, 3, 4)] for r in range(5, 15)}
dd = w6["Subsegment_Drilldown"]
out["drill"] = {dd.cell(r, 1).value: [dd.cell(r, c).value for c in (2, 3, 4)] for r in range(5, 16) if dd.cell(r, 1).value}
w2 = load_workbook(os.path.join(BASE, "02_Market_Sizing.xlsx"), data_only=True)
e2 = w2["Executive_Summary"]
fc2 = w2["Firm_Counts"]
out["estab"] = {fc2.cell(r, 1).value: fc2.cell(r, 4).value for r in range(60, 80) if fc2.cell(r, 1).value and isinstance(fc2.cell(r, 4).value, (int, float))}
out["sizing"] = {e2.cell(r, 1).value: [e2.cell(r, c).value for c in (2, 3, 4)] for r in range(5, 27) if e2.cell(r, 1).value}
json.dump(out, open(os.path.join(BASE, "appendix", "deck_data_from_model.json"), "w"), indent=1, default=str)
print("wrote appendix/deck_data_from_model.json")
