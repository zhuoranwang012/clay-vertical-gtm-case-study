"""
Contractor coverage test (review issue M5; preview of sprint test T1)
U.S. Vertical Expansion & GTM Strategy for Clay - independent case study, not affiliated with Clay Labs Inc.

Two subcommands:

  python3 appendix/coverage_test.py template
      Builds an empty appendix/Coverage_Test_Log.xlsx (log, live results with Wilson CIs,
      pre-registered verdict, re-check sheet, definitions).

  python3 appendix/coverage_test.py sample --input <CSLB download> [--input <another>] [--n 100] [--reserve 20] [--seed 2026]
      Reads the CSLB "List by Classification" download(s) (.xls / .xlsx / .csv / HTML-table .xls),
      keeps currently licensed firms holding C-10, C-20 or C-36, draws a seeded, stratified random
      sample plus reserves, and writes it into the log. Re-running with the same file and seed
      reproduces the same sample.

Nothing in this script creates results. All coverage observations are entered by hand.
"""
import argparse
import os
import random
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "appendix", "Coverage_Test_Log.xlsx")

STRATA = [("C-10", "Electrical contractor", r"\bC[- ]?10\b"),
          ("C-20", "HVAC contractor", r"\bC[- ]?20\b"),
          ("C-36", "Plumbing contractor", r"\bC[- ]?36\b")]
N_ROWS = 120          # 100 primary + 20 reserve (max)
FIRST = 6             # first data row on Log
LAST = FIRST + N_ROWS - 1

F = "Arial"
HDR = Font(name=F, bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1F2A44")
INFILL = PatternFill("solid", fgColor="FFF2CC")      # hand-entered observations
CALCFILL = PatternFill("solid", fgColor="E8ECF3")    # formulas
BOLD = Font(name=F, bold=True, size=10)
NORM = Font(name=F, size=10)
BLUE = Font(name=F, size=10, color="0000FF")
SUB = Font(name=F, italic=True, size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Log columns: (letter, header, width, kind, validation list or None)
COLS = [
    ("A", "ID", 7, "key", None),
    ("B", "Draw", 9, "key", None),
    ("C", "Stratum", 8, "key", None),
    ("D", "License #", 11, "key", None),
    ("E", "Business name", 34, "key", None),
    ("F", "City", 16, "key", None),
    ("G", "ZIP", 8, "key", None),
    ("H", "Business phone", 14, "key", None),
    ("I", "Status", 22, "in", "Done,Skipped - out of business,Skipped - cannot identify firm,Not started"),
    ("J", "Website found", 9, "in", "Y,N"),
    ("K", "Google Business Profile", 10, "in", "Y,N"),
    ("L", "Web DM identifiable", 10, "in", "Y,N"),
    ("M", "Registry person of record", 10, "in", "Y,N"),
    ("N", "Apollo company record", 10, "in", "Y,N"),
    ("O", "Apollo contacts listed", 10, "in", "0,1-2,3+"),
    ("P", "Apollo DM listed", 10, "in", "Y,N"),
    ("Q", "Apollo DM email shown available", 11, "in", "Y,N,Can't tell"),
    ("R", "LinkedIn company page", 10, "in", "Y,N"),
    ("S", "LinkedIn DM profile", 10, "in", "Y,N"),
    ("T", "Minutes", 8, "in", None),
    ("U", "Notes (no personal data)", 30, "in", None),
    ("V", "Done?", 7, "calc", None),
    ("W", "Apollo ∪ web DM", 9, "calc", None),
    ("X", "Apollo ∪ web ∪ registry DM", 10, "calc", None),
    ("Y", "Apollo ∪ LinkedIn DM", 9, "calc", None),
    ("Z", "Business type (CSLB)", 14, "key", None),
]


def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR, HFILL, CENTER, BOX


def put(ws, ref, val, font=NORM, fmt=None, fill=None, wrap=False, border=True, align=None):
    c = ws[ref]
    c.value = val
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    if align:
        c.alignment = align
    if border:
        c.border = BOX
    return c


def title(ws, t, s):
    ws["A1"].value, ws["A1"].font = t, Font(name=F, bold=True, size=14)
    ws["A2"].value, ws["A2"].font = s, SUB


def build_template():
    wb = Workbook()

    # ------------------------------------------------------------ README
    rd = wb.active
    rd.title = "README"
    title(rd, "Contractor coverage test — log and results (review M5; preview of sprint test T1)",
          "Independent case study · not commissioned by or affiliated with Clay Labs Inc. · all observations entered by hand; this workbook contains no results until the test is run")
    rows = [
        ("Question", "For contractors that electrical, HVAC and plumbing distributors (NAICS 4236/4237) sell to, how often does a standard B2B database (Apollo) list a decision-maker, and how much does adding public-web sources raise that share?"),
        ("Why it matters", "The 'data your database misses' positioning (08 v2) and H3 rest on an education-based LinkedIn proxy (05) and one vendor case study (S07). This test replaces the proxy with a measured match rate."),
        ("Sample frame", "CSLB 'List by Classification' download (S56): currently licensed California contractors in C-10, C-20, C-36. Seeded, stratified random draw (appendix/coverage_test.py sample)."),
        ("Protocol", "appendix/Coverage_Test_Protocol.md — definitions, search procedure, stopping and replacement rules, pre-registered verdict. Definitions are frozen after the 10-firm pilot."),
        ("How to use", "1) Run the sampler. 2) Pilot the first 10 IDs; adjust definitions only now and note changes in Definitions. 3) Work through Primary rows in ID order; use Reserve rows only to replace Skipped rows, in order, within the same stratum. 4) Re-check 15 random Done rows on a different day (Recheck sheet). 5) Read Results."),
        ("Colour code", "Yellow = hand-entered observation. Grey-blue = formula. Blue text = input you may change (Results: z value, proxy, thresholds)."),
        ("Data handling", "Record Y/N flags and title categories only. Do not record individuals' names, emails, phone numbers or profile URLs. Business-level registry fields are public (CSLB). No automated scraping; manual lookups only, within each site's terms."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        put(rd, f"A{i}", k, font=BOLD, fill=CALCFILL)
        put(rd, f"B{i}", v, wrap=True)
        rd.row_dimensions[i].height = 46
    rd.column_dimensions["A"].width = 18
    rd.column_dimensions["B"].width = 120

    # ------------------------------------------------------------ Log
    lg = wb.create_sheet("Log")
    title(lg, "Log — one row per sampled contractor",
          "Columns A–H are written by the sampler. Enter observations in the yellow columns; formulas in V–Y. Status must be 'Done' for a row to count.")
    lg["A3"].value = "Sample frame file:"
    lg["A3"].font = BOLD
    lg["D3"].value = "(not yet sampled)"
    lg["D3"].font = BLUE
    lg["I3"].value = "Seed / draw date:"
    lg["I3"].font = BOLD
    lg["L3"].value = ""
    lg["L3"].font = BLUE
    hdr(lg, 5, [c[1] for c in COLS])
    lg.row_dimensions[5].height = 54
    for letter, _, w, kind, opts in COLS:
        lg.column_dimensions[letter].width = w
        if opts:
            dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
            dv.error = "Choose from the list"
            dv.add(f"{letter}{FIRST}:{letter}{LAST}")
            lg.add_data_validation(dv)
    dvm = DataValidation(type="decimal", operator="between", formula1="0", formula2="60", allow_blank=True)
    dvm.add(f"T{FIRST}:T{LAST}")
    lg.add_data_validation(dvm)
    for r in range(FIRST, LAST + 1):
        for letter, _, _, kind, _ in COLS:
            c = lg[f"{letter}{r}"]
            c.border = BOX
            c.font = NORM
            if kind == "in":
                c.fill = INFILL
            elif kind == "calc":
                c.fill = CALCFILL
        put(lg, f"V{r}", f'=IF(I{r}="Done",1,0)', fill=CALCFILL, align=CENTER)
        put(lg, f"W{r}", f'=IF(V{r}=1,IF(OR(P{r}="Y",L{r}="Y"),1,0),"")', fill=CALCFILL, align=CENTER)
        put(lg, f"X{r}", f'=IF(V{r}=1,IF(OR(P{r}="Y",L{r}="Y",M{r}="Y"),1,0),"")', fill=CALCFILL, align=CENTER)
        put(lg, f"Y{r}", f'=IF(V{r}=1,IF(OR(P{r}="Y",S{r}="Y"),1,0),"")', fill=CALCFILL, align=CENTER)
    lg.freeze_panes = f"F{FIRST}"

    # ------------------------------------------------------------ Results
    rs = wb.create_sheet("Results")
    title(rs, "Results — live from Log (share of Done rows; Wilson 95% intervals)",
          "Empty until rows are marked Done. Rates are for currently licensed California C-10/C-20/C-36 contractors; they are not national estimates.")
    put(rs, "A3", "z (95%)", font=BOLD)
    put(rs, "B3", 1.96, font=BLUE, fmt="0.00")
    put(rs, "C3", "LinkedIn proxy used in 05/06 (contractor personas)", font=BOLD)
    put(rs, "F3", 0.20, font=BLUE, fmt="0%")
    put(rs, "G3", "Minimum n for a verdict", font=BOLD)
    put(rs, "I3", 60, font=BLUE, fmt="0")
    rng = lambda col: f"Log!${col}${FIRST}:${col}${LAST}"
    groups = [("All", None)] + [(s[0], s[0]) for s in STRATA]
    hdr(rs, 5, ["Metric", "Definition"] + sum([[f"{g} rate", f"{g} low", f"{g} high"] for g, _ in groups], []))
    rs.column_dimensions["A"].width = 30
    rs.column_dimensions["B"].width = 44
    for j in range(3, 3 + 3 * len(groups)):
        rs.column_dimensions[L(j)].width = 9

    # n rows
    put(rs, "A6", "Firms completed (n)", font=BOLD)
    put(rs, "B6", "Status = Done")
    put(rs, "A7", "Skipped share", font=BOLD)
    put(rs, "B7", "Skipped ÷ (Done + Skipped) — closed or unidentifiable firms")
    for g_i, (g, s) in enumerate(groups):
        c = L(3 + 3 * g_i)
        cond = "" if s is None else f',{rng("C")},"{s}"'
        put(rs, f"{c}6", f'=COUNTIFS({rng("I")},"Done"{cond})', fmt="0", fill=CALCFILL, font=BOLD)
        put(rs, f"{c}7", f'=IFERROR(COUNTIFS({rng("I")},"Skipped*"{cond})/(COUNTIFS({rng("I")},"Skipped*"{cond})+{c}6),"")', fmt="0%", fill=CALCFILL)

    metrics = [
        ("Apollo company record", "N", "Apollo has an organization record matching the firm"),
        ("Apollo decision-maker listed  [primary]", "P", "Apollo lists ≥ 1 named contact with a DM title at the firm"),
        ("Apollo DM email shown available", "Q", "Apollo indicates an email is available for that DM (not revealed)"),
        ("LinkedIn company page", "R", "Firm has a LinkedIn company page"),
        ("LinkedIn decision-maker profile", "S", "≥ 1 LinkedIn profile names the firm as current employer with a DM title"),
        ("Website found", "J", "Firm's own website located"),
        ("Google Business Profile", "K", "Firm has a Google Maps / Business Profile"),
        ("Web DM identifiable", "L", "DM name + role identifiable on website or Google profile"),
        ("Registry person of record", "M", "CSLB license detail lists a qualifier / officer / owner"),
    ]
    MROW = {}
    for k, (lab, col, d) in enumerate(metrics):
        r = 9 + k
        put(rs, f"A{r}", lab, font=BOLD if "primary" in lab else NORM)
        put(rs, f"B{r}", d, wrap=True)
        for g_i, (g, s) in enumerate(groups):
            c0, c1, c2 = L(3 + 3 * g_i), L(4 + 3 * g_i), L(5 + 3 * g_i)
            cond = "" if s is None else f',{rng("C")},"{s}"'
            n = f"{c0}$6"
            put(rs, f"{c0}{r}", f'=IFERROR(COUNTIFS({rng("I")},"Done",{rng(col)},"Y"{cond})/{n},"")', fmt="0%", fill=CALCFILL)
            p, z = f"{c0}{r}", "$B$3"
            core = f"({p}+{z}^2/(2*{n}))/(1+{z}^2/{n})"
            half = f"({z}*SQRT({p}*(1-{p})/{n}+{z}^2/(4*{n}^2)))/(1+{z}^2/{n})"
            put(rs, f"{c1}{r}", f'=IFERROR(MAX(0,{core}-{half}),"")', fmt="0%", fill=CALCFILL)
            put(rs, f"{c2}{r}", f'=IFERROR(MIN(1,{core}+{half}),"")', fmt="0%", fill=CALCFILL)
        MROW[col] = r
    # union metrics from computed columns
    unions = [("Apollo ∪ web DM  [Clay-style waterfall]", "W", "DM found in Apollo or on website / Google profile"),
              ("Apollo ∪ web ∪ registry DM  [upper bound]", "X", "Adds CSLB person of record"),
              ("Apollo ∪ LinkedIn DM", "Y", "Best of the two LinkedIn-derived sources")]
    for k, (lab, col, d) in enumerate(unions):
        r = 9 + len(metrics) + k
        put(rs, f"A{r}", lab, font=BOLD if "waterfall" in lab else NORM)
        put(rs, f"B{r}", d, wrap=True)
        for g_i, (g, s) in enumerate(groups):
            c0, c1, c2 = L(3 + 3 * g_i), L(4 + 3 * g_i), L(5 + 3 * g_i)
            cond = "" if s is None else f',{rng("C")},"{s}"'
            n = f"{c0}$6"
            put(rs, f"{c0}{r}", f'=IFERROR(SUMIFS({rng(col)},{rng("V")},1{cond})/{n},"")', fmt="0%", fill=CALCFILL)
            p, z = f"{c0}{r}", "$B$3"
            core = f"({p}+{z}^2/(2*{n}))/(1+{z}^2/{n})"
            half = f"({z}*SQRT({p}*(1-{p})/{n}+{z}^2/(4*{n}^2)))/(1+{z}^2/{n})"
            put(rs, f"{c1}{r}", f'=IFERROR(MAX(0,{core}-{half}),"")', fmt="0%", fill=CALCFILL)
            put(rs, f"{c2}{r}", f'=IFERROR(MIN(1,{core}+{half}),"")', fmt="0%", fill=CALCFILL)
        MROW[col] = r

    r0 = 9 + len(metrics) + len(unions) + 1
    put(rs, f"A{r0}", "Decision metrics (All strata)", font=BOLD, fill=CALCFILL)
    put(rs, f"B{r0}", None, fill=CALCFILL)
    ap, wf = MROW["P"], MROW["W"]
    dm = [
        ("Uplift: (Apollo ∪ web) ÷ Apollo  [T1 analogue]", f'=IFERROR(C{wf}/C{ap},"")', '0.00"×"', "Point estimate. Sprint threshold T1 ≥ 1.5×."),
        ("Conservative uplift", f'=IFERROR(D{wf}/E{ap},"")', '0.00"×"', "Lower bound of union ÷ upper bound of Apollo (both Wilson 95%)."),
        ("Upper-bound uplift (with registry)", f'=IFERROR(C{MROW["X"]}/C{ap},"")', '0.00"×"', "Includes the CSLB person of record; reported, not used for the verdict."),
        ("Measured LinkedIn DM coverage − proxy", f'=IFERROR(C{MROW["S"]}-$F$3,"")', '+0%;-0%;0%', "Positive = the 05 proxy understated LinkedIn coverage."),
        ("Supported threshold (uplift)", 1.5, '0.00"×"', "Pre-registered; change only before data collection."),
        ("Partial threshold (uplift)", 1.2, '0.00"×"', "Pre-registered."),
        ("Apollo DM coverage, firms with a website", f'=IFERROR(COUNTIFS({rng("I")},"Done",{rng("J")},"Y",{rng("P")},"Y")/COUNTIFS({rng("I")},"Done",{rng("J")},"Y"),"")', "0%", "Rough size proxy: CSLB has no employee counts; firms with a website skew larger."),
        ("Apollo DM coverage, sole owners", f'=IFERROR(COUNTIFS({rng("I")},"Done",{rng("Z")},"Sole Owner",{rng("P")},"Y")/COUNTIFS({rng("I")},"Done",{rng("Z")},"Sole Owner"),"")', "0%", "CSLB business type; sole proprietors are mostly very small firms."),
        ("Apollo DM coverage, corporations / LLCs / partnerships", f'=IFERROR((COUNTIFS({rng("I")},"Done",{rng("P")},"Y")-COUNTIFS({rng("I")},"Done",{rng("Z")},"Sole Owner",{rng("P")},"Y"))/(COUNTIFS({rng("I")},"Done")-COUNTIFS({rng("I")},"Done",{rng("Z")},"Sole Owner")),"")', "0%", "Closer to distributors' mid-market contractor customers; the better size proxy."),
        ("Apollo DM coverage, firms without a website", f'=IFERROR(COUNTIFS({rng("I")},"Done",{rng("J")},"N",{rng("P")},"Y")/COUNTIFS({rng("I")},"Done",{rng("J")},"N"),"")', "0%", "Mostly sole proprietors and very small firms."),
    ]
    for k, (lab, f, fmt, note) in enumerate(dm):
        r = r0 + 1 + k
        put(rs, f"A{r}", lab, font=BOLD if k == 0 else NORM)
        put(rs, f"C{r}", f, fmt=fmt, fill=CALCFILL if isinstance(f, str) else None, font=BLUE if not isinstance(f, str) else BOLD)
        put(rs, f"B{r}", note, wrap=True)
    up, cons, sup, par = r0 + 1, r0 + 2, r0 + 5, r0 + 6
    rv = r0 + len(dm) + 2
    put(rs, f"A{rv}", "Pre-registered verdict on H3 / coverage claim", font=BOLD, fill=CALCFILL)
    verdict = (f'=IF(C6<$I$3,"Insufficient sample: n = "&C6&" (need ≥ "&$I$3&")",'
               f'IF(AND(C{up}>=C{sup},C{cons}>=C{sup}),"SUPPORTED (robust): measured uplift clears 1.5× even at the conservative bound — keep coverage-led positioning, cite the measured rate",'
               f'IF(C{up}>=C{sup},"SUPPORTED (point estimate only): uplift ≥ 1.5× but the conservative bound is below — claim a measured gap, flag the interval",'
               f'IF(C{up}>=C{par},"PARTIAL: uplift 1.2–1.5× — soften to \'adds coverage on a material share of contractors\'; T1 must still be run at scale",'
               f'"NOT SUPPORTED: uplift < 1.2× — withdraw coverage-led positioning (roadmap R4); H3 becomes \'not supported\'"))))')
    put(rs, f"B{rv}", verdict, font=BOLD, wrap=True)
    rs.merge_cells(f"B{rv}:N{rv}")
    rs.row_dimensions[rv].height = 34
    put(rs, f"A{rv + 2}", "Reading", font=BOLD)
    put(rs, f"B{rv + 2}", "With n = 100 the 95% interval on a share near 50% is about ±10 points; by stratum (n ≈ 33) about ±16 points. Treat stratum differences as directional. The sample is California only and contractors only (customers of 4236/4237); plant buyers of 4238 MRO distributors are not covered.", wrap=True)
    rs.merge_cells(f"B{rv + 2}:N{rv + 2}")
    rs.row_dimensions[rv + 2].height = 44
    rs.freeze_panes = "C6"

    # ------------------------------------------------------------ Recheck
    rc = wb.create_sheet("Recheck")
    title(rc, "Re-check — 15 random Done rows, re-recorded on a different day without looking at the first entry",
          "Agreement below 85% on any field means the definition is ambiguous: tighten it in Definitions and re-run that field on all rows.")
    hdr(rc, 4, ["ID", "Apollo DM listed (re-check)", "Web DM identifiable (re-check)", "LinkedIn DM profile (re-check)",
                "Apollo match", "Web match", "LinkedIn match"])
    for j, w in enumerate([8, 14, 14, 14, 10, 10, 10], start=1):
        rc.column_dimensions[L(j)].width = w
    rc.row_dimensions[4].height = 40
    for k in range(15):
        r = 5 + k
        put(rc, f"A{r}", None, fill=INFILL)
        for col in "BCD":
            put(rc, f"{col}{r}", None, fill=INFILL)
        for col, src, logcol in (("E", "B", "P"), ("F", "C", "L"), ("G", "D", "S")):
            put(rc, f"{col}{r}", f'=IF(OR(A{r}="",{src}{r}=""),"",IF(INDEX(Log!${logcol}${FIRST}:${logcol}${LAST},MATCH(A{r},Log!$A${FIRST}:$A${LAST},0))={src}{r},1,0))', fill=CALCFILL, align=CENTER)
    for col in "BCD":
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        dv.add(f"{col}5:{col}19")
        rc.add_data_validation(dv)
    put(rc, "A21", "Agreement", font=BOLD)
    for col in "EFG":
        put(rc, f"{col}21", f'=IFERROR(AVERAGE({col}5:{col}19),"")', fmt="0%", fill=CALCFILL, font=BOLD)

    # ------------------------------------------------------------ Definitions
    df = wb.create_sheet("Definitions")
    title(df, "Definitions — frozen after the 10-firm pilot (record any change below with date and reason)", "Full procedure: appendix/Coverage_Test_Protocol.md")
    hdr(df, 4, ["Field", "Definition", "Where to look"])
    defs = [
        ("Decision-maker (DM) title", "Owner, founder, president, CEO, principal, partner, vice president, general manager, operations manager, purchasing / procurement manager, estimator / chief estimator, project manager. Not: technicians, apprentices, office staff without a buying role.", "—"),
        ("Website found", "The firm's own site (not a directory listing), matched on name and city or license number.", "Google: \"<business name>\" <city>"),
        ("Google Business Profile", "A Maps / Business Profile for the firm at the registry city.", "Google Maps"),
        ("Web DM identifiable", "A DM's name and role appear on the firm's website or Google profile (about / team / owner reply).", "Website, Google profile"),
        ("Registry person of record", "CSLB license detail lists a qualifier (RMO/RME), officer or owner.", "CSLB Instant License Check"),
        ("Apollo company record", "An Apollo organization matches the firm on name and city / domain.", "Apollo free account, Companies search"),
        ("Apollo contacts listed", "Count of people Apollo shows at that organization (any title).", "Apollo company page → People"),
        ("Apollo DM listed", "At least one listed person has a DM title (above). Do not reveal emails or phones.", "Apollo company page → People"),
        ("Apollo DM email shown available", "Apollo's interface indicates an email exists for a listed DM without revealing it; 'Can't tell' if the interface does not show it.", "Apollo"),
        ("LinkedIn company page", "A company page matching name and location.", "LinkedIn search or Google site:linkedin.com/company"),
        ("LinkedIn DM profile", "At least one profile lists the firm as current employer with a DM title.", "LinkedIn people search or Google site:linkedin.com/in"),
        ("Done / Skipped", "Done = all fields checked. Skipped - out of business = phone disconnected and no web presence, or marked closed. Skipped - cannot identify = name too generic to match after 5 minutes.", "—"),
    ]
    for i, (a, b, c) in enumerate(defs, start=5):
        put(df, f"A{i}", a, font=BOLD)
        put(df, f"B{i}", b, wrap=True)
        put(df, f"C{i}", c, wrap=True)
        df.row_dimensions[i].height = 44
    df.column_dimensions["A"].width = 26
    df.column_dimensions["B"].width = 90
    df.column_dimensions["C"].width = 34
    ch = 5 + len(defs) + 1
    put(df, f"A{ch}", "Definition changes (pilot only)", font=BOLD, fill=CALCFILL)
    hdr(df, ch + 1, ["Date", "Change", "Reason"])
    for r in range(ch + 2, ch + 7):
        for col in "ABC":
            put(df, f"{col}{r}", None, fill=INFILL)

    for ws, color in ((rd, "1F2A44"), (lg, "FFC000"), (rs, "4472C4"), (rc, "4472C4"), (df, "A6A6A6")):
        ws.sheet_properties.tabColor = color
    wb.save(OUT)
    print("template written:", OUT)


# ---------------------------------------------------------------- sampling
def read_frame(path):
    import pandas as pd
    errors = []
    readers = []
    low = path.lower()
    if low.endswith(".csv"):
        readers.append(lambda: pd.read_csv(path, dtype=str))
    readers += [lambda: pd.read_excel(path, dtype=str),
                lambda: pd.read_html(path)[0].astype(str),
                lambda: pd.read_csv(path, dtype=str, sep=None, engine="python")]
    for rd in readers:
        try:
            df = rd()
            if len(df.columns) > 2 and len(df) > 0:
                return df
        except Exception as e:  # try the next reader
            errors.append(str(e)[:120])
    sys.exit(f"Could not read {path}. Tried Excel, HTML table and CSV. Errors: {errors}. "
             "If it is a real .xls file, install xlrd (pip install xlrd) or re-save it as .xlsx / .csv.")


def find_col(cols, *patterns, required=True):
    for p in patterns:
        for c in cols:
            if re.search(p, str(c), re.I):
                return c
    if required:
        sys.exit(f"No column matching {patterns}. Columns found: {list(cols)}. Rename the column or edit find_col patterns.")
    return None


def draw_sample(inputs, n, reserve, seed, state="CA"):
    import pandas as pd
    frames = [read_frame(p) for p in inputs]
    df = pd.concat(frames, ignore_index=True)
    cols = df.columns
    c_lic = find_col(cols, r"lic.*(no|num|#)", r"^lic", r"license")
    c_name = find_col(cols, r"business.*name", r"^name", r"business")
    c_cls = find_col(cols, r"class")
    c_status = find_col(cols, r"status", required=False)
    c_city = find_col(cols, r"city", required=False)
    c_zip = find_col(cols, r"zip", r"postal", required=False)
    c_phone = find_col(cols, r"phone", r"tele", required=False)
    c_addr = find_col(cols, r"address", required=False)
    c_state = find_col(cols, r"^state$", required=False)
    c_btype = find_col(cols, r"business.?type", r"entity", required=False)

    df = df.dropna(subset=[c_lic, c_name]).drop_duplicates(subset=[c_lic])
    n_frame = len(df)
    if c_status:
        keep = df[c_status].fillna("").str.contains(r"clear|active|current", case=False, regex=True)
        df = df[keep]
    n_active = len(df)
    if c_state and state:
        df = df[df[c_state].fillna("").str.strip().str.upper() == state.upper()]
    n_state = len(df)

    rng = random.Random(seed)
    per = [n // 3 + (1 if i < n % 3 else 0) for i in range(3)]
    per_res = [reserve // 3 + (1 if i < reserve % 3 else 0) for i in range(3)]
    used = set()
    rows = []
    frame_counts = {}
    for (code, _, pat), k, kr in zip(STRATA, per, per_res):
        pool = df[df[c_cls].fillna("").str.contains(pat, case=False, regex=True)]
        frame_counts[code] = len(pool)
        lic_list = sorted(set(pool[c_lic]) - used)
        if len(lic_list) < k + kr:
            sys.exit(f"Stratum {code} has only {len(lic_list)} eligible licenses; need {k + kr}.")
        pick = rng.sample(lic_list, k + kr)
        used.update(pick)
        sub = pool.set_index(c_lic)
        for j, lic in enumerate(pick):
            rec = sub.loc[lic]
            if isinstance(rec, pd.DataFrame):
                rec = rec.iloc[0]
            city = rec.get(c_city, "") if c_city else ""
            zp = rec.get(c_zip, "") if c_zip else ""
            if (not city or not zp) and c_addr:
                m = re.search(r"([A-Za-z .'-]+),?\s+CA\s+(\d{5})", str(rec.get(c_addr, "")))
                if m:
                    city = city or m.group(1).strip()
                    zp = zp or m.group(2)
            rows.append(dict(draw="Primary" if j < k else "Reserve", stratum=code, lic=str(lic),
                             name=str(rec.get(c_name, "")), city=str(city or ""), zip=str(zp or "")[:5],
                             phone=str(rec.get(c_phone, "") if c_phone else ""),
                             btype=str(rec.get(c_btype, "") if c_btype else "").strip()))
    # order: primary first (interleaved strata by draw order), then reserve
    prim = [r for r in rows if r["draw"] == "Primary"]
    res = [r for r in rows if r["draw"] == "Reserve"]
    rng.shuffle(prim)
    if len(prim) + len(res) > N_ROWS:
        sys.exit(f"Log holds {N_ROWS} rows; reduce --n or --reserve.")
    return prim + res, dict(frame=n_frame, active=n_active, in_state=n_state, strata=frame_counts)


def cmd_sample(args):
    if not os.path.exists(OUT):
        build_template()
    rows, stats = draw_sample(args.input, args.n, args.reserve, args.seed, args.state)
    wb = load_workbook(OUT)
    lg = wb["Log"]
    if any(lg[f"I{r}"].value not in (None, "", "Not started") for r in range(FIRST, LAST + 1)):
        sys.exit("The log already has observations. Rebuild the template first if you really want a new sample.")
    for r in range(FIRST, LAST + 1):
        for col in "ABCDEFGHZ":
            lg[f"{col}{r}"].value = None
    for i, rec in enumerate(rows):
        r = FIRST + i
        vals = [f"{'P' if rec['draw'] == 'Primary' else 'R'}{i + 1:03d}", rec["draw"], rec["stratum"], rec["lic"],
                rec["name"], rec["city"], rec["zip"], rec["phone"]]
        for col, v in zip("ABCDEFGH", vals):
            lg[f"{col}{r}"].value = v
        lg[f"Z{r}"].value = rec["btype"] or None
        lg[f"I{r}"].value = "Not started"
    src = ", ".join(os.path.basename(p) for p in args.input)
    lg["D3"].value = (f"{src} · {stats['frame']:,} licenses read · {stats['active']:,} current · {stats['in_state']:,} with a CA address · "
                      + " · ".join(f"{k}: {v:,}" for k, v in stats["strata"].items()))
    import datetime
    lg["L3"].value = f"seed {args.seed} · drawn {datetime.date.today().isoformat()} · n {args.n} + {args.reserve} reserve"
    wb.save(OUT)
    print(f"sample written to {OUT}: {len(rows)} rows ({args.n} primary + {args.reserve} reserve)")
    print("frame:", stats)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sp = ap.add_subparsers(dest="cmd", required=True)
    sp.add_parser("template")
    s = sp.add_parser("sample")
    s.add_argument("--input", action="append", required=True)
    s.add_argument("--n", type=int, default=100)
    s.add_argument("--reserve", type=int, default=20)
    s.add_argument("--seed", type=int, default=2026)
    s.add_argument("--state", default="CA", help="keep licenses with this mailing state ('' to keep all)")
    a = ap.parse_args()
    if a.cmd == "template":
        build_template()
    else:
        cmd_sample(a)
