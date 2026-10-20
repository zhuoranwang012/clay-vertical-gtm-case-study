"""Build 02_Market_Sizing.xlsx — Clay vertical expansion case study (Phase 1).
All outputs are formulas over raw-data and assumption tabs."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

BASE = "/home/claude/Clay_Strategy_Consulting_Case"
OUT = f"{BASE}/02_Market_Sizing.xlsx"

F = "Arial"
BLUE = Font(name=F, color="0000FF", size=10)
BLACK = Font(name=F, color="000000", size=10)
GREEN = Font(name=F, color="008000", size=10)
BOLD = Font(name=F, bold=True, size=10)
HDR = Font(name=F, bold=True, color="FFFFFF", size=10)
TITLE = Font(name=F, bold=True, size=14, color="1F2A44")
SUB = Font(name=F, italic=True, size=9, color="555555")
HFILL = PatternFill("solid", fgColor="1F2A44")
SFILL = PatternFill("solid", fgColor="E8ECF3")
YFILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

USD = '$#,##0;($#,##0);"-"'
USDM = '$#,##0.0,,"M";($#,##0.0,,"M");"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
X = '0.00"x"'

wb = Workbook()


def hdr(ws, row, labels, col=1, widths=None):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR, HFILL, CENTER, BOX
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[L(col + i)].width = w


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, bold=False, wrap=False, border=True):
    c = ws[ref]
    c.value = val
    c.font = Font(name=F, size=10, bold=bold, color=font.color) if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    if border:
        c.border = BOX
    return c


def title(ws, t, s):
    ws["A1"].value, ws["A1"].font = t, TITLE
    ws["A2"].value, ws["A2"].font = s, SUB
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------- Cover
ws = wb.active
ws.title = "Cover"
title(ws, "02 — Market Sizing: U.S. Vertical Opportunity for Clay",
      "Independent portfolio case study by Zhuoran (Zora) Wang · Not commissioned by or affiliated with Clay Labs Inc. · Research reference date: 2026-10-18")
rows = [
    ("Purpose", "Size and compare three candidate U.S. verticals bottom-up (unique firms × annual contract value), with explicit TAM / SAM / SOM definitions and sensitivity ranges."),
    ("TAM", "Annual revenue if every in-scope firm (20+ employees, segment NAICS) bought Clay at the size-band ACV. A theoretical ceiling, not a forecast."),
    ("SAM", "TAM firms filtered to those with (1) an active B2B outbound / territory sales motion and (2) a CRM/data stack Clay can integrate with today (Growth tier and above)."),
    ("SOM", "Year-3 ARR potential from SAM firms at a uniform 'fair-share' penetration rate. Phase 1 deliberately applies the SAME penetration to every segment; segment-specific right-to-win is assessed in Phases 2–3."),
    ("Unit of analysis", "Unique FIRMS (enterprises), never establishments. Firm counts are de-duplicated across sub-industries using Census-measured overlap (see Firm_Counts)."),
    ("Primary data", "2022 Economic Census — Establishment & Firm Size Statistics, Employment Size of Firms (EC2200SIZEEMPFIRM); BLS OEWS May 2025 industry-specific estimates. See Sources."),
    ("Not a forecast", "Nothing here is a Clay forecast or uses Clay internal data. ACV, readiness and penetration inputs are labeled benchmarks [B] or scenario assumptions [D]."),
    ("Color legend", "Blue text = hardcoded input · Black = formula · Green = formula linking another tab · Yellow fill = key judgment-based assumption."),
    ("Evidence classes", "A = verified public fact · B = external benchmark / third-party estimate · C = analyst estimate (derived) · D = scenario assumption · E = recommendation."),
    ("Tabs", "Executive_Summary → Assumptions → Firm_Counts → Segment_Profile → TAM_SAM_SOM → Sensitivity → Hypothesis_Tests → Segment_Defs → Raw_Census → Raw_OEWS → Sources"),
]
for i, (k, v) in enumerate(rows, start=4):
    put(ws, f"A{i}", k, bold=True, fill=SFILL)
    put(ws, f"B{i}", v, wrap=True)
    ws.row_dimensions[i].height = 42
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 120

# ---------------------------------------------------------------- Raw_Census
rc = wb.create_sheet("Raw_Census")
title(rc, "Raw data — 2022 Economic Census, Employment Size of Firms (EC2200SIZEEMPFIRM), United States",
      "Source S14 · Reference year 2022 · Firms 'operated for the entire year' (FY) by firm employment size · Receipts in $ thousands · D = suppressed by Census (disclosure) · blank = not published")
hdr(rc, 4, ["Key", "NAICS 2022", "Size band", "Firms", "Firm flag", "Establishments", "Employees", "Receipts ($000)", "Parent sector (for de-dup)", "File"],
    widths=[16, 11, 11, 12, 9, 14, 13, 18, 14, 34])
parent_of = {}
for c in ["4231", "4232", "4233", "4234", "4235", "4236", "4237", "4238", "4239"]:
    parent_of[c] = "423"
for c in ["311", "312", "313", "314", "315", "316", "321", "322", "323", "324", "325", "326", "327", "331",
          "332", "333", "334", "335", "336", "337", "339"]:
    parent_of[c] = "31-33"
for c in ["5411", "5412", "5413", "5414", "5415", "5416", "5417", "5418", "5419"]:
    parent_of[c] = "54"
r = 5
for fname, cols in [("appendix/raw_census_ecnsize_2022_firm_size.csv", "main"), ("appendix/overlap_raw.csv", "ovl")]:
    for row in csv.DictReader(open(f"{BASE}/{fname}")):
        n = row.get("naics2022") or row.get("naics")
        b = row.get("size_band") or row.get("band")
        fl = row.get("firm_flag", row.get("flag", ""))
        est = row.get("establishments") or row.get("estab") or ""
        emp = row.get("employees", "")
        rcp = row.get("receipts_usd_thousands", "")
        vals = [f'=B{r}&"|"&C{r}', n, b,
                int(row["firms"]) if row["firms"] else None, fl or None,
                int(est) if est else None, int(emp) if emp else None, int(rcp) if rcp else None,
                parent_of.get(n), fname.split("/")[-1]]
        for j, v in enumerate(vals, start=1):
            cell = rc.cell(row=r, column=j, value=v)
            cell.font = BLACK if j == 1 else BLUE
            if j in (4, 6, 7, 8):
                cell.number_format = NUM
        r += 1
RC_LAST = r - 1
rc.freeze_panes = "A5"
RCK, RCN, RCB, RCF, RCFL, RCE, RCEMP, RCR, RCP = [f"Raw_Census!${c}$5:${c}${RC_LAST}" for c in "ABCDEFGHI"]

# ---------------------------------------------------------------- Raw_OEWS
ro = wb.create_sheet("Raw_OEWS")
title(ro, "Raw data — BLS Occupational Employment and Wage Statistics (OEWS), May 2025, national industry-specific estimates",
      "Source S15 · Employment counts by occupation within industry · RSE = relative standard error (%) · suppressed estimates left blank")
hdr(ro, 4, ["Key", "BLS industry code", "Industry", "Segment proxy", "SOC code", "Occupation", "Employment", "Emp. RSE (%)"],
    widths=[18, 14, 46, 16, 10, 44, 13, 12])
ind_name = {"000000": "Cross-industry (all)", "4230A1": "Merchant Wholesalers, Durable Goods (4232/3/5/6/7/9 only)",
            "423800": "Machinery, Equipment & Supplies Merchant Wholesalers", "326000": "Plastics & Rubber Products Mfg",
            "332000": "Fabricated Metal Product Mfg", "333000": "Machinery Mfg", "335000": "Electrical Equipment, Appliance & Component Mfg",
            "541000": "Professional, Scientific & Technical Services", "541100": "Legal Services", "541200": "Accounting, Tax Prep, Bookkeeping & Payroll",
            "541300": "Architectural, Engineering & Related", "541500": "Computer Systems Design & Related", "541600": "Management, Scientific & Technical Consulting",
            "541800": "Advertising, PR & Related", "513200": "Software Publishers", "561300": "Employment Services"}
seg_proxy = {"4230A1": "Distribution", "423800": "Distribution", "326000": "Manufacturing", "332000": "Manufacturing",
             "333000": "Manufacturing", "335000": "Manufacturing", "541300": "ProfServices", "541500": "ProfServices",
             "541600": "ProfServices", "541800": "ProfServices", "513200": "Software", "000000": "AllIndustries"}
occ_name = {"00-0000": "All occupations", "41-0000": "Sales & related occupations (major group)",
            "41-4011": "Sales reps, wholesale & mfg — technical & scientific products",
            "41-4012": "Sales reps, wholesale & mfg — except technical & scientific",
            "11-2021": "Marketing managers", "11-2022": "Sales managers", "13-1161": "Market research analysts & marketing specialists",
            "15-0000": "Computer & mathematical occupations (major group)", "43-4051": "Customer service representatives"}
r = 5
for row in csv.DictReader(open(f"{BASE}/appendix/raw_bls_oews_may2025_selected.csv")):
    ind, oc = row["bls_industry"], row["occ_code"]
    vals = [f'=B{r}&"|"&E{r}', ind, ind_name[ind], seg_proxy.get(ind, "(not used)"), oc, occ_name[oc],
            int(row["employment"]) if row["employment"] else None,
            float(row["emp_rse_pct"]) if row["emp_rse_pct"] not in ("", "suppressed") else None]
    for j, v in enumerate(vals, start=1):
        cell = ro.cell(row=r, column=j, value=v)
        cell.font = BLACK if j == 1 else BLUE
        if j == 7:
            cell.number_format = NUM
    r += 1
RO_LAST = r - 1
ro.freeze_panes = "A5"
ROSEG, ROOCC, ROEMP = [f"Raw_OEWS!${c}$5:${c}${RO_LAST}" for c in "DEG"]

# ---------------------------------------------------------------- Raw_TYPOP (review M3)
rt = wb.create_sheet("Raw_TYPOP")
title(rt, "Raw data — 2022 Economic Census, Wholesale Trade: Detailed Type of Operation (EC2242TYPOP), United States",
      "Source S58 · establishments, employees and sales by type of operation · TYPOP 21 sales branches (with stock) and 22 sales offices (without stock) = manufacturers' sales branches and offices (MSBOs)")
hdr(rt, 4, ["Key", "NAICS 2022", "Industry", "TYPOP", "Type of operation", "Establishments", "Employees", "Sales ($000)", "MSBO?"],
    widths=[12, 11, 40, 8, 40, 14, 13, 16, 8])
r = 5
for row in csv.DictReader(open(f"{BASE}/appendix/raw_census_ec2242typop_2022_423.csv")):
    vals = [f'=B{r}&"|"&D{r}', row["NAICS2022"], row["NAICS2022_LABEL"], row["TYPOP"], row["TYPOP_LABEL"],
            int(row["ESTAB"]) if row["ESTAB"] else None, int(row["EMP"]) if row["EMP"] else None,
            int(row["RCPTOT"]) if row["RCPTOT"] else None, "Y" if row["TYPOP"] in ("21", "22") else "N"]
    for j, v in enumerate(vals, start=1):
        cell = rt.cell(row=r, column=j, value=v)
        cell.font = BLACK if j == 1 else BLUE
        if j in (6, 7, 8):
            cell.number_format = NUM
    r += 1
RT_LAST = r - 1
rt.freeze_panes = "A5"
RTN, RTE, RTEMP, RTM = [f"Raw_TYPOP!${c}$5:${c}${RT_LAST}" for c in "BFGI"]

# ---------------------------------------------------------------- Segment_Defs
sd = wb.create_sheet("Segment_Defs")
title(sd, "Segment definitions (refined in Phase 1)", "NAICS 2022 codes · 'Core' codes drive the model; 'Broad' codes are used only in the definition sensitivity (Sensitivity tab, block C)")
hdr(sd, 4, ["Segment", "NAICS", "Industry", "Scope", "Rationale"], widths=[22, 9, 52, 9, 90])
defs = [
    ("A. Industrial Distribution", "4235", "Metal & Mineral (except Petroleum) Merchant Wholesalers", "Core", "Industrial MRO / metals service centers; outside-sales-led; serve manufacturers and contractors."),
    ("A. Industrial Distribution", "4236", "Household Appliances & Electrical/Electronic Goods Merchant Wholesalers", "Core", "Electrical distributors (e.g., the Regency Supply case, S07); contractor-facing territory sales."),
    ("A. Industrial Distribution", "4237", "Hardware, Plumbing & Heating Equipment Merchant Wholesalers", "Core", "PHVAC / hardware distributors; branch-heavy networks serving contractors."),
    ("A. Industrial Distribution", "4238", "Machinery, Equipment & Supplies Merchant Wholesalers", "Core", "Industrial machinery, equipment and supplies distributors; largest core sub-segment."),
    ("A. Industrial Distribution", "4233", "Lumber & Other Construction Materials Merchant Wholesalers", "Broad", "Building-materials distribution; adjacent but more residential/contractor commodity-driven."),
    ("A. Industrial Distribution", "4234", "Professional & Commercial Equipment Merchant Wholesalers", "Broad", "Includes computers, medical and office equipment — mixed, partly tech-channel; excluded from core."),
    ("B. Industrial Manufacturing", "326", "Plastics & Rubber Products Manufacturing", "Core", "Component suppliers selling B2B to OEMs; engineered-to-order sales."),
    ("B. Industrial Manufacturing", "332", "Fabricated Metal Product Manufacturing", "Core", "Job shops and fabricators; largest firm count; many small, owner-led firms."),
    ("B. Industrial Manufacturing", "333", "Machinery Manufacturing", "Core", "Capital equipment OEMs; direct + dealer/distributor channels."),
    ("B. Industrial Manufacturing", "335", "Electrical Equipment, Appliance & Component Manufacturing", "Core", "Electrical equipment/components; often sell through electrical distributors (link to Segment A)."),
    ("B. Industrial Manufacturing", "31-33", "All Manufacturing (sector)", "Broad", "Original brief scope. Rejected as core: includes food, beverage, apparel, printing etc. with retail/CPG go-to-market (tests H1)."),
    ("C. B2B Professional Services", "5413", "Architectural, Engineering & Related Services", "Core", "B2B project firms; relationship + RFP-driven business development."),
    ("C. B2B Professional Services", "5415", "Computer Systems Design & Related Services", "Core", "IT services / systems integrators; technically sophisticated, outbound-capable."),
    ("C. B2B Professional Services", "5416", "Management, Scientific & Technical Consulting", "Core", "Consulting firms; partner-led BD with growing marketing/outbound functions."),
    ("C. B2B Professional Services", "5418", "Advertising, PR & Related Services", "Core", "Agencies; heavy outbound/new-business culture; overlaps Clay's agency ecosystem."),
    ("C. B2B Professional Services", "5411", "Legal Services", "Broad", "Referral/relationship-driven; bar advertising rules; excluded from core."),
    ("C. B2B Professional Services", "5412", "Accounting, Tax Prep, Bookkeeping & Payroll", "Broad", "Large consumer/very-small-business tax-prep component; excluded from core."),
    ("Reference (not scored)", "5132", "Software Publishers", "Benchmark", "Proxy for Clay's current core customer base; used only for comparison of density metrics."),
    ("Adjacency (not scored)", "5613", "Employment Services (staffing)", "Adjacency", "Outbound-heavy B2B service; noted for future screening, outside the brief's three segments."),
]
for i, d in enumerate(defs, start=5):
    for j, v in enumerate(d, start=1):
        c = put(sd, f"{L(j)}{i}", v, font=BLUE if j == 2 else BLACK, wrap=(j == 5))
put(sd, f"A{5+len(defs)+1}", "Size bands (firm employment, all U.S. establishments of the firm): SMB = 20–99 · Mid-Market = 100–499 · Enterprise = 500+. Firms with <20 employees are excluded from TAM (rarely staff a dedicated sales/RevOps function; Clay CRM sync requires the Growth tier).", wrap=True, border=False)
sd.merge_cells(start_row=5 + len(defs) + 1, start_column=1, end_row=5 + len(defs) + 1, end_column=5)
sd.row_dimensions[5 + len(defs) + 1].height = 30

# ---------------------------------------------------------------- Assumptions
asm = wb.create_sheet("Assumptions")
title(asm, "Assumptions — every judgment-based input in one place (Low / Base / High)",
      "Low and High are compounded in the scenario outer bounds on TAM_SAM_SOM; one-at-a-time swings are on Sensitivity. Edit blue cells only.")
hdr(asm, 4, ["ID", "Parameter", "Applies to", "Low", "Base", "High", "Unit", "Class", "Source / rationale"],
    widths=[7, 40, 22, 11, 11, 11, 10, 7, 110])
A = {}  # name -> row
arows = [
    # id, name, applies, low, base, high, unit, class, rationale, key
    ("A1-D", "ACV — A. Industrial Distribution", "SMB (20–99)", 5352, 5352, 5352, "$/yr", "C/D",
     "Usage-based ACV at published Growth-tier prices (S01): sellers per firm (employees × 22.0% sales share, S15) × 150 accounts × Clay's Data-Credit and Actions meters; Low/Base/High = usage 0.75×/1.0×/1.5×. SMB usage sits below the Growth floor ($5,352). Enterprise = MAX(Enterprise minimum, usage) less 30% discount (S16). Built in 07 › ACV Build (v2).", "acv_d_smb"),
    ("A2-D", "ACV — A. Industrial Distribution", "Mid-Market (100–499)", 7310, 9747, 14621, "$/yr", "C/D",
     "See SMB row for method.", "acv_d_mm"),
    ("A3-D", "ACV — A. Industrial Distribution", "Enterprise (500+)", 46241, 61655, 92482, "$/yr", "C/D",
     "See SMB row for method.", "acv_d_ent"),
    ("A1-M", "ACV — B. Industrial Manufacturing", "SMB (20–99)", 5352, 5352, 5352, "$/yr", "C/D",
     "Same usage build with 2.7% sales share (S15): too few sellers per firm to exceed the Growth floor in SMB and Mid-Market; Enterprise at the Enterprise minimum less discount. Identical in Low/Base/High because usage stays below the floors.", "acv_m_smb"),
    ("A2-M", "ACV — B. Industrial Manufacturing", "Mid-Market (100–499)", 5352, 5352, 5352, "$/yr", "C/D",
     "See SMB row for method.", "acv_m_mm"),
    ("A3-M", "ACV — B. Industrial Manufacturing", "Enterprise (500+)", 44310, 44310, 44310, "$/yr", "C/D",
     "See SMB row for method.", "acv_m_ent"),
    ("A1-P", "ACV — C. B2B Professional Services", "SMB (20–99)", 5352, 5352, 5352, "$/yr", "C/D",
     "Same usage build with 6.2% sales share (S15): SMB and Mid-Market at the Growth floor; Enterprise at the Enterprise minimum less discount. Identical in Low/Base/High because usage stays below the floors.", "acv_p_smb"),
    ("A2-P", "ACV — C. B2B Professional Services", "Mid-Market (100–499)", 5352, 5352, 5352, "$/yr", "C/D",
     "See SMB row for method.", "acv_p_mm"),
    ("A3-P", "ACV — C. B2B Professional Services", "Enterprise (500+)", 44310, 44310, 44310, "$/yr", "C/D",
     "See SMB row for method.", "acv_p_ent"),
    ("A4", "SAM filter 1 — active B2B outbound / territory sales motion", "A. Industrial Distribution", 0.65, 0.80, 0.90, "% firms", "D",
     "Distribution is the most sales-intensive segment: sales occupations = ~22% of employment vs ~8.6% economy-wide (≈2.5×) (OEWS May 2025, S15 [A]; see Segment_Profile).", "f1_dist"),
    ("A5", "SAM filter 1 — active B2B outbound / territory sales motion", "B. Industrial Manufacturing", 0.30, 0.45, 0.60, "% firms", "D",
     "Sales occupations only ~2.7% of manufacturing employment (S15 [A]) — consistent with many manufacturers selling through distributors and independent reps rather than direct outbound teams.", "f1_mfg"),
    ("A6", "SAM filter 1 — active B2B outbound / territory sales motion", "C. B2B Professional Services", 0.40, 0.55, 0.70, "% firms", "D",
     "Sales occupations ~6.2% of employment (S15 [A]), but BD is often done by billable 'seller-doers' (partners/principals) not counted as sales staff; A&E skews RFP-driven.", "f1_ps"),
    ("A7", "SAM filter 2 — CRM / data stack Clay can integrate with", "A. Industrial Distribution", 0.50, 0.56, 0.65, "% firms", "B/D",
     "v2: Base = 56% of distributors use CRM (Phocas/DSG State of Sales in Distribution 2026, S25 [B]) — the newest evidence. Range from DSG 2022 (S17 [B]): 52% general-purpose CRM; 15% none. v1 used 65% Base.", "f2_dist"),
    ("A8", "SAM filter 2 — CRM / data stack Clay can integrate with", "B. Industrial Manufacturing", 0.45, 0.60, 0.75, "% firms", "D",
     "No reliable public adoption rate found (a 2023 vendor-sponsored survey reported usage frequency, not adoption — not used). Set slightly below distribution given lower sales-function density. KEY DATA GAP.", "f2_mfg"),
    ("A9", "SAM filter 2 — CRM / data stack Clay can integrate with", "C. B2B Professional Services", 0.55, 0.70, 0.85, "% firms", "D",
     "No reliable public adoption rate found. Set above distribution given far higher technical workforce share (computer occupations ~24% of employment vs ~2% in distribution; S15 [A]). KEY DATA GAP.", "f2_ps"),
    ("A10", "3-year penetration of SAM firms (uniform across segments)", "SMB (20–99)", 0.005, 0.015, 0.03, "% SAM firms", "D",
     "Fair-share assumption for a focused vertical program over three years. Deliberately identical across segments in Phase 1. Sanity check vs Clay's 17,000+ total customers on Executive_Summary.", "pen_smb"),
    ("A11", "3-year penetration of SAM firms (uniform across segments)", "Mid-Market (100–499)", 0.01, 0.03, 0.06, "% SAM firms", "D",
     "Higher than SMB: sales-assisted motion concentrates effort on fewer, larger accounts.", "pen_mm"),
    ("A12", "3-year penetration of SAM firms (uniform across segments)", "Enterprise (500+)", 0.02, 0.06, 0.12, "% SAM firms", "D",
     "Highest: named-account coverage; Clay reports >200% enterprise NRR (S02 [A]) indicating strong land-and-expand once in.", "pen_ent"),
]
for i, a in enumerate(arows, start=5):
    aid, name, app, lo, ba, hi, unit, cls, why, key = a
    A[key] = i
    put(asm, f"A{i}", aid)
    put(asm, f"B{i}", name, wrap=True)
    put(asm, f"C{i}", app, wrap=True)
    fmt = USD if unit == "$/yr" else PCT
    for col, v in zip("DEF", (lo, ba, hi)):
        put(asm, f"{col}{i}", v, font=BLUE, fmt=fmt, fill=YFILL if cls.startswith("D") or "D" in cls else None)
    put(asm, f"G{i}", unit)
    put(asm, f"H{i}", cls)
    put(asm, f"I{i}", why, wrap=True)
    asm.row_dimensions[i].height = 44
asm.freeze_panes = "D5"
n = 5 + len(arows) + 1
put(asm, f"A{n}", "v2 (2026-10-19, partner review M1): ACV is now segment-specific and usage-based at list prices (07 v2 ACV Build) instead of a uniform $7.5K / $20K / $60K. Penetration remains uniform across segments by design, so SOM is an assumption-driven sizing, not a finding about right-to-win.", wrap=True, border=False)
asm.merge_cells(f"A{n}:I{n}")
asm.row_dimensions[n].height = 32


def aref(key, col):  # col: D low, E base, F high
    return f"Assumptions!${col}${A[key]}"


# ---------------------------------------------------------------- Firm_Counts
fc = wb.create_sheet("Firm_Counts")
title(fc, "Firm counts — unique U.S. firms by segment and firm size (2022)",
      "Firms operated for the entire year. Suppressed cells imputed (block 2). Sub-industry totals de-duplicated with Census-measured overlap (block 3). Establishments shown for contrast only — never used as the counting unit.")
bands = ["20-49", "50-99", "100-249", "250-499", "500+"]
hdr(fc, 4, ["Segment", "NAICS", "Industry", "20-49", "50-99", "100-249", "250-499", "500+", "Firms 20+",
            "Firms <20", "All FY firms", "FY establishments", "Estab. per firm"],
    widths=[26, 8, 44, 10, 10, 10, 10, 10, 11, 11, 12, 14, 11])
core = [("A. Industrial Distribution", ["4235", "4236", "4237", "4238"], "423"),
        ("B. Industrial Manufacturing", ["326", "332", "333", "335"], "31-33"),
        ("C. B2B Professional Services", ["5413", "5415", "5416", "5418"], "54")]
labels = {d[1]: d[2] for d in defs}
r = 5
code_row = {}
seg_rows = {}
IMP_START = 30  # imputation block row (set below)
for seg, codes, parent in core:
    start = r
    for cd in codes:
        code_row[cd] = r
        put(fc, f"A{r}", seg)
        put(fc, f"B{r}", cd, font=BLUE)
        put(fc, f"C{r}", labels[cd])
        for j, b in enumerate(bands):
            col = L(4 + j)
            if cd == "4235" and b in ("50-99", "100-249", "250-499"):
                put(fc, f"{col}{r}", None, fmt=NUM)  # filled after imputation block known
            else:
                put(fc, f"{col}{r}", f'=SUMIFS({RCF},{RCK},$B{r}&"|"&{col}$4)', font=GREEN, fmt=NUM)
        put(fc, f"I{r}", f"=SUM(D{r}:H{r})", fmt=NUM)
        put(fc, f"J{r}", f'=SUMIFS({RCF},{RCK},$B{r}&"|0-4")+SUMIFS({RCF},{RCK},$B{r}&"|5-9")+SUMIFS({RCF},{RCK},$B{r}&"|10-19")', font=GREEN, fmt=NUM)
        put(fc, f"K{r}", f'=SUMIFS({RCF},{RCK},$B{r}&"|FY")', font=GREEN, fmt=NUM)
        put(fc, f"L{r}", f'=SUMIFS({RCE},{RCK},$B{r}&"|FY")', font=GREEN, fmt=NUM)
        put(fc, f"M{r}", f"=IFERROR(L{r}/K{r},0)", fmt=X)
        r += 1
    put(fc, f"A{r}", f"{seg} — subtotal (before de-dup)", bold=True, fill=SFILL)
    for col in "BC":
        put(fc, f"{col}{r}", None, fill=SFILL)
    for col in "DEFGHIJKL":
        put(fc, f"{col}{r}", f"=SUM({col}{start}:{col}{r-1})", fmt=NUM, fill=SFILL, bold=True)
    put(fc, f"M{r}", f"=IFERROR(L{r}/K{r},0)", fmt=X, fill=SFILL, bold=True)
    seg_rows[seg] = (start, r, parent)
    r += 2

# Block 2: imputation for 4235
r += 0
imp = r
put(fc, f"A{imp}", "Block 2 — Imputation of Census-suppressed cells (NAICS 4235, firms with 50–499 employees)", bold=True, border=False)
put(fc, f"A{imp+1}", "4235 firms operated full year (FY)")
put(fc, f"D{imp+1}", f'=SUMIFS({RCF},{RCK},"4235|FY")', font=GREEN, fmt=NUM)
put(fc, f"A{imp+2}", "less: reported bands (0–4, 5–9, 10–19, 20–49, 500+)")
put(fc, f"D{imp+2}", '=' + "+".join([f'SUMIFS({RCF},{RCK},"4235|{b}")' for b in ["0-4", "5-9", "10-19", "20-49", "500+"]]), font=GREEN, fmt=NUM)
put(fc, f"A{imp+3}", "Suppressed residual (firms with 50–499 employees) [C]", bold=True)
put(fc, f"D{imp+3}", f"=D{imp+1}-D{imp+2}", fmt=NUM, bold=True)
put(fc, f"A{imp+4}", "Allocation key: parent NAICS 423 firm mix across 50–99 / 100–249 / 250–499")
for j, b in enumerate(["50-99", "100-249", "250-499"]):
    col = L(5 + j)
    put(fc, f"{col}{imp+3}", b, bold=True, fill=SFILL)
    put(fc, f"{col}{imp+4}", f'=SUMIFS({RCF},{RCK},"423|{b}")/(SUMIFS({RCF},{RCK},"423|50-99")+SUMIFS({RCF},{RCK},"423|100-249")+SUMIFS({RCF},{RCK},"423|250-499"))', font=GREEN, fmt=PCT)
    put(fc, f"{col}{imp+5}", f"=$D${imp+3}*{col}{imp+4}", fmt=NUM1)
    fc[f"{col}{code_row['4235']}"].value = f"={col}{imp+5}"
    fc[f"{col}{code_row['4235']}"].font = BLACK
    fc[f"{col}{code_row['4235']}"].fill = YFILL
    fc[f"{col}{code_row['4235']}"].comment = Comment("Imputed [C]: Census suppressed this cell (flag D). Value = 4235 residual × NAICS 423 band mix. See Block 2.", "ZW")
put(fc, f"A{imp+5}", "Imputed 4235 firms by band (feeds yellow cells above)")

# Block 3: de-dup factors
dd = imp + 8
put(fc, f"A{dd}", "Block 3 — De-duplication: a firm with establishments in several sub-industries is counted in each. Factor = MIN(1, sector firms ÷ Σ reported sub-industry firms), per band.", bold=True, border=False)
hdr(fc, dd + 1, ["Parent sector", "", "Measure", "20-49", "50-99", "100-249", "250-499", "500+"])
fac_row = {}
rr = dd + 2
for seg, codes, parent in core:
    put(fc, f"A{rr}", parent, font=BLUE)
    put(fc, f"C{rr}", "Sector firms (parent)")
    put(fc, f"A{rr+1}", parent, font=BLUE)
    put(fc, f"C{rr+1}", "Σ reported sub-industry firms")
    put(fc, f"A{rr+2}", parent, font=BLUE)
    put(fc, f"C{rr+2}", "Implied overlap")
    put(fc, f"A{rr+3}", parent, font=BLUE, bold=True)
    put(fc, f"C{rr+3}", "De-dup factor applied [C]", bold=True)
    for j, b in enumerate(bands):
        col = L(4 + j)
        put(fc, f"{col}{rr}", f'=SUMIFS({RCF},{RCK},$A{rr}&"|"&{col}${dd+1})', font=GREEN, fmt=NUM)
        put(fc, f"{col}{rr+1}", f'=SUMIFS({RCF},{RCP},$A{rr+1},{RCB},{col}${dd+1})', font=GREEN, fmt=NUM)
        put(fc, f"{col}{rr+2}", f"=IFERROR({col}{rr+1}/{col}{rr}-1,0)", fmt=PCT)
        put(fc, f"{col}{rr+3}", f"=IFERROR(MIN(1,{col}{rr}/{col}{rr+1}),1)", fmt='0.000', bold=True)
    fac_row[parent] = rr + 3
    rr += 5
put(fc, f"A{rr}", "Reading: overlap is small in distribution (~1–3%) and professional services (~1–5%) but material in mid-size/large manufacturing (~7–11%), where multi-plant firms span subsectors. Where some sub-industry cells are suppressed, the reported sum is a lower bound, so the factor is conservative (removes less overlap than exists).", wrap=True, border=False)
fc.merge_cells(f"A{rr}:M{rr}")
fc.row_dimensions[rr].height = 30

# Block 4: unique firms by segment x band, and size groups
uq = rr + 2
put(fc, f"A{uq}", "Block 4 — Unique firms by segment (de-duplicated) [C]", bold=True, border=False)
hdr(fc, uq + 1, ["Segment", "", "", "20-49", "50-99", "100-249", "250-499", "500+", "Firms 20+"])
uq_row = {}
rr = uq + 2
for seg, codes, parent in core:
    s, sub, _ = seg_rows[seg]
    put(fc, f"A{rr}", seg, bold=True)
    for j, b in enumerate(bands):
        col = L(4 + j)
        put(fc, f"{col}{rr}", f"={col}{sub}*{col}${fac_row[parent]}", fmt=NUM)
    put(fc, f"I{rr}", f"=SUM(D{rr}:H{rr})", fmt=NUM, bold=True)
    uq_row[seg] = rr
    rr += 1

grp = rr + 1
put(fc, f"A{grp}", "Block 5 — Unique firms by size group (feeds TAM_SAM_SOM and Sensitivity)", bold=True, border=False)
hdr(fc, grp + 1, ["Size group", "", "Bands", "A. Industrial Distribution", "B. Industrial Manufacturing", "C. B2B Professional Services"])
grp_row = {}
groups = [("SMB", "20–99", ["D", "E"]), ("Mid-Market", "100–499", ["F", "G"]), ("Enterprise", "500+", ["H"])]
for k, (g, rng, cols) in enumerate(groups):
    row = grp + 2 + k
    put(fc, f"A{row}", g, bold=True)
    put(fc, f"C{row}", rng)
    for j, (seg, _, _) in enumerate(core):
        col = L(4 + j)
        put(fc, f"{col}{row}", "=" + "+".join(f"{c}{uq_row[seg]}" for c in cols), fmt=NUM)
    grp_row[g] = row
tot = grp + 5
put(fc, f"A{tot}", "Total firms 20+", bold=True, fill=SFILL)
for col in "BC":
    put(fc, f"{col}{tot}", None, fill=SFILL)
for j in range(3):
    col = L(4 + j)
    put(fc, f"{col}{tot}", f"=SUM({col}{grp+2}:{col}{grp+4})", fmt=NUM, bold=True, fill=SFILL)
fc.freeze_panes = "D5"
FC_GRP = {g: grp_row[g] for g in grp_row}
FC_SEGCOL = {core[0][0]: "D", core[1][0]: "E", core[2][0]: "F"}
FC_TOT = tot

# Block 6: establishments vs unique firms (contrast only; review m8)
eb = tot + 3
put(fc, f"A{eb}", "Block 6 — Establishments vs unique firms, distribution (contrast only; never a counting unit) [C]", bold=True, border=False)
hdr(fc, eb + 1, ["Measure", "", "", "Value", "Note"])
dcodes = ["4235", "4236", "4237", "4238"]
sumk = lambda rng, keys: "+".join(f'SUMIFS({rng},{RCK},"{k}")' for k in keys)
rows6 = [
    ("Establishments, all sizes (FY), NAICS 4235–4238", "=" + sumk(RCE, [f"{c}|FY" for c in dcodes]), NUM,
     "Each establishment has one NAICS code, so establishments are additive."),
    ("Firms, all sizes (FY), Σ sub-industries (not additive)", "=" + sumk(RCF, [f"{c}|FY" for c in dcodes]), NUM,
     "A multi-industry firm is counted in each sub-industry."),
    ("All-size de-dup factor, NAICS 423 (sector FY firms ÷ Σ sub-industry FY firms)",
     f'=MIN(1,SUMIFS({RCF},{RCK},"423|FY")/SUMIFS({RCF},{RCP},"423",{RCB},"FY"))', '0.000',
     "Same method as Block 3, applied to the all-size total."),
    ("Unique firms, all sizes [C]", f"=D{eb+3}*D{eb+4}", NUM, "De-duplicated."),
    ("Establishments per unique firm, all sizes", f"=D{eb+2}/D{eb+5}", '0.00"x"', "Counting establishments would overstate firms by this ratio − 1."),
    ("Establishments of firms with 20+ employees, 4235–4238 [C]",
     "=" + sumk(RCE, [f"{c}|{b}" for c in dcodes[1:] for b in ["20-49", "50-99", "100-249", "250-499", "500+"]])
     + "+" + sumk(RCE, ["4235|FY"]) + "-(" + sumk(RCE, ["4235|0-4", "4235|5-9", "4235|10-19"]) + ")",
     NUM, "4235 establishments at 50–499 are suppressed; recovered as FY total less the reported small bands [C]."),
    ("Unique firms with 20+ employees (Block 5)", f"=D{FC_TOT}", NUM, "The counting unit used in TAM/SAM/SOM."),
    ("Establishments per unique firm, 20+ employees", f"=D{eb+7}/D{eb+8}", '0.00"x"', "The relevant contrast for the in-scope (20+) universe."),
]
for i, (lab, f, fmt, note) in enumerate(rows6):
    rr = eb + 2 + i
    put(fc, f"A{rr}", lab, bold=(i in (4, 7)))
    fc.merge_cells(f"A{rr}:C{rr}")
    put(fc, f"D{rr}", f, fmt=fmt, bold=(i in (4, 7)))
    put(fc, f"E{rr}", note)
    fc.merge_cells(f"E{rr}:M{rr}")
# Block 7: manufacturers' sales branches and offices inside NAICS 423 (review M3)
mb = eb + 2 + len(rows6) + 2
put(fc, f"A{mb}", "Block 7 — Manufacturers' sales branches and offices (MSBOs) inside 4235–4238: how many 500+ 'distributors' may be manufacturers' sales arms [C]", bold=True, border=False)
hdr(fc, mb + 1, ["Measure", "", "", "Low", "High", "Note"])
dc4 = ["4235", "4236", "4237", "4238"]
sumt = lambda rng, extra="": "+".join(f'SUMIFS({rng},{RTN},"{c}"{extra})' for c in dc4)
rows7 = [
    ("MSBO establishments (TYPOP 21 + 22)", "=" + sumt(RTE, f',{RTM},"Y"'), None, NUM, "Census type of operation, all firm sizes (S58)."),
    ("All establishments, all types of operation", "=" + sumt(RTE), None, NUM, "Equals the size-table total including part-year firms."),
    ("Establishments of 500+ employee firms", "=" + sumk(RCE, [f"{c}|500+" for c in dc4]), None, NUM, "Firm-size table (S14)."),
    ("MSBO share of the 500+ band — bound", f"=D{mb + 2}/D{mb + 3}", f"=MIN(1,D{mb + 2}/D{mb + 4})", PCT,
     "Low: MSBOs spread like all establishments. High: every MSBO belongs to a 500+ firm. Census does not cross type of operation with firm size."),
    ("Unique enterprise (500+) firms, as sized", f"=D{FC_TOT - 1}", f"=D{FC_TOT - 1}", NUM, "Block 5."),
    ("Enterprise firms excluding MSBOs [C]", f"=D{mb + 6}*(1-D{mb + 5})", f"=E{mb + 6}*(1-E{mb + 5})", NUM, "Applies the establishment share to firms (no firm-level split exists)."),
    ("TAM effect, Base enterprise ACV held ($)", f"=-(D{mb + 6}-D{mb + 7})*{aref('acv_d_ent', 'E')}", f"=-(E{mb + 6}-E{mb + 7})*{aref('acv_d_ent', 'E')}", USD, "Before any change in enterprise ACV."),
    ("SAM effect ($)", f"=D{mb + 8}*{aref('f1_dist', 'E')}*{aref('f2_dist', 'E')}", f"=E{mb + 8}*{aref('f1_dist', 'E')}*{aref('f2_dist', 'E')}", USD, "Same filters as TAM_SAM_SOM, Base."),
]
for i, (lab, lo, hi, fmt, note) in enumerate(rows7):
    rr = mb + 2 + i
    put(fc, f"A{rr}", lab, bold=(i in (3, 5)))
    fc.merge_cells(f"A{rr}:C{rr}")
    put(fc, f"D{rr}", lo, fmt=fmt, bold=(i in (3, 5)))
    put(fc, f"E{rr}", hi if hi is not None else None, fmt=fmt, bold=(i in (3, 5)))
    put(fc, f"F{rr}", note)
    fc.merge_cells(f"F{rr}:M{rr}")
put(fc, f"A{mb + 2 + len(rows7)}", "Reading: Base keeps Census as published (NAICS 423 includes MSBOs by definition). Removing them is a sensitivity: on the v2 engine (appendix/msbo_sensitivity.py) it moves the distribution lean-pilot Base NPV by −$0.01M to −$0.03M and the solution play's Upside advantage over Option 0 from +$0.36M to +$0.18–0.27M; no conclusion changes.", wrap=True, border=False)
fc.merge_cells(f"A{mb + 2 + len(rows7)}:M{mb + 2 + len(rows7)}")
fc.row_dimensions[mb + 2 + len(rows7)].height = 30
FC_MSBO = {"share": mb + 5, "ent": mb + 6, "ent_adj": mb + 7, "tam": mb + 8, "sam": mb + 9}
FC_EST = {"est_all": eb + 2, "firms_all_raw": eb + 3, "firms_all_uq": eb + 5, "ratio_all": eb + 6,
          "est_20": eb + 7, "firms_20": eb + 8, "ratio_20": eb + 9}

# ---------------------------------------------------------------- Segment_Profile
sp = wb.create_sheet("Segment_Profile")
title(sp, "Segment profile — structure, sales intensity and 'builder capacity'",
      "Census metrics use the core NAICS codes; OEWS metrics use the closest BLS industry groupings (distribution proxy 4230A1 also includes 4232/4233/4239). Ratios are [C] derived from [A] data.")
hdr(sp, 4, ["Metric", "A. Industrial Distribution", "B. Industrial Manufacturing", "C. B2B Professional Services",
            "Ref: Software Publishers (Clay core proxy)", "Ref: All industries", "Unit / source", "What it indicates"],
    widths=[44, 16, 16, 16, 18, 14, 22, 70])
segkeys = ["Distribution", "Manufacturing", "ProfServices", "Software", "AllIndustries"]
codes_by = {"Distribution": ["4235", "4236", "4237", "4238"], "Manufacturing": ["326", "332", "333", "335"],
            "ProfServices": ["5413", "5415", "5416", "5418"], "Software": ["5132"]}


def census_sum(codes, field, bandlist):
    return "+".join(f'SUMIFS({field},{RCK},"{c}|{b}")' for c in codes for b in bandlist)


big = ["20-49", "50-99", "100-249", "250-499", "500+"]
metrics = []
# row builders
r = 5
def mrow(label, cells, fmt, unit, meaning, bold=False):
    global r
    put(sp, f"A{r}", label, bold=bold)
    for j, v in enumerate(cells):
        put(sp, f"{L(2+j)}{r}", v, fmt=fmt, font=GREEN if (isinstance(v, str) and "!" in v) else BLACK)
    put(sp, f"G{r}", unit, wrap=True)
    put(sp, f"H{r}", meaning, wrap=True)
    sp.row_dimensions[r].height = 30
    r += 1
    return r - 1

put(sp, f"A{r}", "Structure (2022 Economic Census, S14)", bold=True, fill=SFILL); r += 1
row_f20 = mrow("Unique firms, 20+ employees", [f"=Firm_Counts!D{FC_TOT}", f"=Firm_Counts!E{FC_TOT}", f"=Firm_Counts!F{FC_TOT}",
                                               "=" + census_sum(["5132"], RCF, big), None], NUM, "firms [C]", "Size of the reachable firm universe (de-duplicated).")
mrow("Establishments per firm (all FY firms)", [
    "=(" + census_sum(codes_by[s], RCE, ["FY"]) + ")/(" + census_sum(codes_by[s], RCF, ["FY"]) + ")" for s in segkeys[:4]] + [None],
    X, "ratio [C]", "Branch intensity. >1.5 means multi-branch networks → territory planning and branch-level account mapping matter.")
mrow("Avg receipts per firm, 20+ firms ($M)", [
    "=(" + census_sum(codes_by[s], RCR, big) + ")/(" + "+".join(f'IF(SUMIFS({RCR},{RCK},"{c}|{b}")>0,SUMIFS({RCF},{RCK},"{c}|{b}"),0)' for c in codes_by[s] for b in big) + ")/1000"
    for s in segkeys[:4]] + [None],
    NUM1, "$M per firm [C]", "Ability to pay. Denominator excludes cells whose receipts are suppressed.")
mrow("Avg receipts per firm, 20–99 employees ($M)", [
    "=(" + census_sum(codes_by[s], RCR, ["20-49", "50-99"]) + ")/(" + "+".join(f'IF(SUMIFS({RCR},{RCK},"{c}|{b}")>0,SUMIFS({RCF},{RCK},"{c}|{b}"),0)' for c in codes_by[s] for b in ["20-49", "50-99"]) + ")/1000"
    for s in segkeys[:4]] + [None],
    NUM1, "$M per firm [C]", "Same metric for the SMB band only — removes the skew from very large firms.")
row_emp = mrow("Avg employees per firm, 20+ firms", [
    "=(" + census_sum(codes_by[s], RCEMP, big) + ")/(" + "+".join(f'IF(SUMIFS({RCEMP},{RCK},"{c}|{b}")>0,SUMIFS({RCF},{RCK},"{c}|{b}"),0)' for c in codes_by[s] for b in big) + ")"
    for s in segkeys[:4]] + [None],
    NUM, "employees [C]", "Average firm scale within the in-scope universe.")
r += 1
put(sp, f"A{r}", "Sales intensity & builder capacity (BLS OEWS May 2025, S15)", bold=True, fill=SFILL); r += 1


def oe(seg, occs):
    return "(" + "+".join(f'SUMIFS({ROEMP},{ROSEG},"{seg}",{ROOCC},"{o}")' for o in occs) + ")"


row_sales = mrow("Sales occupations (SOC 41) share of employment", [f"={oe(s,['41-0000'])}/{oe(s,['00-0000'])}" for s in segkeys], PCT,
                 "% of employment [C]", "Need proxy: how much of the workforce sells. Distribution ≈ 2.5× the economy-wide share.")
mrow("Wholesale & mfg sales reps per 1,000 employees (41-4011 + 41-4012)", [f"={oe(s,['41-4011','41-4012'])}/{oe(s,['00-0000'])}*1000" for s in segkeys], NUM1,
     "per 1,000 [C]", "Field/inside B2B reps — the users of prospecting, territory and enrichment workflows.")
mrow("Sales managers per 1,000 employees (11-2022)", [f"={oe(s,['11-2022'])}/{oe(s,['00-0000'])}*1000" for s in segkeys], NUM1, "per 1,000 [C]", "Likely economic buyers / champions for sales tooling.")
mrow("Marketing managers per 1,000 employees (11-2021)", [f"={oe(s,['11-2021'])}/{oe(s,['00-0000'])}*1000" for s in segkeys], NUM1, "per 1,000 [C]", "Marketing-led demand generation capacity.")
row_mra = mrow("Market research analysts & marketing specialists per 1,000 (13-1161)", [f"={oe(s,['13-1161'])}/{oe(s,['00-0000'])}*1000" for s in segkeys], NUM1,
               "per 1,000 [C]", "Closest occupational proxy for marketing-ops / data-savvy GTM staff who build Clay workflows.")
row_comp = mrow("Computer & mathematical occupations share of employment (SOC 15)", [f"={oe(s,['15-0000'])}/{oe(s,['00-0000'])}" for s in segkeys], PCT,
                "% of employment [C]", "'Builder capacity' proxy (tests H4): technical talent available to configure and maintain workflows.")
row_spf = mrow("Implied sales staff per 20+ firm", [f"=B{row_emp}*B{row_sales}", f"=C{row_emp}*C{row_sales}", f"=D{row_emp}*D{row_sales}", f"=E{row_emp}*E{row_sales}", None], NUM1,
               "people [C]", "Avg employees per 20+ firm × sales-occupation share. Clay usage (records, signals, sequences) scales with seller count — relevant to ACV in Phase 4.")
r += 1
put(sp, f"A{r}", "Indices vs all industries (=1.00)", bold=True, fill=SFILL); r += 1
mrow("Sales intensity index", [f"={c}{row_sales}/$F{row_sales}" for c in "BCDE"] + [1], X, "index [C]", "")
mrow("Builder capacity index (computer occupations)", [f"={c}{row_comp}/$F{row_comp}" for c in "BCDE"] + [1], X, "index [C]", "")
mrow("Marketing-ops talent index (13-1161)", [f"={c}{row_mra}/$F{row_mra}" for c in "BCDE"] + [1], X, "index [C]", "")
SP_ROWS = dict(f20=row_f20, sales=row_sales, comp=row_comp, spf=row_spf, mra=row_mra)
r += 1
put(sp, f"A{r}", "Caveat: OEWS counts employees in establishments classified to the industry; manufacturers' sales are partly done by independent reps (NAICS 4251) and distributors, which this metric correctly shows as low direct-sales intensity. Professional-services BD by billable staff is not captured in SOC 41.", wrap=True, border=False)
sp.merge_cells(f"A{r}:H{r}")
sp.row_dimensions[r].height = 30
sp.freeze_panes = "B5"

# ---------------------------------------------------------------- TAM_SAM_SOM
ts = wb.create_sheet("TAM_SAM_SOM")
title(ts, "TAM / SAM / SOM by segment — three scenarios shown side by side",
      "Low and High compound every input at its low/high value (outer bounds, not likely cases). SOM = year-3 ARR potential at uniform fair-share penetration — not a Clay forecast. [C]/[D]")
heads = ["Segment", "Size group", "Unique firms", "ACV ($)", "TAM ($)", "SAM filter 1: outbound motion", "SAM filter 2: CRM-ready",
         "SAM firms", "SAM ($)", "3-yr penetration", "SOM customers", "SOM ARR ($)"]
seg_f1 = {core[0][0]: "f1_dist", core[1][0]: "f1_mfg", core[2][0]: "f1_ps"}
seg_f2 = {core[0][0]: "f2_dist", core[1][0]: "f2_mfg", core[2][0]: "f2_ps"}
g_acv = {"SMB": "smb", "Mid-Market": "mm", "Enterprise": "ent"}
seg_code = {core[0][0]: "d", core[1][0]: "m", core[2][0]: "p"}
g_pen = {"SMB": "pen_smb", "Mid-Market": "pen_mm", "Enterprise": "pen_ent"}
scen_cols = {"Low": "D", "Base": "E", "High": "F"}
TS = {}  # (scen, seg) -> total row
r = 4
for scen in ["Base", "Low", "High"]:
    put(ts, f"A{r}", f"Scenario: {scen}", bold=True, border=False)
    r += 1
    hdr(ts, r, heads, widths=[28, 12, 12, 11, 14, 13, 12, 11, 14, 12, 11, 14])
    ts.row_dimensions[r].height = 30
    r += 1
    for seg, _, _ in core:
        first = r
        for g in ["SMB", "Mid-Market", "Enterprise"]:
            ac = scen_cols[scen]
            put(ts, f"A{r}", seg)
            put(ts, f"B{r}", g)
            put(ts, f"C{r}", f"=Firm_Counts!{FC_SEGCOL[seg]}{FC_GRP[g]}", font=GREEN, fmt=NUM)
            put(ts, f"D{r}", f"={aref('acv_' + seg_code[seg] + '_' + g_acv[g], ac)}", font=GREEN, fmt=USD)
            put(ts, f"E{r}", f"=C{r}*D{r}", fmt=USD)
            put(ts, f"F{r}", f"={aref(seg_f1[seg], ac)}", font=GREEN, fmt=PCT)
            put(ts, f"G{r}", f"={aref(seg_f2[seg], ac)}", font=GREEN, fmt=PCT)
            put(ts, f"H{r}", f"=C{r}*F{r}*G{r}", fmt=NUM)
            put(ts, f"I{r}", f"=H{r}*D{r}", fmt=USD)
            put(ts, f"J{r}", f"={aref(g_pen[g], ac)}", font=GREEN, fmt=PCT)
            put(ts, f"K{r}", f"=H{r}*J{r}", fmt=NUM1)
            put(ts, f"L{r}", f"=K{r}*D{r}", fmt=USD)
            r += 1
        put(ts, f"A{r}", f"{seg} — total", bold=True, fill=SFILL)
        put(ts, f"B{r}", None, fill=SFILL)
        for col in "CEHIKL":
            put(ts, f"{col}{r}", f"=SUM({col}{first}:{col}{r-1})", fmt=USD if col in "EIL" else (NUM1 if col == "K" else NUM), bold=True, fill=SFILL)
        for col in "DFGJ":
            put(ts, f"{col}{r}", None, fill=SFILL)
        TS[(scen, seg)] = r
        r += 1
    put(ts, f"A{r}", "All three segments (comparison only — segments are alternatives, not additive bets)", bold=True)
    for col in "CEHIKL":
        put(ts, f"{col}{r}", "=" + "+".join(f"{col}{TS[(scen, s[0])]}" for s in core), fmt=USD if col in "EIL" else (NUM1 if col == "K" else NUM), bold=True)
    TS[(scen, "ALL")] = r
    r += 3
ts.freeze_panes = "C4"

# ---------------------------------------------------------------- Sensitivity
se = wb.create_sheet("Sensitivity")
title(se, "Sensitivity — what moves SOM, and what would have to be true for the ranking to change",
      "Block A: one-at-a-time swings on Base SOM ARR. Block B: Distribution SOM grid vs Prof. Services base. Block C: segment-definition sensitivity (tests H1).")
# Block A tornado
put(se, "A4", "Block A — One-at-a-time sensitivity of Base SOM ARR ($), other inputs held at Base", bold=True, border=False)
hdr(se, 5, ["Input swung", "Segment", "SOM @ Low input", "SOM @ Base", "SOM @ High input", "Swing (High − Low)"],
    widths=[44, 30, 16, 16, 16, 16])
r = 6
fcol = FC_SEGCOL
for seg, _, _ in core:
    c = fcol[seg]
    firms = f"Firm_Counts!{c}{FC_GRP['SMB']}:{c}{FC_GRP['Enterprise']}"
    pen = {k: f"Assumptions!{k}{A['pen_smb']}:{k}{A['pen_ent']}" for k in "DEF"}
    acv = {k: f"Assumptions!{k}{A['acv_' + seg_code[seg] + '_smb']}:{k}{A['acv_' + seg_code[seg] + '_ent']}" for k in "DEF"}
    f1 = {k: aref(seg_f1[seg], k) for k in "DEF"}
    f2 = {k: aref(seg_f2[seg], k) for k in "DEF"}

    def som(a="E", p="E", x="E", y="E"):
        return f"=SUMPRODUCT({firms},{pen[p]},{acv[a]})*{f1[x]}*{f2[y]}"
    for label, lo, hi in [("ACV (all size bands)", som(a="D"), som(a="F")),
                          ("3-yr penetration (all size bands)", som(p="D"), som(p="F")),
                          ("SAM filter 1 — outbound motion", som(x="D"), som(x="F")),
                          ("SAM filter 2 — CRM readiness", som(y="D"), som(y="F"))]:
        put(se, f"A{r}", label)
        put(se, f"B{r}", seg)
        put(se, f"C{r}", lo, fmt=USD)
        put(se, f"D{r}", som(), fmt=USD)
        put(se, f"E{r}", hi, fmt=USD)
        put(se, f"F{r}", f"=E{r}-C{r}", fmt=USD)
        r += 1
    r += 1
check = r
put(se, f"A{check}", "Check: Block A base equals TAM_SAM_SOM base SOM (should be 0)")
put(se, f"B{check}", f"=ROUND(D6-TAM_SAM_SOM!L{TS[('Base', core[0][0])]},2)+ROUND(D11-TAM_SAM_SOM!L{TS[('Base', core[1][0])]},2)+ROUND(D16-TAM_SAM_SOM!L{TS[('Base', core[2][0])]},2)", fmt='0.00')

# Block B grid
gb = check + 3
put(se, f"A{gb}", "Block B — Distribution Base SOM ARR under multipliers on penetration (rows) × ACV (columns). Shaded = matches or exceeds Prof. Services Base SOM.", bold=True, border=False)
put(se, f"A{gb+1}", "Prof. Services Base SOM ARR (benchmark)")
put(se, f"B{gb+1}", f"=TAM_SAM_SOM!L{TS[('Base', core[2][0])]}", font=GREEN, fmt=USD)
put(se, f"A{gb+2}", "Distribution Base SOM ARR")
put(se, f"B{gb+2}", f"=TAM_SAM_SOM!L{TS[('Base', core[0][0])]}", font=GREEN, fmt=USD)
mults = [0.5, 0.75, 1.0, 1.5, 2.0, 3.0]
put(se, f"A{gb+4}", "Penetration multiplier ↓ / ACV multiplier →", bold=True, fill=SFILL)
for j, m in enumerate(mults):
    put(se, f"{L(2+j)}{gb+4}", m, font=BLUE, fmt=X, fill=SFILL)
for i, m in enumerate(mults):
    rr = gb + 5 + i
    put(se, f"A{rr}", m, font=BLUE, fmt=X, fill=SFILL)
    for j in range(len(mults)):
        col = L(2 + j)
        put(se, f"{col}{rr}", f"=$B${gb+2}*$A{rr}*{col}${gb+4}", fmt=USD)
grid_rng = f"B{gb+5}:{L(1+len(mults))}{gb+4+len(mults)}"
se.conditional_formatting.add(grid_rng, CellIsRule(operator="greaterThanOrEqual", formula=[f"$B${gb+1}"], fill=PatternFill("solid", fgColor="C6E0B4")))
for j in range(len(mults)):
    se.column_dimensions[L(2 + j)].width = 16
be = gb + 5 + len(mults)
put(se, f"A{be+1}", "Break-even combined multiplier for Distribution to match Prof. Services (penetration × ACV)")
put(se, f"B{be+1}", f"=B{gb+1}/B{gb+2}", fmt=X, bold=True)

# Block C definitions
dc = be + 4
put(se, f"A{dc}", "Block C — Segment-definition sensitivity: unique-ish firms with 20+ employees (core definition vs broader alternatives)", bold=True, border=False)
hdr(se, dc + 1, ["Definition", "Segment", "Firms 20+", "vs core", "Note"])
se.column_dimensions["E"].width = 16
defsens = [
    ("Core (model)", core[0][0], f"=Firm_Counts!D{FC_TOT}", "De-duplicated core NAICS 4235–4238."),
    ("Broad: + 4233 lumber/construction materials + 4234 prof. & commercial equipment", core[0][0],
     f"=Firm_Counts!D{FC_TOT}+" + "+".join(f'SUMIFS({RCF},{RCK},"{c}|{b}")' for c in ["4233", "4234"] for b in big) + f"*1", "Not de-duplicated (adds ≤3% overlap)."),
    ("Core (model)", core[1][0], f"=Firm_Counts!E{FC_TOT}", "De-duplicated core NAICS 326/332/333/335."),
    ("Broad: all manufacturing (sector 31-33, original brief)", core[1][0], "=" + "+".join(f'SUMIFS({RCF},{RCK},"31-33|{b}")' for b in big), "Sector total is already unique firms."),
    ("Core (model)", core[2][0], f"=Firm_Counts!F{FC_TOT}", "De-duplicated core NAICS 5413/5415/5416/5418."),
    ("Broad: all professional, scientific & technical services (sector 54)", core[2][0], "=" + "+".join(f'SUMIFS({RCF},{RCK},"54|{b}")' for b in big), "Adds legal, accounting, R&D, design, other."),
]
for i, (d, s, f, note) in enumerate(defsens):
    rr = dc + 2 + i
    put(se, f"A{rr}", d, wrap=True)
    put(se, f"B{rr}", s)
    put(se, f"C{rr}", f, fmt=NUM, font=GREEN)
    corerow = dc + 2 + (i // 2) * 2
    put(se, f"D{rr}", f"=C{rr}/C{corerow}", fmt=X)
    put(se, f"E{rr}", note, wrap=True)
    se.row_dimensions[rr].height = 28

# ---------------------------------------------------------------- Hypothesis_Tests
ht = wb.create_sheet("Hypothesis_Tests")
title(ht, "Hypothesis tests addressed in Phase 1", "Hypotheses from 01_Project_Charter.md §8. Verdicts are formula-driven where the test is quantitative.")
hdr(ht, 4, ["ID", "Hypothesis (abridged)", "Test metric", "Value", "Threshold", "Verdict", "Comment"],
    widths=[6, 46, 46, 14, 12, 26, 70])
dist, mfg, ps = (c[0] for c in core)
# H1
put(ht, "A5", "H1")
put(ht, "B5", "'B2B manufacturing suppliers' as written (all of 31-33) is too broad; a narrower industrial subset is more appropriate.", wrap=True)
put(ht, "C5", "Share of all-manufacturing 20+ firms captured by the core industrial subsectors", wrap=True)
put(ht, "D5", f"=Firm_Counts!E{FC_TOT}/Sensitivity!C{dc+5}", fmt=PCT)
put(ht, "E5", "n/a (qualitative)")
put(ht, "F5", "Supported (qualitative)", bold=True)
put(ht, "G5", "Core subsectors are only about two-fifths of manufacturing firms with 20+ employees; the remainder includes food, beverage, apparel, printing, furniture and wood — retail/CPG or local-market go-to-market, not engineered B2B selling. Industry heterogeneity is also visible in receipts per firm across core subsectors ($9M fabricated metal vs $38M electrical equipment).", wrap=True)
# H2
put(ht, "A6", "H2")
put(ht, "B6", "Distribution's mid-market pool (20–499) is ≥ manufacturing's and ≥ 50% of professional services'.", wrap=True)
put(ht, "C6", "Distribution 20–499 firms ÷ Prof. Services 20–499 firms", wrap=True)
put(ht, "D6", f"=(Firm_Counts!D{FC_GRP['SMB']}+Firm_Counts!D{FC_GRP['Mid-Market']})/(Firm_Counts!F{FC_GRP['SMB']}+Firm_Counts!F{FC_GRP['Mid-Market']})", fmt=PCT)
put(ht, "E6", 0.5, font=BLUE, fmt=PCT)
put(ht, "F6", '=IF(AND(D6>=E6,D7>=1),"Supported","Rejected")', bold=True)
put(ht, "G6", "Distribution is the SMALLEST segment by firm count. The case for distribution, if any, must rest on need intensity and per-firm value, not market breadth.", wrap=True)
put(ht, "C7", "Distribution 20–499 firms ÷ Manufacturing 20–499 firms", wrap=True)
put(ht, "D7", f"=(Firm_Counts!D{FC_GRP['SMB']}+Firm_Counts!D{FC_GRP['Mid-Market']})/(Firm_Counts!E{FC_GRP['SMB']}+Firm_Counts!E{FC_GRP['Mid-Market']})", fmt=PCT)
put(ht, "E7", 1.0, font=BLUE, fmt=PCT)
# H4
put(ht, "A8", "H4")
put(ht, "B8", "Binding constraint in industrial segments is builder capacity, not need.", wrap=True)
put(ht, "C8", "Builder-capacity index (computer occupations share vs all industries): Distribution / Manufacturing / Prof. Services", wrap=True)
put(ht, "D8", f'=TEXT(Segment_Profile!B{SP_ROWS["comp"]}/Segment_Profile!F{SP_ROWS["comp"]},"0.00")&" / "&TEXT(Segment_Profile!C{SP_ROWS["comp"]}/Segment_Profile!F{SP_ROWS["comp"]},"0.00")&" / "&TEXT(Segment_Profile!D{SP_ROWS["comp"]}/Segment_Profile!F{SP_ROWS["comp"]},"0.00")')
put(ht, "E8", "<1.0 = below avg")
put(ht, "F8", f'=IF(AND(Segment_Profile!B{SP_ROWS["comp"]}<Segment_Profile!F{SP_ROWS["comp"]},Segment_Profile!B{SP_ROWS["sales"]}>Segment_Profile!F{SP_ROWS["sales"]}),"Preliminary support (distribution)","Not supported")', bold=True)
put(ht, "G8", "Distribution combines the highest sales intensity with below-average technical talent — the 'high need, low builder capacity' profile. Manufacturing is low on both (weak need AND weak capacity). Professional services is the mirror image of distribution. Occupational proxies only; confirm with job-posting / partner evidence in Phase 2.", wrap=True)
# H6
put(ht, "A9", "H6")
put(ht, "B9", "Industrial mid-market ACV sits in the Growth-tier range (~$5K–$25K).", wrap=True)
put(ht, "C9", "Distribution Base ACV for SMB / Mid-Market (Assumptions A1-D, A2-D; usage-based)", wrap=True)
put(ht, "D9", f'=TEXT({aref("acv_d_smb","E")},"$#,##0")&" / "&TEXT({aref("acv_d_mm","E")},"$#,##0")')
put(ht, "E9", "$5K–$25K")
put(ht, "F9", f'=IF(AND({aref("acv_d_mm","E")}>=5000,{aref("acv_d_mm","E")}<=25000),"Supported (usage build)","Not supported")', bold=True)
put(ht, "G9", "v2: tested with the Phase 4 usage build at list prices. Distribution SMB sits at the Growth floor; Mid-Market ~$9.7K. Seller density lifts ACV only from Mid-Market up. No contract data is public, so this remains an estimate [C/D].", wrap=True)
for i in range(5, 10):
    ht.row_dimensions[i].height = 62

# ---------------------------------------------------------------- Executive_Summary
ex = wb.create_sheet("Executive_Summary", 1)
title(ex, "Executive summary — Phase 1 market sizing (Base case unless noted)",
      "All $ are annual. SOM = year-3 ARR potential at uniform fair-share penetration (scenario output, not a forecast). Segments are alternatives; do not add them.")
hdr(ex, 4, ["Metric", "A. Industrial Distribution", "B. Industrial Manufacturing", "C. B2B Professional Services", "Unit / class"],
    widths=[52, 22, 22, 22, 22])
r = 5
def exrow(label, f, fmt, unit, bold=False):
    global r
    put(ex, f"A{r}", label, bold=bold)
    for j, seg in enumerate([dist, mfg, ps]):
        put(ex, f"{L(2+j)}{r}", f(seg, j), fmt=fmt, font=GREEN, bold=bold)
    put(ex, f"E{r}", unit)
    r += 1
    return r - 1
e_f = exrow("Unique firms, 20+ employees", lambda s, j: f"=Firm_Counts!{L(4+j)}{FC_TOT}", NUM, "firms · [C] from [A]")
exrow("   of which SMB (20–99)", lambda s, j: f"=Firm_Counts!{L(4+j)}{FC_GRP['SMB']}", NUM, "firms")
exrow("   of which Mid-Market (100–499)", lambda s, j: f"=Firm_Counts!{L(4+j)}{FC_GRP['Mid-Market']}", NUM, "firms")
exrow("   of which Enterprise (500+)", lambda s, j: f"=Firm_Counts!{L(4+j)}{FC_GRP['Enterprise']}", NUM, "firms")
e_tam = exrow("TAM (Base)", lambda s, j: f"=TAM_SAM_SOM!E{TS[('Base', s)]}", USDM, "$/yr · [C]/[D]", bold=True)
e_sam = exrow("SAM (Base)", lambda s, j: f"=TAM_SAM_SOM!I{TS[('Base', s)]}", USDM, "$/yr · [C]/[D]", bold=True)
e_som = exrow("SOM — year-3 ARR potential (Base)", lambda s, j: f"=TAM_SAM_SOM!L{TS[('Base', s)]}", USDM, "$/yr · [D]", bold=True)
exrow("SOM range — compounded Low", lambda s, j: f"=TAM_SAM_SOM!L{TS[('Low', s)]}", USDM, "$/yr · outer bound")
exrow("SOM range — compounded High", lambda s, j: f"=TAM_SAM_SOM!L{TS[('High', s)]}", USDM, "$/yr · outer bound")
e_cust = exrow("SOM customers, year 3 (Base)", lambda s, j: f"=TAM_SAM_SOM!K{TS[('Base', s)]}", NUM, "customers · [D]")
exrow("SAM as % of TAM", lambda s, j: f"={L(2+j)}{e_sam}/{L(2+j)}{e_tam}", PCT, "% · [D]")
r += 1
put(ex, f"A{r}", "Structural profile", bold=True, fill=SFILL); r += 1
exrow("Sales occupations share of employment (all industries ≈ 8.6%)", lambda s, j: f"=Segment_Profile!{L(2+j)}{SP_ROWS['sales']}", PCT, "% · [C] from OEWS")
exrow("Computer & math occupations share (all industries ≈ 3.4%)", lambda s, j: f"=Segment_Profile!{L(2+j)}{SP_ROWS['comp']}", PCT, "% · [C] from OEWS")
exrow("Implied sales staff per 20+ firm", lambda s, j: f"=Segment_Profile!{L(2+j)}{SP_ROWS['spf']}", NUM1, "people · [C]")
r += 1
put(ex, f"A{r}", "Sanity checks", bold=True, fill=SFILL); r += 1
put(ex, f"A{r}", "Clay customers reported (Sep 2026, S03) — for scale")
put(ex, f"B{r}", 17000, font=BLUE, fmt=NUM)
put(ex, f"E{r}", "customers · [B] (trade-press report of company statement)")
cl = r; r += 1
put(ex, f"A{r}", "Base SOM customers as % of Clay's total customer count")
for j in range(3):
    put(ex, f"{L(2+j)}{r}", f"={L(2+j)}{e_cust}/$B${cl}", fmt=PCT)
put(ex, f"E{r}", "% · [C]")
r += 1
put(ex, f"A{r}", "Clay ARR, third-party estimate (Sep 2026, S23; was ~$150M in May 2026, S04) — for scale")
put(ex, f"B{r}", 200000000, font=BLUE, fmt=USDM)
put(ex, f"E{r}", "$/yr · [B] estimate")
arr = r; r += 1
put(ex, f"A{r}", "Base SOM ARR as % of Clay's estimated ARR")
for j in range(3):
    put(ex, f"{L(2+j)}{r}", f"={L(2+j)}{e_som}/$B${arr}", fmt=PCT)
put(ex, f"E{r}", "% · [C] (H8 context)")
r += 2
put(ex, f"A{r}", "Key takeaways (Phase 1)", bold=True, fill=SFILL); r += 1
takeaways = [
    "1. Industrial distribution is the smallest segment by firm count (H2 rejected): professional services offers nearly 3× the 20+ firm universe and the largest TAM, SAM and SOM at uniform penetration.",
    "2. Distribution is the most sales-intensive vertical in scope — sales occupations are ~22% of employment (~2.5× the economy-wide share) and an in-scope firm employs ~28 sales staff on average. Because so many distributors run active sales motions, distribution's SAM is roughly equal to manufacturing's despite having half the firms.",
    "3. v2 (usage-based ACV at list prices): seller density lifts distribution's Mid-Market ACV to ~$9.7K, but SMB sits at the $5,352 Growth floor in every segment, and manufacturing and professional services sit at the floor in Mid-Market too. v1's uniform $7.5K / $20K / $60K put TAM and SAM roughly 40–75% above the v2 usage-based values.",
    "4. Manufacturing is large in firm count but has the weakest direct-sales motion (~2.7% sales occupations); much of its selling runs through distributors and reps — a reason to reach manufacturers through the distribution channel rather than as a stand-alone vertical.",
    "5. At uniform fair-share penetration, every segment's 3-year SOM is ~0.6–1.2% of Clay's estimated ARR, and even the compounded High case stays below ~4%. Because penetration is assumed, SOM is a sizing of the prize, not a finding about right-to-win; any vertical program is an incremental option, not a pivot (H8).",
]
for t in takeaways:
    put(ex, f"A{r}", t, wrap=True)
    ex.merge_cells(f"A{r}:E{r}")
    ex.row_dimensions[r].height = 32
    r += 1
ex.freeze_panes = "B5"

# ---------------------------------------------------------------- Sources
so = wb.create_sheet("Sources")
title(so, "Sources used in this workbook", "Full register: Research_Sources.csv. Access date for all: 2026-10-18/19.")
hdr(so, 4, ["ID", "Source", "Publisher", "Reference period", "URL", "Used for"], widths=[6, 60, 28, 16, 70, 40])
srcs = [
    ("S01", "Clay pricing page", "Clay Labs Inc.", "Live, Oct 2026", "https://www.clay.com/pricing", "ACV floor (Growth tier list price)"),
    ("S02", "Clay second employee tender offer press release", "Business Wire", "Dec 2025–Jan 2026", "https://www.businesswire.com/news/home/20260128514638/en", "Enterprise NRR >200% (penetration rationale)"),
    ("S03", "Clay Series D coverage (17,000+ customers)", "Autobound / mrweb (secondary)", "Sep 2026", "https://mrweb.com/drno/news40306.htm", "Sanity check: customer count"),
    ("S04", "Clay crosses $150M/year (estimate)", "Sacra", "May 2026", "https://sacra.com/research/clay-crosses-150m-year/", "Superseded ARR estimate [B]"),
    ("S23", "Apollo.io at $265M ARR (incl. Clay ~$200M ARR estimate, Sep 2026)", "Sacra", "2026-10-14", "https://sacra.com/research/apollo-at-265m-arr/", "Sanity check: ARR scale [B]"),
    ("S14", "2022 Economic Census — Employment Size of Firms (EC2200SIZEEMPFIRM)", "U.S. Census Bureau", "2022", "https://data.census.gov/table/ECNSIZE2022.EC2200SIZEEMPFIRM", "All firm, establishment, employment and receipts counts"),
    ("S15", "OEWS May 2025 national industry-specific estimates", "U.S. Bureau of Labor Statistics", "May 2025", "https://data.bls.gov/oes/#/industry/423800/2025", "Sales intensity and builder-capacity proxies"),
    ("S16", "Clay pricing benchmarks (Vendr marketplace)", "Vendr", "Updated Feb 2026", "https://www.vendr.com/marketplace/clay", "ACV benchmarks: median $59,996; range $16.5K–$170K; n=96"),
    ("S58", "2022 Economic Census — Wholesale Trade: Detailed Type of Operation (EC2242TYPOP)", "U.S. Census Bureau", "2022", "https://www2.census.gov/programs-surveys/economic-census/data/2022/sector42/EC2242TYPOP.zip", "Manufacturers' sales branches and offices inside NAICS 423 (Firm_Counts Block 7)"),
    ("S17", "The State of CRM in Distribution (2022)", "Distribution Strategy Group", "2022", "https://archive.distributionstrategy.com/wp-content/uploads/2022/10/DSG-The-State-of-CRM-in-Distribution-Making-the-Software-Work-for-You.pdf", "Distributor CRM adoption mix (SAM filter 2)"),
]
for i, s in enumerate(srcs, start=5):
    for j, v in enumerate(s, start=1):
        put(so, f"{L(j)}{i}", v, wrap=True)
    so.row_dimensions[i].height = 30

# sheet order
order = ["Cover", "Executive_Summary", "Assumptions", "Firm_Counts", "Segment_Profile", "TAM_SAM_SOM", "Sensitivity",
         "Hypothesis_Tests", "Segment_Defs", "Raw_Census", "Raw_OEWS", "Raw_TYPOP", "Sources"]
wb._sheets = [wb[n] for n in order]
for wsx in wb.worksheets:
    wsx.sheet_properties.tabColor = {"Cover": "1F2A44", "Executive_Summary": "1F2A44", "Assumptions": "FFC000",
                                     "Raw_Census": "A6A6A6", "Raw_OEWS": "A6A6A6", "Raw_TYPOP": "A6A6A6", "Sources": "A6A6A6"}.get(wsx.title, "4472C4")
wb.save(OUT)
print("saved", OUT, "census rows", RC_LAST - 4, "oews rows", RO_LAST - 4)
