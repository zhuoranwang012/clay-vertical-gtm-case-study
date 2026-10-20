"""Export the 07_Financial_Model.xlsx outputs used by the deck charts to JSON.
Run after recalculating the workbook:  python3 appendix/export_deck_data.py
Then build the deck:  node appendix/build_deck.js appendix/deck_data_from_model.json Executive_Strategy_Deck.pptx"""
import json, os
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
wb = load_workbook(os.path.join(BASE, "07_Financial_Model.xlsx"), data_only=True)
s = wb["Scenarios"]
out, blk = {}, None
for r in range(30, 62):
    a = s.cell(r, 1).value
    if not isinstance(a, str):
        continue
    if a.startswith("ARR trajectory"):
        blk = "arr"; continue
    if a.startswith("Cumulative program P&L by"):
        blk = "cum"; continue
    if blk and (a.startswith("Partner-led") or a.startswith("Sales-led")):
        out.setdefault(blk, {})[a] = [s.cell(r, c).value for c in range(4, 16)]
out["scen"] = {s.cell(r, 1).value: [s.cell(r, c).value for c in range(4, 11)] for r in range(5, 21)}
out["scen_cols"] = [s.cell(4, c).value for c in range(4, 11)]
out["dec"] = {s.cell(r, 1).value: [s.cell(r, c).value for c in range(4, 8)] for r in range(26, 37) if s.cell(r, 1).value}
t = wb["Sensitivity"]
out["tornado"] = [(t.cell(r, 1).value, t.cell(r, 6).value, t.cell(r, 4).value, t.cell(r, 5).value) for r in range(5, 18)]
out["base_tv"] = t.cell(18, 4).value
json.dump(out, open(os.path.join(BASE, "appendix", "deck_data_from_model.json"), "w"), indent=1, default=str)
print("wrote appendix/deck_data_from_model.json")
